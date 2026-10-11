"""Exchange-reference candidates to be confirmed separately against Futu.

HKEX supplies the trading currency for each counter. Nasdaq Trader's symbol
directories have no currency field: USD below is the standard US listed-stock/
ETF market convention, explicitly recorded as such. These are current listing
references, not proof of Futu coverage or a point-in-time security master.
"""

from __future__ import annotations

import csv
from datetime import datetime
from io import BytesIO, StringIO
import re
from typing import Iterable
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import openpyxl

from studio_market.numeric.providers.http_client import PublicResponse, get_public_bytes


HKEX_LISTINGS_URL = "https://www.hkex.com.hk/eng/services/trading/securities/securitieslists/ListOfSecurities.xlsx"
NASDAQ_LISTED_URL = "https://www.nasdaqtrader.com/dynamic/SymDir/nasdaqlisted.txt"
OTHER_LISTED_URL = "https://www.nasdaqtrader.com/dynamic/SymDir/otherlisted.txt"
_US_EXCHANGES = {"A": "XASE", "N": "XNYS", "P": "ARCX", "Z": "BATS"}
_HK_EQUITY_SUBTYPES = {
    "Equity Securities (Main Board)", "Equity Securities (GEM)", "Depositary Receipts",
}
_US_SYMBOL = re.compile(r"[A-Z]{1,5}(?:\.[A-Z])?\Z")
_COMMON_EQUITY = re.compile(
    r"\b(?:common|ordinary)\s+(?:stock|shares?)\b"
    r"|\bAmerican\s+Depositar[yi]\s+(?:shares?|receipts?)\b|\bADRs?\b",
    re.IGNORECASE,
)
_NON_COMMON = re.compile(
    r"\b(?:preferred|preference)\s+(?:stock|shares?)\b"
    r"|\b(?:warrants?|rights?|units?|debentures?|fund|portfolio)\b"
    r"|\bnotes?\s+due\b|\b(?:senior|subordinated)\s+(?:notes?|bonds?)\b"
    r"|\bclosed[- ]end\b",
    re.IGNORECASE,
)
_EXCHANGE_TRADED_NOTE = re.compile(r"\bETNs?\b|\bexchange[- ]traded\s+notes?\b", re.IGNORECASE)


class ListingReferenceError(ValueError):
    pass


def _candidate(*, symbol: str, ticker: str, name: str, instrument_type: str,
               exchange: str, currency: str, source: str, endpoint: str,
               currency_basis: str) -> dict[str, str]:
    return {
        "symbol": symbol,
        "exchange_ticker": ticker,
        "futu_stock_type": "ETF" if instrument_type == "etf" else "STOCK",
        "instrument_type": instrument_type,
        "currency": currency,
        "currency_basis": currency_basis,
        "exchange_code": exchange,
        "name": name,
        "source_dataset": source,
        "source_endpoint": endpoint,
    }


def _unique(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    seen = set()
    for row in rows:
        if row["symbol"] in seen:
            raise ListingReferenceError("Listing references contain duplicate security symbols.")
        seen.add(row["symbol"])
    if not rows:
        raise ListingReferenceError("Listing reference contains no supported stock or ETF candidates.")
    return rows


def parse_hkex_listings(body: bytes) -> list[dict[str, str]]:
    """Read HKEX's explicit stock/ETF classes and each counter's own currency.

    REITs, warrants, debt, investment companies, trading-only securities and the
    separately classified Leveraged and Inverse products are outside this list.
    """
    try:
        workbook = openpyxl.load_workbook(BytesIO(body), read_only=True, data_only=True)
    except (ValueError, OSError, KeyError, BadZipFile, ParseError, openpyxl.utils.exceptions.InvalidFileException):
        raise ListingReferenceError("HKEX listing reference is not a readable XLSX workbook.") from None
    try:
        if "ListOfSecurities" not in workbook.sheetnames:
            raise ListingReferenceError("HKEX listing workbook is missing ListOfSecurities.")
        sheet = workbook["ListOfSecurities"]
        # The current official workbook declares A1:R8 despite containing over
        # 17,000 rows. Using its dimension would silently truncate the directory.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        required = {"Stock Code", "Name of Securities", "Category", "Sub-Category", "Trading Currency"}
        columns = None
        for values in iterator:
            labels = [str(value).strip() if value is not None else "" for value in values]
            if required.issubset(labels):
                if any(labels.count(label) != 1 for label in required):
                    raise ListingReferenceError("HKEX listing workbook has duplicate column labels.")
                columns = {label: labels.index(label) for label in required}
                break
        if columns is None:
            raise ListingReferenceError("HKEX listing workbook is missing required columns.")
        rows = []
        for values in iterator:
            if not any(value is not None and str(value).strip() for value in values):
                continue
            row = {key: values[index] if index < len(values) else None for key, index in columns.items()}
            category = str(row["Category"] or "").strip()
            subtype = str(row["Sub-Category"] or "").strip()
            if category == "Equity" and subtype in _HK_EQUITY_SUBTYPES:
                instrument_type = "equity"
            elif category == "Exchange Traded Products" and subtype == "Exchange Traded Funds":
                instrument_type = "etf"
            else:
                continue
            code_value = row["Stock Code"]
            if isinstance(code_value, float) and code_value.is_integer():
                code_value = int(code_value)
            code = str(code_value).strip()
            name = str(row["Name of Securities"] or "").strip()
            currency = str(row["Trading Currency"] or "").strip()
            currency = "CNY" if currency == "RMB" else currency
            if (not re.fullmatch(r"[0-9]{1,5}", code) or not 1 <= int(code) <= 99999
                    or not name or not re.fullmatch(r"[A-Z]{3}", currency)):
                raise ListingReferenceError("HKEX stock/ETF row has invalid code, name or trading currency.")
            rows.append(_candidate(
                symbol=f"HK.{int(code):05d}", ticker=f"{int(code):04d}.HK", name=name,
                instrument_type=instrument_type, exchange="XHKG", currency=currency,
                source="hkex_securities_list", endpoint=HKEX_LISTINGS_URL,
                currency_basis="exchange_directory",
            ))
        return _unique(rows)
    finally:
        workbook.close()


def parse_nasdaq_listings(body: bytes, *, other_listed: bool = False) -> list[dict[str, str]]:
    """Select standard US symbols without translating ambiguous suffixes.

    ETF=N is not sufficient evidence of ordinary shares: preferred securities,
    rights, units, notes, funds and uncertain descriptions are excluded. Exact
    class symbols such as BRK.B are preserved for Futu's subsequent confirmation.
    """
    try:
        text = body.decode("utf-8-sig")
    except UnicodeError:
        raise ListingReferenceError("Nasdaq listing reference is not UTF-8 text.") from None
    try:
        records = list(csv.reader(StringIO(text), delimiter="|", strict=True))
    except csv.Error:
        raise ListingReferenceError("Nasdaq listing reference is not valid delimited text.") from None
    records = [record for record in records if any(value.strip() for value in record)]
    if len(records) < 3:
        raise ListingReferenceError("Nasdaq listing reference is empty or incomplete.")
    footer = records[-1]
    if not footer[0].startswith("File Creation Time: ") or any(value.strip() for value in footer[1:]):
        raise ListingReferenceError("Nasdaq listing reference is missing its completion timestamp.")
    try:
        datetime.strptime(footer[0].removeprefix("File Creation Time: ").strip(), "%m%d%Y%H:%M")
    except ValueError:
        raise ListingReferenceError("Nasdaq listing reference has an invalid completion timestamp.") from None
    header = [value.strip() for value in records[0]]
    symbol_field = "ACT Symbol" if other_listed else "Symbol"
    required = {symbol_field, "Security Name", "Test Issue", "ETF"}
    if other_listed:
        required.add("Exchange")
    if not required.issubset(header) or len(header) != len(set(header)):
        raise ListingReferenceError("Nasdaq listing reference is missing unique required columns.")
    rows = []
    for values in records[1:-1]:
        if len(values) != len(header):
            raise ListingReferenceError("Nasdaq listing reference contains a malformed row.")
        row = dict(zip(header, (value.strip() for value in values)))
        if row["Test Issue"] not in {"Y", "N"} or row["ETF"] not in {"Y", "N"}:
            raise ListingReferenceError("Nasdaq listing reference has an invalid Test Issue or ETF flag.")
        if row["Test Issue"] == "Y":
            continue
        if row.get("NextShares") == "Y":
            continue
        exchange = _US_EXCHANGES.get(row["Exchange"]) if other_listed else "XNAS"
        symbol, name = row[symbol_field], row["Security Name"]
        if exchange is None or not _US_SYMBOL.fullmatch(symbol):
            continue
        if not name:
            raise ListingReferenceError("Nasdaq listing reference contains an unnamed security.")
        if row["ETF"] == "Y":
            if _EXCHANGE_TRADED_NOTE.search(name):
                continue
            instrument_type = "etf"
        else:
            if not _COMMON_EQUITY.search(name) or _NON_COMMON.search(name):
                continue
            # A trust without a corporate issuer description may be a closed-end
            # fund, not an ordinary operating-company share. Leave it unclassified.
            if (re.search(r"\btrust\b", name, re.IGNORECASE)
                    and not re.search(r"\b(?:corporation|corp|inc|limited|ltd|plc)\b", name, re.IGNORECASE)):
                continue
            instrument_type = "equity"
        rows.append(_candidate(
            symbol=f"US.{symbol}", ticker=symbol, name=name, instrument_type=instrument_type,
            exchange=exchange, currency="USD", currency_basis="us_listing_convention",
            source="nasdaq_other_listed" if other_listed else "nasdaq_listed",
            endpoint=OTHER_LISTED_URL if other_listed else NASDAQ_LISTED_URL,
        ))
    return _unique(rows)


def load_listings(markets: Iterable[str] = ("hk", "us")) -> tuple[list[dict[str, str]], list[PublicResponse]]:
    """Download current exchange references; do not call Futu or persist data."""
    selected = tuple(dict.fromkeys(markets))
    if not selected or set(selected) - {"hk", "us"}:
        raise ValueError("Listing-reference markets must be hk and/or us.")
    rows: list[dict[str, str]] = []
    captures: list[PublicResponse] = []
    for market in selected:
        if market == "hk":
            capture = get_public_bytes(HKEX_LISTINGS_URL)
            captures.append(capture)
            rows.extend(parse_hkex_listings(capture.body))
        else:
            for endpoint, other_listed in ((NASDAQ_LISTED_URL, False), (OTHER_LISTED_URL, True)):
                capture = get_public_bytes(endpoint)
                captures.append(capture)
                rows.extend(parse_nasdaq_listings(capture.body, other_listed=other_listed))
    return _unique(rows), captures
