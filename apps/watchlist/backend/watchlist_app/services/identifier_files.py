from __future__ import annotations

import csv
import re
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_IDENTIFIER_FILE_BYTES = 5 * 1024 * 1024
MAX_IDENTIFIERS = 2_000
IDENTIFIER_HEADERS = frozenset(
    {"identifier", "ticker", "isin", "tickerisin", "instrumentid"}
)


def _normalized_header(value: object) -> str:
    return re.sub(r"[\s_\-/]+", "", str(value or "").strip().lower())


def _identifiers_from_rows(rows: list[list[object]]) -> list[str]:
    non_empty_rows = [
        row for row in rows if any(str(value or "").strip() for value in row)
    ]
    if not non_empty_rows:
        raise ValueError("The file contains no identifier rows.")
    header = [_normalized_header(value) for value in non_empty_rows[0]]
    identifier_index = next(
        (index for index, value in enumerate(header) if value in IDENTIFIER_HEADERS),
        None,
    )
    data_rows = non_empty_rows[1:] if identifier_index is not None else non_empty_rows
    column_index = identifier_index or 0
    identifiers: list[str] = []
    for row in data_rows:
        if column_index >= len(row):
            continue
        identifier = str(row[column_index] or "").strip()
        if identifier and identifier not in identifiers:
            identifiers.append(identifier)
        if len(identifiers) > MAX_IDENTIFIERS:
            raise ValueError("The file exceeds the 2,000 identifier limit.")
    if not identifiers:
        raise ValueError(
            "No valid rows found. Expected an Identifier, Ticker, or ISIN column."
        )
    return identifiers


def _parse_text(content: bytes) -> list[list[object]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("CSV, TSV, and text files must use UTF-8 encoding.") from error
    first_line = next((line for line in text.splitlines() if line.strip()), "")
    delimiter = "\t" if "\t" in first_line else ","
    return [list(row) for row in csv.reader(StringIO(text, newline=""), delimiter=delimiter)]


def _parse_xlsx(content: bytes) -> list[list[object]]:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ValueError("The Excel file is not a valid .xlsx workbook.") from error
    try:
        worksheet = workbook.active
        rows: list[list[object]] = []
        for row_number, cells in enumerate(worksheet.iter_rows(), start=1):
            if row_number > MAX_IDENTIFIERS + 1:
                raise ValueError("The file exceeds the 2,000 identifier limit.")
            values: list[object] = []
            for cell in cells:
                if cell.data_type == "f":
                    raise ValueError(
                        f"{worksheet.title}!{cell.coordinate} contains a formula; "
                        "identifier imports require literal values."
                    )
                if cell.data_type == "e":
                    raise ValueError(
                        f"{worksheet.title}!{cell.coordinate} contains an Excel error value."
                    )
                values.append(cell.value)
            rows.append(values)
        return rows
    finally:
        workbook.close()


def parse_identifier_file(filename: str | None, content: bytes) -> list[str]:
    if not content:
        raise ValueError("Identifier file content is required.")
    if len(content) > MAX_IDENTIFIER_FILE_BYTES:
        raise ValueError("Identifier file exceeds the 5 MB limit.")
    extension = Path(str(filename or "")).suffix.lower()
    if extension in {".csv", ".tsv", ".txt"}:
        return _identifiers_from_rows(_parse_text(content))
    if extension == ".xlsx":
        return _identifiers_from_rows(_parse_xlsx(content))
    raise ValueError("Identifier files must use .csv, .tsv, .txt, or .xlsx.")
