from __future__ import annotations

import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from portfolio_app.services.transaction_csv import (
    DERIVATIVE_DEFINITION_COLUMNS,
    IMPORT_COLUMNS,
    parse_transaction_csv,
    render_transaction_csv,
    transaction_export_rows,
)
from portfolio_app.services.transaction_import import (
    ASSET_TYPE_VALUES,
    TRANSACTION_ACTIONS,
)
from portfolio_app.services.transaction_xlsx import (
    CURRENCY_VALUES,
    EXAMPLES_SHEET_NAME,
    FIELD_GUIDE_SHEET_NAME,
    INSTRUCTIONS_SHEET_NAME,
    LISTS_SHEET_NAME,
    TEMPLATE_INPUT_LAST_ROW,
    TEMPLATE_EXAMPLE_ROWS,
    TRANSACTION_SHEET_NAME,
    XLSX_MEDIA_TYPE,
    render_transaction_xlsx,
    render_transaction_xlsx_template,
    transaction_xlsx_to_csv,
)


def _deposit_record(external_reference: str) -> dict[str, object]:
    return {
        "transaction_sequence": 9001,
        "transaction_type": "deposit",
        "trade_date": "2026-05-25",
        "account_id": "cash-usd-main",
        "gross_amount": 25.12345678,
        "source_gross_amount": "25.12345678",
        "fees": 0,
        "taxes": 0,
        "fee_category": "unknown",
        "currency": "USD",
        "source_system": "transaction_file_test",
        "external_reference": external_reference,
        "note": "=literal note",
    }


def test_transaction_xlsx_is_typed_safe_and_round_trip_importable() -> None:
    content = render_transaction_xlsx([_deposit_record("xlsx-round-trip-1")])
    workbook = load_workbook(BytesIO(content), data_only=False)
    worksheet = workbook[TRANSACTION_SHEET_NAME]

    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == f"A1:{get_column_letter(len(IMPORT_COLUMNS))}2"
    assert worksheet.cell(1, 1).value == "asset_type"
    assert worksheet.cell(2, 1).value == "cash"
    assert worksheet.cell(2, 2).value == "deposit"
    note_cell = worksheet.cell(2, IMPORT_COLUMNS.index("note") + 1)
    assert note_cell.value == "=literal note"
    assert note_cell.data_type == "s"
    workbook.close()

    _headers, rows = parse_transaction_csv(transaction_xlsx_to_csv(content))
    assert rows[0].errors == ()
    assert rows[0].transaction is not None
    assert str(rows[0].transaction.gross_amount) == "25.12345678"
    assert rows[0].transaction.note == "=literal note"


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
@pytest.mark.parametrize("estimated", [False, True])
def test_file_round_trip_preserves_known_versus_estimated_trade_time(file_format, estimated):
    record = {**_deposit_record("trade-time-roundtrip"), "trade_time": "12:00", "trade_time_is_estimated": estimated}
    csv_content = render_transaction_csv([record]) if file_format == "csv" else transaction_xlsx_to_csv(render_transaction_xlsx([record]))
    _, rows = parse_transaction_csv(csv_content)
    assert rows[0].errors == ()
    assert rows[0].transaction.trade_time == (None if estimated else "12:00")


def test_transaction_xlsx_rejects_formula_cells() -> None:
    workbook = load_workbook(BytesIO(render_transaction_xlsx_template()))
    worksheet = workbook[TRANSACTION_SHEET_NAME]
    values = {
        "asset_type": "cash",
        "transaction_action": "deposit",
        "trade_date": "2026-05-25",
        "account_id": "cash-usd-main",
        "gross_amount": 25,
        "currency": "USD",
        "note": "=1+1",
    }
    for column, value in values.items():
        worksheet.cell(2, IMPORT_COLUMNS.index(column) + 1).value = value
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    with pytest.raises(ValueError, match="contains a formula"):
        transaction_xlsx_to_csv(output.getvalue())


def test_transaction_xlsx_template_guides_manual_entry_without_importing_examples() -> (
    None
):
    content = render_transaction_xlsx_template()
    workbook = load_workbook(BytesIO(content), data_only=False)

    assert workbook.sheetnames == [
        INSTRUCTIONS_SHEET_NAME,
        TRANSACTION_SHEET_NAME,
        FIELD_GUIDE_SHEET_NAME,
        EXAMPLES_SHEET_NAME,
        LISTS_SHEET_NAME,
    ]
    assert workbook.active.title == INSTRUCTIONS_SHEET_NAME

    worksheet = workbook[TRANSACTION_SHEET_NAME]
    assert worksheet.max_row == 1
    assert worksheet.max_column == len(IMPORT_COLUMNS)
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == (
        f"A1:{get_column_letter(len(IMPORT_COLUMNS))}{TEMPLATE_INPUT_LAST_ROW}"
    )
    assert worksheet.cell(1, 1).comment is not None
    assert "资产类型" in str(worksheet.cell(1, 1).comment.text)

    named_ranges = set(workbook.defined_names)
    assert {
        "AssetTypeValues",
        "security_actions",
        "fcn_actions",
        "option_actions",
        "cash_actions",
        "FeeCategoryValues",
        "CurrencyValues",
        "OptionTypeValues",
    } <= named_ranges
    validations = worksheet.data_validations.dataValidation
    asset_type_validation = next(
        item
        for item in validations
        if str(item.sqref) == f"A2:A{TEMPLATE_INPUT_LAST_ROW}"
    )
    transaction_action_validation = next(
        item
        for item in validations
        if str(item.sqref) == f"B2:B{TEMPLATE_INPUT_LAST_ROW}"
    )
    currency_column = get_column_letter(IMPORT_COLUMNS.index("currency") + 1)
    currency_validation = next(
        item
        for item in validations
        if str(item.sqref)
        == f"{currency_column}2:{currency_column}{TEMPLATE_INPUT_LAST_ROW}"
    )
    assert asset_type_validation.type == "list"
    assert asset_type_validation.formula1 == "AssetTypeValues"
    assert transaction_action_validation.type == "list"
    assert transaction_action_validation.formula1 == 'INDIRECT($A2&"_actions")'
    assert currency_validation.type == "list"
    assert currency_validation.formula1 == "CurrencyValues"

    lists = workbook[LISTS_SHEET_NAME]
    for asset_type, column in {"security": 4, "fcn": 7, "option": 10, "cash": 13}.items():
        for row_index, action in enumerate(TRANSACTION_ACTIONS[asset_type], start=2):
            assert lists.cell(row_index, column).value == action
            assert lists.cell(row_index, column + 1).value, f"Missing template guidance for {action}"

    examples = workbook[EXAMPLES_SHEET_NAME]
    assert examples.cell(1, 1).value == "资产类型与交易动作示例（仅供参考，不会上传）"
    assert examples.cell(2, 1).value is not None
    assert [
        examples.cell(4, index).value for index in range(2, len(IMPORT_COLUMNS) + 2)
    ] == list(IMPORT_COLUMNS)
    example_scenarios = {
        str(examples.cell(row_index, 1).value)
        for row_index in range(5, examples.max_row + 1)
    }
    example_asset_types = {
        str(examples.cell(row_index, 2).value)
        for row_index in range(5, examples.max_row + 1)
    }
    example_actions = {
        str(examples.cell(row_index, 3).value)
        for row_index in range(5, examples.max_row + 1)
    }
    for row_index in range(5, examples.max_row + 1):
        for column_index in range(1, len(IMPORT_COLUMNS) + 1):
            worksheet.cell(row_index, column_index).value = examples.cell(
                row_index,
                column_index + 1,
            ).value
    example_workbook = BytesIO()
    workbook.save(example_workbook)
    workbook.close()

    _headers, rows = parse_transaction_csv(
        transaction_xlsx_to_csv(example_workbook.getvalue())
    )
    assert len(rows) >= 40
    assert all(row.errors == () for row in rows)
    assert {
        str(row.transaction.currency) for row in rows if row.transaction is not None
    } == {"USD", "HKD", "CNY"}
    assert set(CURRENCY_VALUES) == {"USD", "HKD", "CNY", "EUR", "GBP", "CHF"}
    assert example_asset_types == set(ASSET_TYPE_VALUES)
    assert {
        "knock_in_close",
        "knock_out_close",
        "maturity_close",
        "expire_long",
        "cash_settle_long",
        "expire_written",
        "cash_settle_written",
    } <= example_actions
    assert {
        (str(row["asset_type"]), str(row["transaction_action"]))
        for row in TEMPLATE_EXAMPLE_ROWS
    } == {
        (asset_type, action)
        for asset_type, actions in TRANSACTION_ACTIONS.items()
        for action in actions
    } - {("cash", "transfer_in"), ("security", "transfer_in")}
    assert {
        "股票｜USD 股票买入",
        "ETF｜HKD ETF 卖出",
        "公募基金｜CNY 申购确认",
        "私募基金｜CNY 赎回确认",
        "证券｜HKD 独立交易费用",
        "证券｜HKD 独立税费",
        "Option｜CNY 合约期初多头",
        "Option｜HKD 独立费用",
        "Option｜HKD 独立税费",
        "FCN｜新合约进入",
        "FCN｜HKD 合约期初持仓",
        "FCN｜提前退出",
        "FCN｜票息收入",
        "FCN｜独立费用",
        "FCN｜独立税费",
        "FCN｜敲入后实际现金兑付",
    } <= example_scenarios
    assert not any("Option 实物行权拆分" in scenario for scenario in example_scenarios)
    assert not any("Option 空头 Call 指派拆分" in scenario for scenario in example_scenarios)


def test_transaction_csv_rejects_old_protocol_and_cross_asset_fields() -> None:
    with pytest.raises(ValueError, match="Unsupported CSV columns: transaction_type"):
        parse_transaction_csv(
            "transaction_type,trade_date,account_id,gross_amount,currency\n"
            "deposit,2026-01-01,cash-usd-main,10,USD\n"
        )

    _headers, invalid_rows = parse_transaction_csv(
        "asset_type,transaction_action,trade_date,account_id,instrument_id,"
        "gross_amount,currency\n"
        "cash,deposit,2026-01-01,cash-usd-main,equity-us-abbv,10,USD\n"
    )
    assert invalid_rows[0].errors == (
        "Cash actions must not carry security or derivative fields.",
    )

    _headers, invalid_rows = parse_transaction_csv(
        "asset_type,transaction_action,trade_date,account_id,"
        "derivative_contract_id,derivative_contract_name,gross_amount,currency\n"
        "option,cash_settle_long,2026-01-01,options-us,option-1,New option,10,USD\n"
    )
    assert invalid_rows[0].errors == (
        "A new OPTION contract can only be defined on: "
        "buy_to_open, opening_balance, opening_written, sell_to_open.",
    )

    _headers, invalid_rows = parse_transaction_csv(
        "asset_type,transaction_action,trade_date,account_id,instrument_id,"
        "gross_amount,currency\n"
        "security,coupon,2026-01-01,broker-us-core,equity-us-abbv,10,USD\n"
    )
    assert "transaction_action 'coupon' is not supported for asset_type 'security'" in (
        invalid_rows[0].errors[0]
    )

    _headers, invalid_rows = parse_transaction_csv(
        "asset_type,transaction_action,trade_date,account_id,gross_amount,currency\n"
        "security,buy,2026-01-01,broker-us-core,10,USD\n"
        "fcn,coupon,2026-01-01,broker-us-fcn,10,USD\n"
    )
    assert invalid_rows[0].errors == ("Security actions require instrument_id.",)
    assert invalid_rows[1].errors == ("FCN actions require derivative_contract_id.",)


def test_transaction_transfer_actions_resolve_direction_from_the_selected_side() -> None:
    _headers, rows = parse_transaction_csv(
        "asset_type,transaction_action,trade_date,account_id,"
        "counterparty_account_id,gross_amount,currency\n"
        "cash,transfer_out,2026-01-01,cash-a,cash-b,10,USD\n"
        "cash,transfer_in,2026-01-02,cash-b,cash-a,10,USD\n"
    )

    assert all(row.errors == () and row.internal_transfer is not None for row in rows)
    assert [
        (row.internal_transfer.from_account_id, row.internal_transfer.to_account_id)
        for row in rows
        if row.internal_transfer is not None
    ] == [("cash-a", "cash-b"), ("cash-a", "cash-b")]


def test_transaction_file_preview_checks_asset_account_category(client) -> None:
    response = client.post(
        "/api/portfolios/investment-studio/transactions/csv/preview",
        json={
            "csv_text": (
                "asset_type,transaction_action,trade_date,account_id,instrument_id,"
                "quantity,price,gross_amount,currency\n"
                "security,buy,2026-05-25,cash-usd-main,equity-us-abbv,1,10,10,USD\n"
            )
        },
    )

    assert response.status_code == 200
    preview = response.json()
    assert preview["error_count"] == 1
    assert preview["rows"][0]["errors"] == ["Buy requires securities_account."]


def test_transaction_export_keeps_derivative_definitions_on_opening_rows_only() -> None:
    option_terms = {
        "underlying_instrument_id": "equity-us-abbv",
        "option_type": "call",
        "expiry_date": "2026-12-18",
        "strike": 220,
        "contract_multiplier": 100,
    }
    records = [
        {
            "transaction_type": "option_write",
            "trade_date": "2026-05-01",
            "account_id": "options-us",
            "settlement_cash_account_id": "cash-usd-main",
            "derivative_contract_id": "option-short-call-1",
            "derivative_contract": {
                "contract_name": "ABBV Dec 220 Call",
                "contract_type": "option",
                "external_reference": "BROKER-OPTION-001",
                "terms": option_terms,
            },
            "quantity": 1,
            "price": 5,
            "gross_amount": 500,
            "fees": 0,
            "taxes": 0,
            "fee_category": "unknown",
            "currency": "USD",
        },
        {
            "transaction_type": "lifecycle_event",
            "lifecycle_event_type": "option_writer_cash_settlement",
            "trade_date": "2026-12-18",
            "account_id": "options-us",
            "settlement_cash_account_id": "cash-usd-main",
            "derivative_contract_id": "option-short-call-1",
            "derivative_contract": {
                "contract_name": "ABBV Dec 220 Call",
                "contract_type": "option",
                "external_reference": "BROKER-OPTION-001",
                "terms": option_terms,
            },
            "quantity": 1,
            "gross_amount": 1000,
            "fees": 0,
            "taxes": 0,
            "fee_category": "unknown",
            "currency": "USD",
        },
    ]

    exported_rows = transaction_export_rows(records)

    assert [row["transaction_action"] for row in exported_rows] == [
        "sell_to_open",
        "cash_settle_written",
    ]
    assert exported_rows[0]["derivative_contract_name"] == "ABBV Dec 220 Call"
    assert all(exported_rows[1][column] == "" for column in DERIVATIVE_DEFINITION_COLUMNS)

    _headers, parsed_rows = parse_transaction_csv(render_transaction_csv(records))
    assert all(row.errors == () for row in parsed_rows)


def test_transaction_file_api_supports_csv_and_xlsx_with_one_validation_path(
    client,
) -> None:
    record = _deposit_record("xlsx-api-import-1")
    csv_content = render_transaction_csv([record]).encode("utf-8")
    xlsx_content = render_transaction_xlsx([record])

    csv_preview_response = client.post(
        "/api/portfolios/investment-studio/transactions/files/preview",
        files={"file": ("transactions.csv", csv_content, "text/csv")},
        data={"default_source_system": "portfolio_file_upload"},
    )
    assert csv_preview_response.status_code == 200, csv_preview_response.text
    assert csv_preview_response.json()["error_count"] == 0

    xlsx_preview_response = client.post(
        "/api/portfolios/investment-studio/transactions/files/preview",
        files={"file": ("transactions.xlsx", xlsx_content, XLSX_MEDIA_TYPE)},
        data={"default_source_system": "portfolio_file_upload"},
    )
    assert xlsx_preview_response.status_code == 200, xlsx_preview_response.text
    preview = xlsx_preview_response.json()
    assert preview["valid_count"] == 1
    assert preview["error_count"] == 0

    import_response = client.post(
        "/api/portfolios/investment-studio/transactions/files/import",
        headers={"Idempotency-Key": "xlsx-api-import-1"},
        files={"file": ("transactions.xlsx", xlsx_content, XLSX_MEDIA_TYPE)},
        data={
            "default_source_system": "portfolio_file_upload",
            "preview_digest": preview["preview_digest"],
        },
    )
    assert import_response.status_code == 200, import_response.text
    assert import_response.json()["created_count"] == 1

    export_response = client.get("/api/portfolios/investment-studio/transactions.xlsx")
    assert export_response.status_code == 200
    assert export_response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    _headers, exported_rows = parse_transaction_csv(
        transaction_xlsx_to_csv(export_response.content)
    )
    assert any(
        row.transaction is not None
        and row.transaction.external_reference == "xlsx-api-import-1"
        for row in exported_rows
    )

    template_response = client.get(
        "/api/portfolios/investment-studio/transactions/xlsx-template"
    )
    assert template_response.status_code == 200
    assert template_response.headers["content-disposition"] == (
        'attachment; filename="transaction-import-template.xlsx"'
    )
    template_workbook = load_workbook(BytesIO(template_response.content))
    assert template_workbook.sheetnames == [
        INSTRUCTIONS_SHEET_NAME,
        TRANSACTION_SHEET_NAME,
        FIELD_GUIDE_SHEET_NAME,
        EXAMPLES_SHEET_NAME,
        LISTS_SHEET_NAME,
    ]
    transaction_sheet = template_workbook[TRANSACTION_SHEET_NAME]
    assert transaction_sheet.max_row == 1
    xlsx_template_headers = tuple(
        transaction_sheet.cell(1, column_index).value
        for column_index in range(1, transaction_sheet.max_column + 1)
    )
    assert xlsx_template_headers == IMPORT_COLUMNS
    template_workbook.close()

    csv_template_response = client.get(
        "/api/portfolios/investment-studio/transactions/csv-template"
    )
    assert csv_template_response.status_code == 200
    assert csv_template_response.headers["content-disposition"] == (
        'attachment; filename="transaction-import-template.csv"'
    )
    csv_template_rows = list(
        csv.reader(
            StringIO(
                csv_template_response.content.decode("utf-8-sig"),
                newline="",
            )
        )
    )
    assert csv_template_rows == [list(xlsx_template_headers)]


def test_transaction_file_api_rejects_unsupported_extensions(client) -> None:
    response = client.post(
        "/api/portfolios/investment-studio/transactions/files/preview",
        files={"file": ("transactions.xls", b"legacy", "application/vnd.ms-excel")},
    )

    assert response.status_code == 422
    assert (
        response.json()["detail"] == "Transaction import files must use .csv or .xlsx."
    )


def test_transaction_file_preview_enforces_inception_and_accepts_transfer_identity(client) -> None:
    early_deposit = client.post(
        "/api/portfolios/investment-studio/transactions/csv/preview",
        json={
            "csv_text": (
                "asset_type,transaction_action,trade_date,account_id,gross_amount,currency\n"
                "cash,deposit,2026-01-01,cash-usd-main,1000,USD\n"
            )
        },
    )
    assert early_deposit.status_code == 200
    assert early_deposit.json()["error_count"] == 1
    assert "earlier than portfolio inception_date 2026-01-02" in (
        early_deposit.json()["rows"][0]["errors"][0]
    )

    late_opening = client.post(
        "/api/portfolios/investment-studio/transactions/csv/preview",
        json={
            "csv_text": (
                "asset_type,transaction_action,trade_date,account_id,gross_amount,currency\n"
                "cash,opening_balance,2026-01-03,cash-usd-main,1000,USD\n"
            )
        },
    )
    assert late_opening.status_code == 200
    assert late_opening.json()["error_count"] == 1
    assert "inception_date 2026-01-02" in late_opening.json()["rows"][0][
        "errors"
    ][0]

    transfer = client.post(
        "/api/portfolios/investment-studio/transactions/csv/preview",
        json={
            "csv_text": (
                "asset_type,transaction_action,trade_date,account_id,"
                "counterparty_account_id,gross_amount,currency,source_system,"
                "external_reference\n"
                "cash,transfer_out,2026-04-01,cash-usd-main,"
                "cash-usd-reserve,10,USD,custodian,TRANSFER-001\n"
            )
        },
    )
    assert transfer.status_code == 200
    assert transfer.json()["error_count"] == 0
    assert transfer.json()["warnings"] == []
    transfer_request = transfer.json()["rows"][0]["internal_transfer"]
    assert transfer_request["source_system"] == "custodian"
    assert transfer_request["external_reference"] == "TRANSFER-001"


def test_transfer_file_source_identity_prevents_replay_across_batches(client) -> None:
    csv_text = (
        "asset_type,transaction_action,trade_date,account_id,"
        "counterparty_account_id,gross_amount,currency,source_system,"
        "external_reference\n"
        "cash,transfer_out,2026-04-01,cash-usd-main,cash-usd-reserve,"
        "10,USD,custodian,TRANSFER-REPLAY-001\n"
    )
    preview = client.post(
        "/api/portfolios/investment-studio/transactions/csv/preview",
        json={"csv_text": csv_text},
    )
    assert preview.status_code == 200
    assert preview.json()["error_count"] == 0

    imported = client.post(
        "/api/portfolios/investment-studio/transactions/csv/import",
        headers={"Idempotency-Key": "transfer-file-batch-1"},
        json={
            "csv_text": csv_text,
            "preview_digest": preview.json()["preview_digest"],
        },
    )
    assert imported.status_code == 200
    assert imported.json()["created_count"] == 2
    transfer_out, transfer_in = sorted(
        imported.json()["transactions"],
        key=lambda row: row["transaction_type"],
        reverse=True,
    )
    assert transfer_out["transaction_type"] == "transfer_out"
    assert transfer_out["source_system"] == "custodian"
    assert transfer_out["external_reference"] == "TRANSFER-REPLAY-001"
    assert transfer_in["transaction_type"] == "transfer_in"
    assert transfer_in["source_system"] is None
    assert transfer_in["external_reference"] is None

    repeated_preview = client.post(
        "/api/portfolios/investment-studio/transactions/csv/preview",
        json={"csv_text": csv_text},
    )
    assert repeated_preview.status_code == 200
    assert repeated_preview.json()["error_count"] == 1
    assert repeated_preview.json()["batch_errors"] == [
        "Source identities already exist in this portfolio: "
        "custodian/TRANSFER-REPLAY-001."
    ]


def test_csv_and_xlsx_round_trip_preserve_cross_domain_transaction_semantics() -> None:
    option_terms = {
        "underlying_instrument_id": "equity-us-abbv",
        "option_type": "call",
        "expiry_date": "2026-12-18",
        "strike": 220,
        "contract_multiplier": 100,
    }
    records = [
        {
            **_deposit_record("cross-domain-cash"),
            "gross_amount": 1000,
            "source_gross_amount": "1000.00000000",
        },
        {
            "transaction_sequence": 9002,
            "transaction_type": "buy",
            "trade_date": "2026-05-25",
            "account_id": "broker-us-core",
            "settlement_cash_account_id": "cash-usd-main",
            "instrument_id": "equity-us-abbv",
            "instrument_ref": {
                "instrument_id": "equity-us-abbv",
                "instrument_name": "AbbVie Inc",
                "instrument_type": "equity",
                "exchange_code": "XNYS",
                "currency": "USD",
                "identifiers": [],
                "broker_identifiers": [],
            },
            "quantity": 2,
            "price": 100,
            "gross_amount": 200,
            "fees": 1,
            "taxes": 0,
            "fee_category": "transaction_cost",
            "currency": "USD",
            "source_system": "transaction_file_test",
            "external_reference": "cross-domain-security",
        },
        {
            "transaction_sequence": 9003,
            "transaction_type": "coupon",
            "trade_date": "2026-05-25",
            "entitlement_date": "2026-05-24",
            "account_id": "fcn-us",
            "settlement_cash_account_id": "cash-usd-main",
            "derivative_contract_id": "fcn-cross-domain",
            "derivative_contract": {
                "contract_name": "Cross-domain FCN",
                "contract_type": "fcn",
                "external_reference": "FCN-CROSS-DOMAIN",
                "terms": {
                    "notional": 100000,
                    "annual_coupon_rate_pct": 8,
                    "issue_date": "2026-01-01",
                    "final_observation_date": "2026-12-29",
                    "maturity_date": "2026-12-31",
                    "issuer": "Test Issuer",
                    "counterparty": "Test Broker",
                    "underlyings": [
                        {
                            "instrument_id": "equity-us-abbv",
                            "initial_reference_price": 100,
                            "deliverable": True,
                        }
                    ],
                },
            },
            "gross_amount": 2000,
            "fees": 0,
            "taxes": 0,
            "fee_category": "unknown",
            "currency": "USD",
            "source_system": "transaction_file_test",
            "external_reference": "cross-domain-fcn",
        },
        {
            "transaction_sequence": 9004,
            "transaction_type": "option_write",
            "trade_date": "2026-05-25",
            "account_id": "options-us",
            "settlement_cash_account_id": "cash-usd-main",
            "derivative_contract_id": "option-cross-domain",
            "derivative_contract": {
                "contract_name": "Cross-domain Call",
                "contract_type": "option",
                "external_reference": "OPTION-CROSS-DOMAIN",
                "terms": option_terms,
            },
            "quantity": 1,
            "price": 5,
            "gross_amount": 500,
            "fees": 2,
            "taxes": 1,
            "fee_category": "transaction_cost",
            "currency": "USD",
            "source_system": "transaction_file_test",
            "external_reference": "cross-domain-option",
        },
        {
            "transaction_sequence": 9005,
            "transaction_type": "transfer_out",
            "trade_date": "2026-05-26",
            "settlement_date": "2026-05-26",
            "account_id": "cash-usd-main",
            "counterparty_account_id": "cash-usd-reserve",
            "gross_amount": 100,
            "fees": 0,
            "taxes": 0,
            "currency": "USD",
            "transfer_scope": "internal_portfolio",
            "transfer_object_type": "cash",
            "transfer_group_id": "cross-domain-transfer",
            "source_system": "transaction_file_test",
            "external_reference": "cross-domain-transfer",
        },
        {
            "transaction_sequence": 9006,
            "transaction_type": "transfer_in",
            "trade_date": "2026-05-26",
            "settlement_date": "2026-05-26",
            "account_id": "cash-usd-reserve",
            "counterparty_account_id": "cash-usd-main",
            "gross_amount": 100,
            "fees": 0,
            "taxes": 0,
            "currency": "USD",
            "transfer_scope": "internal_portfolio",
            "transfer_object_type": "cash",
            "transfer_group_id": "cross-domain-transfer",
        },
    ]

    parsed_by_format = []
    for content in (
        render_transaction_csv(records),
        transaction_xlsx_to_csv(render_transaction_xlsx(records)),
    ):
        _headers, parsed_rows = parse_transaction_csv(content)
        assert all(row.errors == () for row in parsed_rows)
        parsed_by_format.append([
            (
                "transaction",
                row.transaction.transaction_type,
                row.transaction.instrument_id,
                row.transaction.derivative_contract_id,
                row.transaction.external_reference,
            )
            if row.transaction is not None
            else (
                "transfer",
                row.internal_transfer.transfer_object_type,
                row.internal_transfer.from_account_id,
                row.internal_transfer.to_account_id,
                float(row.internal_transfer.gross_amount or 0),
                row.internal_transfer.source_system,
                row.internal_transfer.external_reference,
            )
            for row in parsed_rows
            if row.transaction is not None or row.internal_transfer is not None
        ])

    assert parsed_by_format[0] == parsed_by_format[1] == [
        ("transaction", "deposit", None, None, "cross-domain-cash"),
        ("transaction", "buy", "equity-us-abbv", None, "cross-domain-security"),
        ("transaction", "coupon", None, "fcn-cross-domain", "cross-domain-fcn"),
        (
            "transaction",
            "option_write",
            None,
            "option-cross-domain",
            "cross-domain-option",
        ),
        (
            "transfer",
            "cash",
            "cash-usd-main",
            "cash-usd-reserve",
            100.0,
            "transaction_file_test",
            "cross-domain-transfer",
        ),
    ]
