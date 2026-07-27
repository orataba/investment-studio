from __future__ import annotations

from contextlib import nullcontext
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
import zipfile

from openpyxl import Workbook
import pytest

from platform_app.services import market_data_ops
from platform_app.services.market_data_ops import (
    _build_fund_nav_publication,
    _parse_nav_rows_from_attachment,
    _parse_nav_rows_from_label_snapshot_matrix,
    _parse_nav_rows_from_xlsx,
)


def _workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _multi_sheet_workbook_bytes(
    sheets: list[tuple[str, list[list[object]]]],
) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        sheet = workbook.create_sheet(title=title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_complete_fund_nav_publication(**kwargs: object):
    instrument = dict(kwargs["instrument"])
    instrument.setdefault("currency", "CNY")
    instrument.setdefault("fund_nav_reinvestment_evidence", [])
    kwargs["instrument"] = instrument
    rows = market_data_ops._stamp_nav_value_statuses(
        [
            {"currency": "CNY", **dict(row)}
            for row in list(kwargs.pop("rows"))
        ],
        status="complete",
    )
    return _build_fund_nav_publication(rows=rows, **kwargs)


def test_factor_quantization_matches_postgres_half_up_rounding() -> None:
    assert market_data_ops._quantized_factor_level(
        Decimal("1.0000000000000000005")
    ) == Decimal("1.000000000000000001")


def test_xshg_calendar_filters_weekends_and_spring_festival() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="daily-private-fund",
        instrument={
            "source_settings": {
                "expected_frequency": "daily",
                "market_calendar": "XSHG",
            },
            "fund_nav_events": [],
            "fund_nav_adjustment_factors": [],
        },
        rows=[
            {
                "as_of_date": "2026-02-13",
                "nav": "1.0000",
                "nav_with_dividend": "1.0000",
            },
            {
                "as_of_date": "2026-02-14",
                "nav": "1.0100",
                "nav_with_dividend": "1.0100",
            },
            {
                "as_of_date": "2026-02-16",
                "nav": "1.0200",
                "nav_with_dividend": "1.0200",
            },
            {
                "as_of_date": "2026-02-23",
                "nav": "1.0300",
                "nav_with_dividend": "1.0300",
            },
            {
                "as_of_date": "2026-02-24",
                "nav": "1.0400",
                "nav_with_dividend": "1.0400",
            },
        ],
    )

    assert [row["as_of_date"] for row in publication.rows] == [
        "2026-02-13",
        "2026-02-24",
    ]
    assert publication.projection_run["projection_status"] == "complete"
    calendar_evidence = publication.projection_run["evidence"][
        "market_calendar_filter"
    ]
    assert calendar_evidence == {
        "market_calendar": "XSHG",
        "unfiltered_source_observation_count": 5,
        "calendar_included_source_observation_count": 2,
        "calendar_excluded_source_observation_count": 3,
        "calendar_excluded_first_date": "2026-02-14",
        "calendar_excluded_last_date": "2026-02-23",
        "calendar_excluded_date_sample": [
            "2026-02-14",
            "2026-02-16",
            "2026-02-23",
        ],
        "calendar_excluded_dates_sha256": (
            "b772312118dfd9cb5942162147bed1626c7ab3a92d2f4190ab73105ad5815827"
        ),
        "unfiltered_source_observation_fingerprint": (
            "f5992551822d382fa478e04ae1988afd41727991603efe35d718e639b2d5f516"
        ),
        "filtered_source_observation_fingerprint": (
            "cc23e163203c3da043a4e407164c875d276edc995eb7b1c09dfcc085dfeed698"
        ),
    }


def test_unconfigured_market_calendar_keeps_reported_weekend_nav() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="unconfigured-private-fund",
        instrument={
            "source_settings": {"expected_frequency": "daily"},
            "fund_nav_events": [],
            "fund_nav_adjustment_factors": [],
        },
        rows=[
            {
                "as_of_date": "2026-02-13",
                "nav": "1.0000",
                "nav_with_dividend": "1.0000",
            },
            {
                "as_of_date": "2026-02-14",
                "nav": "1.0100",
                "nav_with_dividend": "1.0100",
            },
        ],
    )

    assert [row["as_of_date"] for row in publication.rows] == [
        "2026-02-13",
        "2026-02-14",
    ]
    assert "market_calendar_filter" not in publication.projection_run["evidence"]


def test_non_fund_nav_preview_and_file_import_fail_before_parsing_or_persistence(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        market_data_ops,
        "get_instrument",
        lambda instrument_id: {
            "instrument_id": instrument_id,
            "instrument_type": "fx",
        },
    )

    def fail_if_called(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("NAV parsing or persistence must not run for a non-fund instrument")

    monkeypatch.setattr(market_data_ops, "_parse_nav_rows_from_text", fail_if_called)
    monkeypatch.setattr(market_data_ops, "_parse_nav_rows_from_uploaded_file", fail_if_called)
    monkeypatch.setattr(market_data_ops, "publish_fund_nav_history", fail_if_called)

    with pytest.raises(ValueError, match="only supported for fund instruments"):
        market_data_ops.preview_nav_import(
            instrument_id="fx-usd-cny",
            raw_text="not-a-nav-row",
        )

    with pytest.raises(ValueError, match="only supported for fund instruments"):
        market_data_ops.import_nav_file(
            instrument_id="fx-usd-cny",
            file_name="not-a-workbook.xlsx",
            file_bytes=b"not-a-workbook",
            provider=None,
            status="complete",
            updated_by="pytest",
        )


def _with_broken_dimension(file_bytes: bytes) -> bytes:
    source = BytesIO(file_bytes)
    output = BytesIO()
    with zipfile.ZipFile(source) as reader, zipfile.ZipFile(output, "w") as writer:
        for info in reader.infolist():
            payload = reader.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b'dimension ref="A1:E2"', b'dimension ref="A1"')
            writer.writestr(info, payload)
    return output.getvalue()


def test_parse_nav_rows_from_xlsx_supports_chinese_headers() -> None:
    rows = [
        ["净值日期", "产品代码", "产品名称", "单位净值", "累计单位净值"],
        ["2026-04-14", "SBCJ69", "国泰君安期货CTA因子组合2号集合资产管理计划", "1.1002", "1.1002"],
    ]

    parsed = _parse_nav_rows_from_xlsx(_workbook_bytes(rows))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2026-04-14"
    assert str(parsed[0]["nav"]) == "1.1002"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.1002"
    assert parsed[0].get("nav_with_dividend") is None
    assert parsed[0]["instrument_code"] == "SBCJ69"
    assert parsed[0]["instrument_name"] == "国泰君安期货CTA因子组合2号集合资产管理计划"


def test_attachment_uses_xlsx_content_when_filename_is_mislabeled_xls() -> None:
    parsed = _parse_nav_rows_from_attachment(
        attachment_name="mislabeled.xls",
        attachment_bytes=_workbook_bytes(
            [
                ["净值日期", "产品代码", "产品名称", "单位净值"],
                ["2026-07-15", "SNJ545", "九慕黄龙1号私募证券投资基金", "1.0123"],
            ]
        ),
        parser_profile="generic_nav_table",
    )

    assert len(parsed) == 1
    assert parsed[0]["instrument_code"] == "SNJ545"
    assert str(parsed[0]["nav"]) == "1.0123"


def test_generic_attachment_parses_bilingual_administrator_headers() -> None:
    parsed = _parse_nav_rows_from_attachment(
        attachment_name="bilingual.xlsx",
        attachment_bytes=_workbook_bytes(
            [
                [
                    "日期\n（NAV As Of Date）",
                    "产品名称\n（Fund Name）",
                    "单位净值\n（NAV/Share）",
                    "累计单位净值\n（Accumulated NAV/Share）",
                    "协会备案编码\n（Fund Filling Code）",
                ],
                ["2026-07-17", "九木宏观对冲1号私募证券投资基金A", "1.2085", "1.2085", "SBAE63"],
            ]
        ),
        parser_profile="generic_nav_table",
    )

    assert parsed[0]["as_of_date"] == "2026-07-17"
    assert parsed[0]["instrument_code"] == "SBAE63"
    assert str(parsed[0]["nav"]) == "1.2085"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.2085"


def test_generic_attachment_falls_back_to_label_snapshot_layout() -> None:
    parsed = _parse_nav_rows_from_attachment(
        attachment_name="SAZL37-2026年07月16日-发送每日净值信息.xls",
        attachment_bytes=_workbook_bytes(
            [
                ["资产净值公告", None],
                ["彬元量化市场中性1号私募证券投资基金C专用表", None],
                ["2026年07月16日", None],
                ["基金代码：", "AZL37C"],
                ["基金名称：", "彬元量化市场中性1号私募证券投资基金C"],
                ["基金份额净值：", "1.2345"],
                ["基金份额累计净值：", "1.2345"],
            ]
        ),
        parser_profile="generic_nav_table",
    )

    assert parsed[0]["as_of_date"] == "2026-07-16"
    assert parsed[0]["instrument_code"] == "AZL37C"
    assert str(parsed[0]["nav"]) == "1.2345"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.2345"


def test_attachment_uses_xls_content_when_filename_is_mislabeled_xlsx(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        market_data_ops,
        "_xls_matrices",
        lambda payload: [
            (
                "Sheet1",
                [
                    ["净值日期", "产品代码", "单位净值"],
                    ["2026-07-15", "SNJ545", "1.0123"],
                ],
            )
        ],
    )
    monkeypatch.setattr(
        market_data_ops,
        "_xlsx_matrices",
        lambda payload: (_ for _ in ()).throw(
            AssertionError("OLE content must not be dispatched to the XLSX parser")
        ),
    )

    parsed = _parse_nav_rows_from_attachment(
        attachment_name="mislabeled.xlsx",
        attachment_bytes=market_data_ops._XLS_OLE_SIGNATURE + b"payload",
        parser_profile="generic_nav_table",
    )

    assert len(parsed) == 1
    assert parsed[0]["instrument_code"] == "SNJ545"
    assert str(parsed[0]["nav"]) == "1.0123"
    assert "currency" not in parsed[0]
    assert "frequency" not in parsed[0]


def test_parse_nav_rows_from_xlsx_merges_all_sheets_and_deduplicates_exact_rows() -> None:
    header = ["净值日期", "产品代码", "产品名称", "单位净值"]
    fund_a = ["2026-07-10", "FUND-A", "基金甲", "1.001"]
    fund_b = ["2026-07-10", "FUND-B", "基金乙", "0.998"]
    parsed = _parse_nav_rows_from_xlsx(
        _multi_sheet_workbook_bytes(
            [
                ("产品甲", [header, fund_a]),
                ("产品乙", [header, fund_b]),
                ("产品甲副本", [header, fund_a]),
            ]
        )
    )

    assert [row["instrument_code"] for row in parsed] == ["FUND-A", "FUND-B"]
    assert parsed[0]["_source_sheet"] == "产品甲"
    assert parsed[0]["_source_sheets"] == ["产品甲", "产品甲副本"]
    assert parsed[1]["_source_sheets"] == ["产品乙"]


def test_same_email_occurrence_conflicting_rows_fail_closed() -> None:
    with pytest.raises(ValueError, match="lack a strict source revision order"):
        market_data_ops._merge_rows_by_date(
            [
                {
                    "as_of_date": "2026-07-10",
                    "nav": "1.001",
                    "currency": "CNY",
                    "_email_candidate_route_id": 11,
                    "_email_folder": "INBOX",
                    "_email_uid": 42,
                    "_email_sent_at": "2026-07-11T01:00:00+00:00",
                },
                {
                    "as_of_date": "2026-07-10",
                    "nav": "1.002",
                    "currency": "CNY",
                    "_email_candidate_route_id": 12,
                    "_email_folder": "INBOX",
                    "_email_uid": 42,
                    "_email_sent_at": "2026-07-11T01:00:00+00:00",
                },
            ]
        )


def test_busy_email_ingestion_is_retryable_failure_not_configuration_block(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    instrument = {
        "instrument_id": "fund-a",
        "instrument_name": "基金甲",
        "instrument_type": "fund",
        "source_settings": {"source_mode": "email"},
    }

    def raise_busy(**kwargs: object) -> None:
        del kwargs
        raise market_data_ops.EmailIngestionBusyError("mailbox lease is active")

    monkeypatch.setattr(
        market_data_ops,
        "ingest_email_nav",
        raise_busy,
    )
    monkeypatch.setattr(
        market_data_ops,
        "update_refresh_status",
        lambda **kwargs: {**instrument, "refresh_status": kwargs},
    )

    response = market_data_ops._run_email_ingestion_batch(
        instruments=[instrument],
        settings=object(),
        updated_by="pytest",
        full_history=False,
    )

    assert response["results"][0]["status"] == "failed"
    assert response["results"][0]["message"] == "mailbox lease is active"


def test_durable_nav_loader_applies_email_history_boundary_to_every_raw_source(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    raw_rows = [
        {"as_of_date": "2025-12-25", "nav": "0.99", "currency": "CNY"},
        {"as_of_date": "2025-12-26", "nav": "1.00", "currency": "CNY"},
    ]
    email_rows = [
        {"as_of_date": "2025-12-24", "nav": "0.98", "currency": "CNY"},
        {"as_of_date": "2026-01-02", "nav": "1.01", "currency": "CNY"},
    ]

    class FakeRepository:
        def __init__(self, session_factory: object) -> None:
            del session_factory

        def matched_nav_history(self, instrument_id: str) -> list[dict[str, object]]:
            assert instrument_id == "fund-a"
            return email_rows

    monkeypatch.setattr(
        market_data_ops,
        "list_raw_nav_observations",
        lambda *, instrument_id: raw_rows,
    )
    monkeypatch.setattr(market_data_ops, "EmailIngestionRepository", FakeRepository)
    monkeypatch.setattr(market_data_ops, "get_session_factory", lambda: object())
    monkeypatch.setattr(
        market_data_ops,
        "get_settings",
        lambda: SimpleNamespace(email_history_start_date=date(2025, 12, 26)),
    )

    email_result = market_data_ops._load_all_durable_nav_source_rows(
        instrument_id="fund-a",
        instrument={
            "currency": "CNY",
            "source_settings": {"source_mode": "email"},
        },
    )
    manual_result = market_data_ops._load_all_durable_nav_source_rows(
        instrument_id="fund-a",
        instrument={
            "currency": "CNY",
            "source_settings": {"source_mode": "manual"},
        },
    )

    assert [row["as_of_date"] for row in email_result] == [
        "2025-12-26",
        "2026-01-02",
    ]
    assert [row["as_of_date"] for row in manual_result] == [
        "2025-12-24",
        "2025-12-25",
        "2025-12-26",
        "2026-01-02",
    ]


def test_projection_reconciliation_rebuilds_only_outdated_current_runs(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    stale = {
        "instrument_id": "fund-stale",
        "instrument_name": "旧方法基金",
        "instrument_type": "fund",
        "source_settings": {"source_mode": "email"},
        "current_fund_nav_projection_run_id": "run-v2",
        "fund_nav_projection_runs": [
            {
                "fund_nav_projection_run_id": "run-v2",
                "method_version": "fund_nav_reinvestment_projection/v2",
                "source_provider": "email",
            }
        ],
    }
    current = {
        "instrument_id": "fund-current",
        "instrument_name": "新方法基金",
        "instrument_type": "fund",
        "source_settings": {"source_mode": "manual"},
        "current_fund_nav_projection_run_id": "run-v3",
        "fund_nav_projection_runs": [
            {
                "fund_nav_projection_run_id": "run-v3",
                "method_version": market_data_ops.FUND_NAV_PROJECTION_METHOD_VERSION,
                "source_provider": "manual",
            }
        ],
    }
    captured: dict[str, object] = {}

    def fake_publish(**kwargs: object):
        captured.update(kwargs)
        return stale, SimpleNamespace(rows=[{"as_of_date": "2026-01-02"}], action_candidates=[])

    monkeypatch.setattr(market_data_ops, "list_instruments", lambda **kwargs: [stale, current])
    monkeypatch.setattr(
        market_data_ops,
        "list_stale_current_fund_nav_projections",
        lambda **kwargs: [
            {
                "instrument_id": "fund-stale",
                "method_version": "fund_nav_reinvestment_projection/v2",
                "source_provider": "email",
            }
        ],
    )
    monkeypatch.setattr(market_data_ops, "_publish_fund_nav_projection", fake_publish)
    monkeypatch.setattr(
        market_data_ops,
        "market_data_item_timeout",
        lambda instrument_id: nullcontext(),
    )

    response = market_data_ops.rebuild_stale_fund_nav_projections(
        updated_by="pytest"
    )

    assert captured["instrument_id"] == "fund-stale"
    assert captured["source_provider"] == "email"
    assert captured["mode"] == "projection_reconciliation"
    assert response["refreshed_count"] == 1
    assert response["skipped_count"] == 1
    assert response["results"][0]["instrument_id"] == "fund-stale"


def test_email_batch_rebuilds_existing_publication_that_crosses_history_boundary(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    instrument = {
        "instrument_id": "fund-a",
        "instrument_name": "基金甲",
        "instrument_type": "fund",
        "currency": "CNY",
        "source_settings": {"source_mode": "email"},
    }
    ingestion = SimpleNamespace(
        batches={},
        failed_folders=[],
        folders=[],
        unmatched_candidates=0,
        ambiguous_candidates=0,
        invalid_candidates=0,
        out_of_scope_candidates=0,
    )
    captured: dict[str, object] = {}

    class FakeRepository:
        def __init__(self, session_factory: object) -> None:
            del session_factory

        def mark_imported(self, route_ids: list[int]) -> None:
            captured["route_ids"] = route_ids

    def fake_publish(**kwargs: object):
        captured["published_instrument_id"] = kwargs["instrument_id"]
        publication = SimpleNamespace(rows=[{"as_of_date": "2025-12-26"}], action_candidates=[])
        return instrument, publication

    def fake_boundary_query(**kwargs: object) -> set[str]:
        captured["boundary_query"] = kwargs
        return {"fund-a"}

    monkeypatch.setattr(market_data_ops, "ingest_email_nav", lambda **kwargs: ingestion)
    monkeypatch.setattr(
        market_data_ops,
        "list_instrument_ids_with_nav_history_before",
        fake_boundary_query,
    )
    monkeypatch.setattr(market_data_ops, "_publish_fund_nav_projection", fake_publish)
    monkeypatch.setattr(market_data_ops, "EmailIngestionRepository", FakeRepository)
    monkeypatch.setattr(market_data_ops, "get_session_factory", lambda: object())

    response = market_data_ops._run_email_ingestion_batch(
        instruments=[instrument],
        settings=SimpleNamespace(email_history_start_date=date(2025, 12, 26)),
        updated_by="pytest",
        full_history=False,
    )

    assert captured["boundary_query"] == {
        "instrument_ids": {"fund-a"},
        "before_date": date(2025, 12, 26),
    }
    assert captured["published_instrument_id"] == "fund-a"
    assert captured["route_ids"] == []
    assert response["refreshed_count"] == 1
    assert response["results"][0]["status"] == "imported"
    assert "2025-12-26" in response["results"][0]["message"]


def test_label_snapshot_does_not_misclassify_longer_nav_labels_as_unit_nav() -> None:
    parsed = _parse_nav_rows_from_label_snapshot_matrix(
        [
            ["净值日期：2026-07-10"],
            ["累计单位净值：1.4200"],
            ["复权单位净值：1.5300"],
            ["单位净值：1.1200"],
        ]
    )

    assert len(parsed) == 1
    assert str(parsed[0]["nav"]) == "1.1200"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.4200"
    assert str(parsed[0]["nav_with_dividend"]) == "1.5300"


def test_normalize_nav_rows_uses_selected_fund_currency_and_rejects_mismatch() -> None:
    instrument = {"currency": "HKD"}
    source_row = {
        "as_of_date": "2026-04-14",
        "nav": "1.1002",
        "frequency": "daily",
    }

    prepared = market_data_ops._normalize_nav_rows_for_instrument(
        instrument=instrument,
        rows=[source_row],
    )

    assert prepared[0]["currency"] == "HKD"
    with pytest.raises(ValueError, match="does not match the selected fund currency"):
        market_data_ops._normalize_nav_rows_for_instrument(
            instrument=instrument,
            rows=[{**source_row, "currency": "USD"}],
        )


def test_parse_nav_rows_from_xlsx_supports_product_code_and_trade_date_aliases() -> None:
    rows = [
        ["产品编码", "产品名称", "交易日期", "单位净值", "累计净值"],
        ["SBMM07", "国泰君安期货CTA因子组合3号集合资产管理计划", "20260529", "1.0148", "1.0148"],
    ]

    parsed = _parse_nav_rows_from_xlsx(_workbook_bytes(rows))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2026-05-29"
    assert parsed[0]["instrument_code"] == "SBMM07"
    assert parsed[0]["instrument_name"] == "国泰君安期货CTA因子组合3号集合资产管理计划"
    assert str(parsed[0]["nav"]) == "1.0148"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.0148"


def test_parse_nav_rows_from_xlsx_ignores_broken_dimension_metadata() -> None:
    rows = [
        ["净值日期", "产品代码", "产品名称", "单位净值", "累计单位净值"],
        ["2026-04-14", "B3935B", "孝庸市场中性一号私募证券投资基金B", "1.3699", "1.3699"],
    ]

    parsed = _parse_nav_rows_from_xlsx(_with_broken_dimension(_workbook_bytes(rows)))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2026-04-14"
    assert str(parsed[0]["nav"]) == "1.3699"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.3699"


def test_parse_nav_rows_from_xlsx_supports_total_nav_aliases_and_datetime_values() -> None:
    rows = [
        ["产品代码", "产品名称", "净值日期", "单位净值  (元)", "累计净值  (元)"],
        ["ANZ73A(A级)", "盈怀香柏树1号私募证券投资基金A类", "2024-09-05 00:00:00", "1", "1"],
    ]

    parsed = _parse_nav_rows_from_xlsx(_workbook_bytes(rows))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2024-09-05"
    assert str(parsed[0]["nav"]) == "1"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1"
    assert parsed[0]["instrument_code"] == "ANZ73A(A级)"
    assert parsed[0]["instrument_name"] == "盈怀香柏树1号私募证券投资基金A类"


def test_parse_nav_rows_from_xlsx_supports_yuan_per_share_headers() -> None:
    rows = [
        ["产品基金净值数据"],
        ["产品代码", "产品名称", "净值日期", "单位净值(元/份)", "累计单位净值(元/份)"],
        ["SAZN63", "孝庸混合策略配置一号私募证券投资基金", "2026-06-01", "1.0147", "1.0147"],
    ]

    parsed = _parse_nav_rows_from_xlsx(_workbook_bytes(rows))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2026-06-01"
    assert str(parsed[0]["nav"]) == "1.0147"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.0147"
    assert parsed[0]["instrument_code"] == "SAZN63"
    assert parsed[0]["instrument_name"] == "孝庸混合策略配置一号私募证券投资基金"


def test_parse_nav_rows_from_xlsx_supports_chinese_date_strings_after_title_rows() -> None:
    rows = [
        ["每日净值表"],
        ["日期：2024年12月05日至2025年10月12日"],
        ["日期", "产品代码", "产品名称", "单位净值", "累计单位净值"],
        ["2024年12月05日", "ARE77A", "盈怀香柏树7号私募证券投资基金A", "1.0", "1.0"],
    ]

    parsed = _parse_nav_rows_from_xlsx(_workbook_bytes(rows))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2024-12-05"
    assert parsed[0]["instrument_code"] == "ARE77A"
    assert parsed[0]["instrument_name"] == "盈怀香柏树7号私募证券投资基金A"
    assert str(parsed[0]["nav"]) == "1.0"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.0"


def test_parse_nav_rows_from_xlsx_supports_yyyymmdd_business_dates() -> None:
    rows = [
        [
            "产品代码",
            "产品名称",
            "业务日期",
            "客户资产净值",
            "客户资产份额",
            "单位净值",
            "累计单位净值",
        ],
        [
            "ZB945A",
            "润洲正行11号私募证券投资基金A",
            "20260422",
            "1,028,139.61",
            "1,059,391.66",
            "0.9705",
            "1.5268",
        ],
    ]

    parsed = _parse_nav_rows_from_xlsx(_workbook_bytes(rows))

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2026-04-22"
    assert parsed[0]["instrument_code"] == "ZB945A"
    assert parsed[0]["instrument_name"] == "润洲正行11号私募证券投资基金A"
    assert str(parsed[0]["nav"]) == "0.9705"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.5268"


def test_reinvested_total_return_normalizes_at_first_observed_cash_window() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0500",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "1.0100",
                "cash_cumulative_nav": "1.0600",
            },
        ],
    )

    assert publication.has_derived_total_return is True
    assert [str(row.get("nav_with_dividend")) for row in publication.rows] == [
        "1.0000000000000000",
        "1.0100000000000000",
    ]
    assert publication.projection_run["projection_status"] == "complete"
    assert publication.projection_run["anchor_date"] == "2026-05-27"
    assert publication.action_candidates == []
    assert [
        factor["evidence_kind"] for factor in publication.adjustment_factors
    ] == ["window_normalized_anchor"]
    assert publication.adjustment_factors[0]["evidence"][
        "starting_cash_balance"
    ] == "0.0500"


def test_positive_cash_cumulative_jump_continues_total_return() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "0.9000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-29",
                "nav": "0.9900",
                "cash_cumulative_nav": "1.0900",
            },
            {
                "as_of_date": "2026-05-30",
                "nav": "1.0800",
                "cash_cumulative_nav": "1.1800",
            },
            {
                "as_of_date": "2026-05-31",
                "nav": "1.1700",
                "cash_cumulative_nav": "1.2700",
            },
        ],
    )

    assert str(publication.rows[0]["nav_with_dividend"]) == "1.0000000000000000"
    assert str(publication.rows[1]["nav_with_dividend"]) == "1.0000000000000000"
    assert str(publication.rows[2]["nav_with_dividend"]) == "1.1000000000000000"
    assert str(publication.rows[3]["nav_with_dividend"]) == "1.2000000000000000"
    assert str(publication.rows[4]["nav_with_dividend"]) == "1.3000000000000000"
    assert publication.derived_total_return_count == 5
    assert publication.action_candidates == []
    assert publication.projection_run["projection_status"] == "complete"
    assert publication.projection_run["projection_kind"] == "hybrid_reanchored"
    assert [
        factor["evidence_kind"] for factor in publication.adjustment_factors
    ] == ["window_normalized_anchor", "provider_cash_cumulative"]
    cash_factor = publication.adjustment_factors[-1]
    assert cash_factor["factor_level"] == "1.111111111111111111"
    assert cash_factor["evidence"]["cash_per_unit"] == "0.1000"
    assert cash_factor["evidence"]["reinvestment_nav"] == "0.9000"
    assert publication.projection_run["evidence"]["auto_cash_distributions"] == [
        {
            "interval_start_date": "2026-05-27",
            "interval_end_date": "2026-05-28",
            "cash_per_unit": "0.1000",
            "reinvestment_nav": "0.9000",
            "implied_interval_return": "0",
            "resolved_by": "cash_cumulative_disclosure",
        }
    ]


def test_single_cumulative_jump_remains_unresolved_until_semantics_are_clear() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="short-cumulative-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "0.9000",
                "cash_cumulative_nav": "1.0000",
            },
        ],
    )

    assert publication.projection_run["evidence"]["cumulative_nav_semantics"][
        "kind"
    ] == "ambiguous"
    assert publication.rows[0].get("nav_with_dividend") is not None
    assert publication.rows[1].get("nav_with_dividend") is None
    assert len(publication.action_candidates) == 1
    assert "auto_cash_distributions" not in publication.projection_run["evidence"]


def test_daily_cash_distribution_matches_real_private_fund_disclosure_shape() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="daily-private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-01-12",
                "nav": "1.7202",
                "cash_cumulative_nav": "1.9830",
            },
            {
                "as_of_date": "2026-01-13",
                "nav": "1.6227",
                "cash_cumulative_nav": "1.9855",
            },
            {
                "as_of_date": "2026-01-14",
                "nav": "1.6300",
                "cash_cumulative_nav": "1.9928",
            },
            {
                "as_of_date": "2026-01-15",
                "nav": "1.6400",
                "cash_cumulative_nav": "2.0028",
            },
            {
                "as_of_date": "2026-01-16",
                "nav": "1.6500",
                "cash_cumulative_nav": "2.0128",
            },
        ],
    )

    assert [str(row["nav_with_dividend"]) for row in publication.rows[:2]] == [
        "1.7202000000000000",
        "1.7227000000000000",
    ]
    assert publication.projection_run["projection_status"] == "complete"
    assert publication.action_candidates == []
    assert publication.adjustment_factors[-1]["evidence"]["cash_per_unit"] == (
        "0.1000"
    )


def test_reinvested_cumulative_nav_is_used_directly_without_double_adjustment() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="reinvested-cumulative-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "0.9000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-29",
                "nav": "0.9900",
                "cash_cumulative_nav": "1.1000",
            },
            {
                "as_of_date": "2026-05-30",
                "nav": "1.0800",
                "cash_cumulative_nav": "1.2000",
            },
            {
                "as_of_date": "2026-05-31",
                "nav": "1.1700",
                "cash_cumulative_nav": "1.3000",
            },
        ],
    )

    assert [str(row["nav_with_dividend"]) for row in publication.rows] == [
        "1.0000000000000000",
        "1.0000000000000000",
        "1.1000000000000000",
        "1.2000000000000000",
        "1.3000000000000000",
    ]
    assert publication.derived_total_return_count == 0
    assert publication.explicit_total_return_count == 5
    assert publication.action_candidates == []
    assert all(
        factor["evidence_kind"] == "provider_total_return"
        for factor in publication.adjustment_factors
    )
    semantics = publication.projection_run["evidence"][
        "cumulative_nav_semantics"
    ]
    assert semantics["kind"] == "dividend_reinvested"
    assert semantics["dividend_reinvested_votes"] >= 3
    assert publication.rows[-1]["nav_with_dividend_lineage"]["evidence"][
        "source_field"
    ] == "cumulative_nav_dividend_reinvested"


def test_provider_total_return_is_preserved_with_explicit_lineage() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-28",
                "nav": "1.0000",
                "nav_with_dividend": "1.0500",
                "_total_return_source_field": "复权净值",
                "_email_candidate_route_id": 17,
            }
        ],
    )

    assert publication.has_derived_total_return is False
    assert publication.explicit_total_return_count == 1
    assert str(publication.rows[0]["nav_with_dividend"]) == "1.0500000000000000"
    assert publication.rows[0]["nav_source_provider"] == "email-route:17"
    assert (
        publication.rows[0]["nav_with_dividend_source_provider"]
        == "email-route:17"
    )
    assert publication.rows[0]["nav_with_dividend_lineage"] == {
        "kind": "provider_explicit",
        "evidence": {
            "factor_logical_key": "provider:2026-05-28",
            "source_field": "复权净值",
            "provider_reported_total_return_nav": "1.0500",
            "email_candidate_route_id": 17,
        },
    }
    assert len(publication.adjustment_factors) == 1
    assert publication.adjustment_factors[0]["factor_kind"] == "provider_implied"
    assert publication.adjustment_factors[0]["factor_level"] == "1.050000000000000000"
    assert publication.projection_run["projection_status"] == "complete"


def test_newer_row_revision_retracts_omitted_total_return_instead_of_splicing() -> None:
    merged = market_data_ops._merge_rows_by_date(
        [
            {
                "as_of_date": "2026-05-28",
                "nav": "1.0000",
                "nav_with_dividend": "1.0500",
                "currency": "CNY",
                "_raw_source_kind": "manual_import",
                "_raw_source_ref": "manual.csv:row:old",
                "_raw_source_provider": "manual:old",
                "_raw_captured_at": "2026-05-29T10:00:00+00:00",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "0.9000",
                "currency": "CNY",
                "_raw_source_kind": "manual_import",
                "_raw_source_ref": "manual.csv:row:new",
                "_raw_source_provider": "manual:new",
                "_raw_captured_at": "2026-05-29T11:00:00+00:00",
            },
        ]
    )

    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=merged,
    )

    row = publication.rows[0]
    assert row["nav"] == "0.9000"
    assert row.get("nav_with_dividend") is None
    assert row["nav_source_provider"] == "manual:new|manual.csv:row:new"
    assert row["nav_lineage"]["evidence"]["raw_source_provider"] == "manual:new"
    assert publication.adjustment_factors == []


def test_tushare_nav_rows_use_announcement_date_as_strict_revision_order() -> None:
    merged = market_data_ops._tushare_nav_rows(
        [
            {
                "ts_code": "018201.OF",
                "ann_date": "20260724",
                "nav_date": "20260723",
                "unit_nav": "0.793",
                "accum_nav": "0.793",
                "adj_nav": None,
                "update_flag": "0",
            },
            {
                "ts_code": "018201.OF",
                "ann_date": "20260723",
                "nav_date": "20260723",
                "unit_nav": "0.793",
                "accum_nav": "0.793",
                "adj_nav": "0.793",
                "update_flag": "0",
            },
        ],
        instrument_currency="CNY",
        latest_date=date(2026, 7, 22),
    )

    assert len(merged) == 1
    assert str(merged[0]["nav"]) == "0.793"
    assert str(merged[0]["cash_cumulative_nav"]) == "0.793"
    assert str(merged[0]["nav_with_dividend"]) == "0.793"
    assert merged[0]["_provider_revision_at"] == "2026-07-24"
    assert merged[0]["_provider_revision_scope"] == "tushare:fund_nav:018201.OF"


def test_cash_distribution_reversal_breaks_the_derived_chain() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "0.9000",
                "cash_cumulative_nav": "0.9000",
            },
            {
                "as_of_date": "2026-05-29",
                "nav": "0.9500",
                "cash_cumulative_nav": "0.9000",
            },
        ],
    )

    assert publication.has_derived_total_return is True
    assert publication.rows[0].get("nav_with_dividend") is not None
    assert publication.rows[1].get("nav_with_dividend") is not None
    assert publication.rows[2].get("nav_with_dividend") is None
    assert publication.projection_run["projection_status"] == "partial"
    assert len(publication.action_candidates) == 1
    assert (
        publication.action_candidates[0]["candidate_type"]
        == "cash_balance_discontinuity"
    )


def test_sustained_cash_cumulative_reporting_reset_preserves_unit_return_chain() -> None:
    rows = [
        {
            "as_of_date": "2025-08-01",
            "nav": "1.179",
            "cash_cumulative_nav": "1.363",
        },
        {
            "as_of_date": "2025-08-08",
            "nav": "1.207",
            "cash_cumulative_nav": "1.391",
        },
        {
            "as_of_date": "2025-08-15",
            "nav": "1.235",
            "cash_cumulative_nav": "1.419",
        },
        {
            "as_of_date": "2025-08-26",
            "nav": "1.301",
            "cash_cumulative_nav": "1.301",
        },
        {
            "as_of_date": "2025-08-29",
            "nav": "1.286",
            "cash_cumulative_nav": "1.286",
        },
        {
            "as_of_date": "2025-09-02",
            "nav": "1.274",
            "cash_cumulative_nav": "1.274",
        },
    ]

    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=rows,
    )

    assert publication.projection_run["projection_status"] == "complete"
    assert publication.action_candidates == []
    assert all(row.get("nav_with_dividend") for row in publication.rows)
    resets = publication.projection_run["evidence"][
        "auto_cash_reporting_basis_resets"
    ]
    assert resets == [
        {
            "kind": "cash_cumulative_reporting_basis_reset_to_unit_nav",
            "interval_start_date": "2025-08-15",
            "interval_end_date": "2025-08-26",
            "stable_pre_reset_start_date": "2025-08-08",
            "cash_balance_before": "0.184",
            "cash_balance_after": "0.000",
            "unit_nav_before": "1.235",
            "unit_nav_after": "1.301",
            "unit_return_at_reset": "0.053441295546558704",
            "confirmation_dates": [
                "2025-08-26",
                "2025-08-29",
                "2025-09-02",
            ],
        }
    ]


def test_confirmed_distribution_advances_factor_and_total_return_curve() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="private-fund",
        instrument={
            "fund_nav_events": [
                {
                    "fund_nav_event_id": "distribution-2026-05-28",
                    "event_type": "cash_distribution",
                    "effective_date": "2026-05-28",
                    "cash_per_unit": "0.1000",
                    "source": "provider-notice:2026-05-28",
                }
            ],
            "fund_nav_reinvestment_evidence": [
                {
                    "fund_nav_reinvestment_evidence_id": "distribution-evidence",
                    "fund_nav_event_id": "distribution-2026-05-28",
                    "reinvestment_nav": "0.9000",
                }
            ],
            "fund_nav_adjustment_factors": [],
        },
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "0.9000",
                "cash_cumulative_nav": "1.0000",
            },
        ],
    )

    assert publication.action_candidates == []
    assert publication.derived_total_return_count == 2
    assert str(publication.rows[1]["nav_with_dividend"]) == "1.0000000000000000"
    assert publication.rows[1]["nav_with_dividend_lineage"]["evidence"][
        "factor_logical_key"
    ] == publication.adjustment_factors[1]["factor_logical_key"]
    assert publication.adjustment_factors[1]["evidence_kind"] == "fund_nav_event"


def test_weekly_nav_applies_confirmed_event_inside_observation_interval() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="weekly-private-fund",
        instrument={
            "fund_nav_events": [
                {
                    "fund_nav_event_id": "weekly-distribution",
                    "event_type": "cash_distribution",
                    "effective_date": "2026-05-26",
                    "cash_per_unit": "0.1000",
                    "source": "provider-notice:weekly-distribution",
                }
            ],
            "fund_nav_reinvestment_evidence": [
                {
                    "fund_nav_reinvestment_evidence_id": "weekly-evidence",
                    "fund_nav_event_id": "weekly-distribution",
                    "reinvestment_nav": "0.9000",
                }
            ],
            "fund_nav_adjustment_factors": [],
        },
        rows=[
            {
                "as_of_date": "2026-05-23",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-30",
                "nav": "0.9500",
                "cash_cumulative_nav": "1.0500",
            },
        ],
    )

    assert publication.action_candidates == []
    assert publication.derived_total_return_count == 2
    assert str(publication.rows[1]["nav_with_dividend"]) == "1.0555555555555556"
    event_factor = publication.adjustment_factors[-1]
    assert event_factor["as_of_date"] == "2026-05-26"
    assert event_factor["fund_nav_event_id"] == "weekly-distribution"
    assert publication.rows[1]["nav_with_dividend_lineage"]["evidence"][
        "factor_logical_key"
    ] == event_factor["factor_logical_key"]


def test_missing_cash_disclosure_is_na_without_losing_complete_baseline() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="sparse-private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "1.0100",
            },
            {
                "as_of_date": "2026-05-29",
                "nav": "1.0200",
                "cash_cumulative_nav": "1.0200",
            },
        ],
    )

    assert publication.rows[0].get("nav_with_dividend") is not None
    assert publication.rows[1].get("nav_with_dividend") is None
    assert str(publication.rows[2]["nav_with_dividend"]) == "1.0200000000000000"
    assert publication.action_candidates == []
    assert publication.projection_run["projection_status"] == "partial"


def test_independent_nav_rounding_jitter_does_not_create_false_action() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="rounded-private-fund",
        instrument={"fund_nav_events": [], "fund_nav_adjustment_factors": []},
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-28",
                "nav": "1.0001",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-29",
                "nav": "0.9999",
                "cash_cumulative_nav": "1.0000",
            },
        ],
    )

    assert publication.action_candidates == []
    assert publication.derived_total_return_count == 3
    assert all(row.get("nav_with_dividend") is not None for row in publication.rows)


def test_same_day_cash_and_split_events_fail_closed_without_sequence_evidence() -> None:
    publication = _build_complete_fund_nav_publication(
        instrument_id="mixed-event-private-fund",
        instrument={
            "fund_nav_events": [
                {
                    "fund_nav_event_id": "same-day-cash",
                    "event_type": "cash_distribution",
                    "effective_date": "2026-05-28",
                    "cash_per_unit": "0.1000",
                    "source": "provider-notice:same-day-cash",
                },
                {
                    "fund_nav_event_id": "same-day-split",
                    "event_type": "unit_split",
                    "effective_date": "2026-05-28",
                    "unit_ratio": "2",
                    "source": "provider-notice:same-day-split",
                },
            ],
            "fund_nav_reinvestment_evidence": [
                {
                    "fund_nav_reinvestment_evidence_id": "same-day-evidence",
                    "fund_nav_event_id": "same-day-cash",
                    "reinvestment_nav": "0.9000",
                }
            ],
            "fund_nav_adjustment_factors": [],
        },
        rows=[
            {
                "as_of_date": "2026-05-27",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
            },
            {
                "as_of_date": "2026-05-29",
                "nav": "0.4500",
                "cash_cumulative_nav": "0.5000",
            },
        ],
    )

    assert publication.rows[1].get("nav_with_dividend") is None
    assert publication.projection_run["projection_status"] == "partial"
    assert len(publication.action_candidates) == 1
    assert publication.action_candidates[0]["source_evidence"][
        "confirmed_event_ids"
    ] == ["same-day-cash", "same-day-split"]
    assert all(
        factor.get("fund_nav_event_id") is None
        for factor in publication.adjustment_factors
    )


def test_missing_reinvestment_evidence_cuts_only_the_affected_segment_and_provider_reanchors() -> None:
    instrument = {
        "currency": "CNY",
        "fund_nav_events": [
            {
                "fund_nav_event_id": "cash-event-without-reinvestment-price",
                "event_type": "cash_distribution",
                "effective_date": "2026-05-28",
                "cash_per_unit": "0.1000",
                "source": "fund-manager-notice",
            }
        ],
        "fund_nav_reinvestment_evidence": [],
    }
    source_rows = [
        {
            "as_of_date": "2026-05-27",
            "currency": "CNY",
            "nav": "1.0000",
            "cash_cumulative_nav": "1.0000",
        },
        {
            "as_of_date": "2026-05-28",
            "currency": "CNY",
            "nav": "0.9000",
            "cash_cumulative_nav": "1.0000",
        },
        {
            "as_of_date": "2026-05-29",
            "currency": "CNY",
            "nav": "1.0000",
            "cash_cumulative_nav": "1.1000",
            "nav_with_dividend": "1.2000",
            "_total_return_source_field": "复权单位净值",
        },
        {
            "as_of_date": "2026-05-30",
            "currency": "CNY",
            "nav": "1.1000",
            "cash_cumulative_nav": "1.2000",
        },
    ]

    publication = _build_fund_nav_publication(
        instrument_id="reanchored-private-fund",
        instrument=instrument,
        rows=market_data_ops._stamp_nav_value_statuses(
            source_rows,
            status="complete",
        ),
    )

    assert [
        row.get("nav_with_dividend") is not None for row in publication.rows
    ] == [False, False, True, True]
    assert str(publication.rows[-1]["nav_with_dividend"]) == "1.3200000000000000"
    assert publication.projection_run["projection_kind"] == "provider_explicit"
    assert publication.projection_run["projection_status"] == "partial"
    assert publication.projection_run["evidence"]["missing_total_return_dates"] == [
        "2026-05-27",
        "2026-05-28",
    ]
    assert publication.projection_run["anchor_date"] == "2026-05-29"
    assert "nav_with_dividend" not in publication.rows[0]
    assert "nav_with_dividend_status" not in publication.rows[0]
    assert "nav_with_dividend_lineage" not in publication.rows[0]
    assert publication.rows[2]["nav_with_dividend_status"] == "complete"
    assert publication.current_fund_nav_event_ids == [
        "cash-event-without-reinvestment-price"
    ]
    assert publication.current_fund_nav_reinvestment_evidence_ids == []
    assert [
        factor["factor_logical_key"] for factor in publication.adjustment_factors
    ] == ["provider:2026-05-29"]
    assert publication.rows[-1]["nav_with_dividend_lineage"]["evidence"][
        "factor_logical_key"
    ] == "provider:2026-05-29"
    assert publication.projection_run["evidence"]["segment_breaks"] == [
        {
            "as_of_date": "2026-05-28",
            "reason": "reinvestment_nav_not_observed",
            "fund_nav_event_id": "cash-event-without-reinvestment-price",
        },
        {
            "as_of_date": "2026-05-29",
            "reason": "provider_reanchor_without_verified_overlap",
            "excluded_complete_points": 1,
        },
    ]


def test_provider_total_return_normalizes_onto_verified_derived_history() -> None:
    publication = _build_fund_nav_publication(
        instrument_id="scale-changing-private-fund",
        instrument={
            "currency": "CNY",
            "fund_nav_events": [],
            "fund_nav_reinvestment_evidence": [],
        },
        rows=market_data_ops._stamp_nav_value_statuses(
            [
                {
                    "as_of_date": "2026-05-27",
                    "currency": "CNY",
                    "nav": "1.0000",
                    "cash_cumulative_nav": "1.0000",
                },
                {
                    "as_of_date": "2026-05-28",
                    "currency": "CNY",
                    "nav": "1.1000",
                    "cash_cumulative_nav": "1.1000",
                },
                {
                    "as_of_date": "2026-05-29",
                    "currency": "CNY",
                    "nav": "1.2000",
                    "cash_cumulative_nav": "1.2000",
                    "nav_with_dividend": "2.4000",
                    "_total_return_source_field": "复权单位净值",
                },
                {
                    "as_of_date": "2026-05-30",
                    "currency": "CNY",
                    "nav": "1.3000",
                    "cash_cumulative_nav": "1.3000",
                },
            ],
            status="complete",
        ),
    )

    assert [
        row.get("nav_with_dividend_status") for row in publication.rows
    ] == ["complete", "complete", "complete", "complete"]
    assert all(
        row.get("nav_with_dividend") is None
        or row.get("nav_with_dividend_status") == "complete"
        for row in publication.rows
    )
    assert [
        factor["factor_logical_key"] for factor in publication.adjustment_factors
    ] == ["window:2026-05-27", "provider:2026-05-29"]
    assert publication.projection_run["anchor_date"] == "2026-05-27"
    assert publication.projection_run["evidence"]["published_total_return_dates"] == [
        "2026-05-27",
        "2026-05-28",
        "2026-05-29",
        "2026-05-30",
    ]
    assert publication.projection_run["evidence"]["missing_total_return_dates"] == []
    assert publication.projection_run["evidence"]["segment_breaks"] == []
    assert str(publication.rows[0]["nav_with_dividend"]) == "1.0000000000000000"
    assert str(publication.rows[1]["nav_with_dividend"]) == "1.1000000000000000"
    assert str(publication.rows[2]["nav_with_dividend"]) == "1.2000000000000000"
    assert str(publication.rows[3]["nav_with_dividend"]) == "1.3000000000000000"
    assert publication.adjustment_factors[-1]["evidence"][
        "canonical_normalization_scale"
    ] == "0.500000000000000000"


def test_projection_source_fingerprint_is_order_stable_and_changes_with_source_evidence() -> None:
    rows = market_data_ops._stamp_nav_value_statuses(
        [
            {
                "as_of_date": "2026-05-27",
                "currency": "CNY",
                "nav": "1.0000",
                "cash_cumulative_nav": "1.0000",
                "_raw_observation_id": 11,
            },
            {
                "as_of_date": "2026-05-28",
                "currency": "CNY",
                "nav": "1.0100",
                "cash_cumulative_nav": "1.0100",
                "_raw_observation_id": 12,
            },
        ],
        status="complete",
    )
    instrument = {
        "currency": "CNY",
        "fund_nav_events": [],
        "fund_nav_reinvestment_evidence": [],
    }

    first = _build_fund_nav_publication(
        instrument_id="fingerprinted-fund",
        instrument=instrument,
        rows=rows,
    )
    reordered = _build_fund_nav_publication(
        instrument_id="fingerprinted-fund",
        instrument=instrument,
        rows=list(reversed(rows)),
    )
    revised_rows = [dict(row) for row in rows]
    revised_rows[-1]["_raw_observation_id"] = 13
    revised = _build_fund_nav_publication(
        instrument_id="fingerprinted-fund",
        instrument=instrument,
        rows=revised_rows,
    )

    fingerprint = first.projection_run["source_observation_fingerprint"]
    assert fingerprint == reordered.projection_run["source_observation_fingerprint"]
    assert fingerprint != revised.projection_run["source_observation_fingerprint"]
    assert len(str(fingerprint)) == 64


def test_public_item_timeout_uses_configured_limit_and_returns_failed_status(monkeypatch) -> None:
    captured: dict[str, object] = {}

    class FakeSettings:
        market_data_batch_item_timeout_seconds = 17

    class ImmediateTimeout:
        def __init__(self, seconds: int, instrument_id: str) -> None:
            captured["seconds"] = seconds
            captured["instrument_id"] = instrument_id

        def __enter__(self) -> None:
            raise market_data_ops.MarketDataItemTimeout(
                "Market data refresh timed out after 17 seconds for fund-a."
            )

        def __exit__(self, exc_type: object, exc: object, tb: object) -> None:
            return None

    def fail_unwrapped_refresh(**kwargs: object) -> dict[str, object]:
        raise AssertionError(f"unwrapped refresh unexpectedly called: {kwargs}")

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured["refresh_status"] = kwargs
        return {"instrument_id": kwargs["instrument_id"]}

    monkeypatch.setattr(market_data_ops, "get_settings", lambda: FakeSettings())
    monkeypatch.setattr(market_data_ops, "_BatchItemTimeout", ImmediateTimeout)
    monkeypatch.setattr(market_data_ops, "refresh_market_data", fail_unwrapped_refresh)
    monkeypatch.setattr(market_data_ops, "update_refresh_status", fake_update_refresh_status)

    record = market_data_ops.refresh_market_data_with_timeout(
        instrument_id="fund-a",
        updated_by="pytest",
        source="tushare",
    )

    assert record == {"instrument_id": "fund-a"}
    assert captured["seconds"] == 17
    assert captured["instrument_id"] == "fund-a"
    assert captured["refresh_status"] == {
        "instrument_id": "fund-a",
        "status": "failed",
        "message": "Market data refresh timed out after 17 seconds for fund-a.",
        "updated_by": "pytest",
        "mode": "api",
    }


def test_tushare_api_uses_sdk_with_configured_proxy_url(monkeypatch) -> None:
    captured: dict[str, object] = {}

    class FakeSettings:
        tushare_ready = True
        tushare_token = "secret-token"
        tushare_api_url = "https://ttx.dailyfetch.top/"

    class FakeFrame:
        def to_dict(self, orient: str) -> list[dict[str, object]]:
            assert orient == "records"
            return [{"ts_code": "000300.SH", "trade_date": "20260615", "close": "4200.12"}]

    class FakePro:
        def __init__(self) -> None:
            self._DataApi__http_url = ""

        def index_daily(self, **kwargs: object) -> FakeFrame:
            captured["http_url"] = self._DataApi__http_url
            captured["index_daily_kwargs"] = kwargs
            return FakeFrame()

    class FakeTushareModule:
        def __init__(self) -> None:
            self.pro = FakePro()

        def pro_api(self, token: str, timeout: int) -> FakePro:
            captured["token"] = token
            captured["timeout"] = timeout
            return self.pro

    monkeypatch.setattr(market_data_ops, "get_settings", lambda: FakeSettings())
    monkeypatch.setattr(
        market_data_ops.tushare_client,
        "ts",
        FakeTushareModule(),
    )

    rows = market_data_ops._call_tushare_api(
        api_name="index_daily",
        params={"ts_code": "000300.SH"},
        fields="ts_code,trade_date,close",
    )

    assert captured["token"] == "secret-token"
    assert captured["timeout"] == 30
    assert captured["http_url"] == "https://ttx.dailyfetch.top"
    assert captured["index_daily_kwargs"] == {
        "ts_code": "000300.SH",
        "fields": "ts_code,trade_date,close",
    }
    assert rows == [{"ts_code": "000300.SH", "trade_date": "20260615", "close": "4200.12"}]


def test_unchanged_fund_nav_publication_clears_stale_refresh_failure(
    monkeypatch,
) -> None:
    captured: dict[str, object] = {}

    class FakePublication:
        rows: list[dict[str, object]] = []
        projection_run = {"projection_run_id": "projection-run"}
        current_fund_nav_event_ids: list[str] = []
        current_fund_nav_reinvestment_evidence_ids: list[str] = []
        adjustment_factors: list[dict[str, object]] = []
        action_candidates: list[dict[str, object]] = []

    class FakeFundNavActionCandidateRepository:
        def __init__(self, _session_factory: object) -> None:
            pass

        def project_current(
            self,
            *,
            instrument_id: str,
            candidates: list[dict[str, object]],
        ) -> None:
            captured["candidate_projection"] = {
                "instrument_id": instrument_id,
                "candidates": candidates,
            }

    monkeypatch.setattr(
        market_data_ops,
        "get_instrument",
        lambda instrument_id: {
            "instrument_id": instrument_id,
            "market_data_updated_at": "2026-07-23T00:00:00Z",
            "refresh_status": {"status": "failed"},
        },
    )
    monkeypatch.setattr(
        market_data_ops,
        "_build_fund_nav_publication",
        lambda **_kwargs: FakePublication(),
    )
    monkeypatch.setattr(
        market_data_ops,
        "publish_fund_nav_history",
        lambda **_kwargs: {
            "record": {
                "instrument_id": "018654-of",
                "refresh_status": {"status": "failed"},
            },
            "changed": False,
        },
    )

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured["refresh_status"] = kwargs
        return {
            "instrument_id": kwargs["instrument_id"],
            "refresh_status": {
                "status": kwargs["status"],
                "message": kwargs["message"],
            },
        }

    monkeypatch.setattr(
        market_data_ops,
        "update_refresh_status",
        fake_update_refresh_status,
    )
    monkeypatch.setattr(
        market_data_ops,
        "FundNavActionCandidateRepository",
        FakeFundNavActionCandidateRepository,
    )

    record, publication = market_data_ops._publish_fund_nav_projection(
        instrument_id="018654-of",
        load_source_rows=lambda _instrument: [],
        source_provider="tushare:fund_nav",
        refresh_status="imported",
        updated_by="launchd-scheduler",
        mode="api",
        message_factory=lambda _publication: "Provider fetch succeeded.",
    )

    assert publication.rows == []
    assert record["refresh_status"]["status"] == "imported"
    assert captured["refresh_status"] == {
        "instrument_id": "018654-of",
        "status": "imported",
        "message": "Provider fetch succeeded.",
        "updated_by": "launchd-scheduler",
        "mode": "api",
    }
    assert captured["candidate_projection"] == {
        "instrument_id": "018654-of",
        "candidates": [],
    }


def test_tushare_refresh_imports_public_fund_nav(monkeypatch) -> None:
    captured: dict[str, object] = {}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        captured["api_call"] = kwargs
        return [
            {
                "ts_code": "018654.OF",
                "ann_date": "20260616",
                "end_date": None,
                "nav_date": "20260615",
                "unit_nav": "1.2345",
                "accum_nav": "1.3456",
                "adj_nav": "1.4567",
            },
            {
                "ts_code": "018654.OF",
                "ann_date": "20260613",
                "end_date": "20260612",
                "nav_date": "20260612",
                "unit_nav": "1.2000",
                "accum_nav": "1.3000",
                "adj_nav": "1.4000",
            },
        ]

    def fake_publish_fund_nav_history(**kwargs: object) -> dict[str, object]:
        captured["publish"] = kwargs
        return {
            "record": {"instrument_id": kwargs["instrument_id"]},
            "changed": True,
            "dirty_from": "2026-06-12",
        }

    def fake_record_raw_nav_observations(**kwargs: object) -> int:
        captured["raw"] = kwargs
        return len(kwargs["rows"])

    def fake_list_raw_nav_observations(**kwargs: object) -> list[dict[str, object]]:
        captured["list_raw"] = kwargs
        return market_data_ops._stamp_nav_value_statuses(
            [dict(row) for row in captured["raw"]["rows"]],
            status=str(captured["raw"]["status"]),
        )

    class FakeFundNavActionCandidateRepository:
        def __init__(self, _session_factory: object) -> None:
            pass

        def project_current(
            self,
            *,
            instrument_id: str,
            candidates: list[dict[str, object]],
        ) -> None:
            captured["candidate_projection"] = {
                "instrument_id": instrument_id,
                "candidates": candidates,
            }

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(
        market_data_ops,
        "get_instrument",
        lambda instrument_id: {
            "instrument_id": instrument_id,
            "currency": "CNY",
            "market_data": [
                {
                    "metric_family": "nav",
                    "quote_basis": "official_nav",
                    "as_of_date": "2026-06-12",
                    "value": "1.2000",
                    "status": "complete",
                },
                {
                    "metric_family": "nav",
                    "quote_basis": "total_return_nav",
                    "as_of_date": "2026-06-12",
                    "value": "1.4000",
                    "status": "complete",
                },
            ],
        },
    )
    monkeypatch.setattr(
        market_data_ops,
        "publish_fund_nav_history",
        fake_publish_fund_nav_history,
    )
    monkeypatch.setattr(
        market_data_ops,
        "record_raw_nav_observations",
        fake_record_raw_nav_observations,
    )
    monkeypatch.setattr(
        market_data_ops,
        "list_raw_nav_observations",
        fake_list_raw_nav_observations,
    )
    monkeypatch.setattr(
        market_data_ops,
        "_load_all_durable_nav_source_rows",
        lambda *, instrument_id, instrument: fake_list_raw_nav_observations(
            instrument_id=instrument_id,
        ),
    )
    monkeypatch.setattr(
        market_data_ops,
        "FundNavActionCandidateRepository",
        FakeFundNavActionCandidateRepository,
    )

    record = market_data_ops._refresh_from_tushare(
        instrument_id="018654-of",
        instrument={
            "instrument_id": "018654-of",
            "instrument_type": "fund",
            "currency": "CNY",
            "identifiers": [
                {"identifier_type": "ticker", "identifier_value": "018654.OF", "is_primary": True}
            ],
            "market_data": [
                {
                    "metric_family": "nav",
                    "quote_basis": "official_nav",
                    "as_of_date": "2026-06-12",
                    "value": "1.2000",
                    "status": "complete",
                },
                {
                    "metric_family": "nav",
                    "quote_basis": "total_return_nav",
                    "as_of_date": "2026-06-12",
                    "value": "1.4000",
                    "status": "complete",
                },
            ],
        },
        updated_by="test",
        full_history=False,
    )

    assert record == {"instrument_id": "018654-of"}
    assert captured["api_call"]["api_name"] == "fund_nav"
    assert captured["api_call"]["params"] == {"ts_code": "018654.OF"}
    assert captured["api_call"]["fields"] == "ts_code,ann_date,end_date,nav_date,unit_nav,accum_nav,adj_nav,update_flag"
    publish_payload = captured["publish"]
    assert publish_payload["mode"] == "api"
    assert [row["as_of_date"] for row in publish_payload["rows"]] == [
        "2026-06-12",
        "2026-06-15",
    ]
    latest_row = publish_payload["rows"][-1]
    assert latest_row["nav"] == "1.2345"
    assert "cash_cumulative_nav" not in latest_row
    assert str(latest_row["nav_with_dividend"]) == "1.4567000000000000"
    assert latest_row["nav_source_provider"] == "tushare:fund_nav"
    assert latest_row["nav_with_dividend_source_provider"] == "tushare:fund_nav"
    assert latest_row["nav_lineage"] == {
        "kind": "provider_explicit",
        "evidence": {"source_field": "unit_nav"},
    }
    assert latest_row["nav_with_dividend_lineage"] == {
        "kind": "provider_explicit",
        "evidence": {
            "factor_logical_key": "provider:2026-06-15",
            "source_field": "adj_nav",
            "provider_reported_total_return_nav": "1.4567",
        },
    }
    assert publish_payload["projection_run"]["projection_status"] == "complete"
    assert publish_payload["current_fund_nav_event_ids"] == []
    assert publish_payload["current_fund_nav_reinvestment_evidence_ids"] == []
    adjustment_factors = publish_payload["adjustment_factors"]
    assert len(adjustment_factors) == 2
    assert all(
        factor["factor_kind"] == "provider_implied"
        for factor in adjustment_factors
    )


def test_tushare_refresh_imports_index_close(monkeypatch) -> None:
    captured: dict[str, object] = {}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        captured["api_call"] = kwargs
        return [
            {"ts_code": "000300.SH", "trade_date": "20260615", "close": "4200.12"},
            {"ts_code": "000300.SH", "trade_date": "20260612", "close": "4190.00"},
        ]

    def fake_upsert_market_data_points(**kwargs: object) -> int:
        captured["upsert"] = kwargs
        return len(kwargs["rows"])

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured["refresh_status"] = kwargs
        return {"instrument_id": kwargs["instrument_id"]}

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(
        market_data_ops,
        "get_price_bar_coverage",
        lambda **_kwargs: {
            "row_count": 1,
            "first_date": "2026-06-12",
            "latest_date": "2026-06-12",
            "adjustment_factor_count": 0,
        },
    )
    monkeypatch.setattr(
        market_data_ops,
        "upsert_market_data_points",
        fake_upsert_market_data_points,
    )
    monkeypatch.setattr(market_data_ops, "update_refresh_status", fake_update_refresh_status)

    record = market_data_ops._refresh_from_tushare(
        instrument_id="000300-sh",
        instrument={
            "instrument_id": "000300-sh",
            "instrument_type": "index",
            "currency": "CNY",
            "identifiers": [
                {"identifier_type": "ticker", "identifier_value": "000300.SH", "is_primary": True}
            ],
            "market_data": [
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2026-06-12",
                    "value": "4190.00",
                    "status": "complete",
                }
            ],
        },
        updated_by="test",
        full_history=False,
    )

    assert record == {"instrument_id": "000300-sh"}
    assert captured["api_call"]["api_name"] == "index_daily"
    assert captured["api_call"]["params"]["ts_code"] == "000300.SH"
    assert captured["api_call"]["params"]["start_date"] == "20260613"
    assert captured["upsert"] == {
        "instrument_id": "000300-sh",
        "rows": [
            {
                "metric_family": "price",
                "quote_basis": "close",
                "as_of_date": market_data_ops.date(2026, 6, 15),
                "value": market_data_ops.Decimal("4200.12"),
                "currency": "CNY",
                "provider": "tushare:index_daily",
                "status": "complete",
            }
        ],
    }
    assert captured["refresh_status"]["status"] == "refreshed"
    assert captured["refresh_status"]["mode"] == "api"


def test_tushare_index_refresh_repairs_ohlcv_history_behind_existing_closes(
    monkeypatch,
) -> None:
    captured: dict[str, object] = {}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        captured["api_call"] = kwargs
        return [
            {
                "ts_code": "000300.SH",
                "trade_date": "20240102",
                "open": "4100",
                "high": "4210",
                "low": "4090",
                "close": "4200",
                "pre_close": "4095",
                "vol": "100",
                "amount": "420000",
            },
            {
                "ts_code": "000300.SH",
                "trade_date": "20260615",
                "open": "4195",
                "high": "4220",
                "low": "4180",
                "close": "4210",
                "pre_close": "4200",
                "vol": "120",
                "amount": "505200",
            },
        ]

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(
        market_data_ops,
        "get_price_bar_coverage",
        lambda **_kwargs: {
            "row_count": 0,
            "first_date": None,
            "latest_date": None,
            "adjustment_factor_count": 0,
        },
    )
    monkeypatch.setattr(
        market_data_ops,
        "upsert_market_data_points",
        lambda **kwargs: captured.setdefault("market_rows", kwargs["rows"])
        and len(kwargs["rows"]),
    )
    monkeypatch.setattr(
        market_data_ops,
        "upsert_price_bars",
        lambda **kwargs: captured.setdefault("bar_rows", kwargs["rows"])
        and len(kwargs["rows"]),
    )
    monkeypatch.setattr(
        market_data_ops,
        "update_refresh_status",
        lambda **kwargs: {"instrument_id": kwargs["instrument_id"]},
    )

    result = market_data_ops._refresh_from_tushare(
        instrument_id="000300-sh",
        instrument={
            "instrument_id": "000300-sh",
            "instrument_type": "index",
            "currency": "CNY",
            "identifiers": [
                {
                    "identifier_type": "ticker",
                    "identifier_value": "000300.SH",
                    "is_primary": True,
                }
            ],
            "market_data": [
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2024-01-02",
                    "value": "4200",
                    "status": "complete",
                },
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2026-06-12",
                    "value": "4200",
                    "status": "complete",
                },
            ],
        },
        updated_by="test",
        full_history=False,
    )

    assert result == {"instrument_id": "000300-sh"}
    assert captured["api_call"]["params"]["start_date"] == "20240102"
    assert [row["as_of_date"] for row in captured["market_rows"]] == [
        market_data_ops.date(2024, 1, 2),
        market_data_ops.date(2026, 6, 15),
    ]
    assert [row["as_of_date"] for row in captured["bar_rows"]] == [
        market_data_ops.date(2024, 1, 2),
        market_data_ops.date(2026, 6, 15),
    ]


def test_listed_security_refresh_repairs_missing_factor_and_bar_history(
    monkeypatch,
) -> None:
    captured: dict[str, object] = {"api_calls": []}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        captured["api_calls"].append(kwargs)
        if kwargs["api_name"] == "fund_adj":
            return [
                {"ts_code": "513050.SH", "trade_date": "20240102", "adj_factor": "1"},
                {"ts_code": "513050.SH", "trade_date": "20260716", "adj_factor": "2"},
            ]
        return [
            {
                "ts_code": "513050.SH",
                "trade_date": "20240102",
                "open": "0.90",
                "high": "0.92",
                "low": "0.89",
                "close": "0.91",
                "pre_close": "0.90",
                "vol": "1000",
                "amount": "910",
            },
            {
                "ts_code": "513050.SH",
                "trade_date": "20260716",
                "open": "1.14",
                "high": "1.16",
                "low": "1.13",
                "close": "1.15",
                "pre_close": "1.14",
                "vol": "2000",
                "amount": "2300",
            },
        ]

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(
        market_data_ops,
        "get_price_bar_coverage",
        lambda **_kwargs: {
            "row_count": 6,
            "first_date": "2026-07-09",
            "latest_date": "2026-07-16",
            "adjustment_factor_count": 5,
        },
    )
    monkeypatch.setattr(
        market_data_ops,
        "upsert_market_data_points",
        lambda **kwargs: len(kwargs["rows"]),
    )
    monkeypatch.setattr(
        market_data_ops,
        "upsert_price_bars",
        lambda **kwargs: captured.setdefault("bar_rows", kwargs["rows"])
        and len(kwargs["rows"]),
    )
    monkeypatch.setattr(
        market_data_ops,
        "update_refresh_status",
        lambda **kwargs: {"instrument_id": kwargs["instrument_id"]},
    )
    monkeypatch.setattr(
        market_data_ops,
        "upsert_corporate_action_event",
        lambda **_kwargs: None,
    )

    result = market_data_ops._refresh_from_tushare(
        instrument_id="513050-sh",
        instrument={
            "instrument_id": "513050-sh",
            "instrument_type": "etf",
            "currency": "CNY",
            "identifiers": [
                {
                    "identifier_type": "ticker",
                    "identifier_value": "513050.SH",
                    "is_primary": True,
                }
            ],
            "quote_selection_policy": {
                "trading": ["last", "close"],
                "valuation": ["close", "last"],
                "total_return": ["adjusted_close", "close", "last"],
                "chart": ["adjusted_close", "close", "last"],
                "reference": ["close", "last"],
            },
            "market_data": [
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2024-01-02",
                    "value": "0.91",
                    "status": "complete",
                },
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2026-07-16",
                    "value": "1.15",
                    "status": "complete",
                },
                {
                    "metric_family": "price",
                    "quote_basis": "adjusted_close",
                    "as_of_date": "2026-07-16",
                    "value": "1.15",
                    "provider": "tushare:fund_adj:qfq:latest_factor=2",
                    "status": "complete",
                },
            ],
        },
        updated_by="test",
        full_history=False,
    )

    assert result == {"instrument_id": "513050-sh"}
    assert captured["api_calls"][0]["params"]["start_date"] == "20240102"
    assert [row["as_of_date"] for row in captured["bar_rows"]] == [
        market_data_ops.date(2024, 1, 2),
        market_data_ops.date(2026, 7, 16),
    ]
    assert all(row["adjustment_factor"] is not None for row in captured["bar_rows"])


def test_tushare_price_refresh_starts_at_2024_when_no_existing_history(monkeypatch) -> None:
    captured: dict[str, object] = {}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        captured.setdefault("api_calls", []).append(kwargs)
        if kwargs["api_name"] == "fund_adj":
            return [
                {"ts_code": "513050.SH", "trade_date": "20240102", "adj_factor": "1.0"},
            ]
        return [
            {"ts_code": "513050.SH", "trade_date": "20231229", "close": "0.90"},
            {"ts_code": "513050.SH", "trade_date": "20240102", "close": "0.91"},
        ]

    def fake_upsert_market_data_points(**kwargs: object) -> int:
        captured["upsert"] = kwargs
        return len(kwargs["rows"])

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured["refresh_status"] = kwargs
        return {"instrument_id": kwargs["instrument_id"]}

    def fake_upsert_quote_selection_policy(**kwargs: object) -> dict[str, object]:
        captured["quote_selection_policy"] = kwargs
        return {"instrument_id": kwargs["instrument_id"]}

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(
        market_data_ops,
        "upsert_market_data_points",
        fake_upsert_market_data_points,
    )
    monkeypatch.setattr(market_data_ops, "update_refresh_status", fake_update_refresh_status)
    monkeypatch.setattr(
        market_data_ops,
        "upsert_quote_selection_policy",
        fake_upsert_quote_selection_policy,
    )

    market_data_ops._refresh_from_tushare(
        instrument_id="513050-sh",
        instrument={
            "instrument_id": "513050-sh",
            "instrument_type": "fund",
            "currency": "CNY",
            "identifiers": [
                {"identifier_type": "ticker", "identifier_value": "513050.SH", "is_primary": True}
            ],
            "market_data": [],
        },
        updated_by="test",
        full_history=False,
    )

    assert captured["api_calls"][0]["api_name"] == "fund_daily"
    assert captured["api_calls"][0]["params"]["start_date"] == "20240101"
    assert captured["quote_selection_policy"] == {
        "instrument_id": "513050-sh",
        "quote_selection_policy": {
            "trading": ["last", "close"],
            "valuation": ["close", "last"],
            "total_return": ["adjusted_close", "close", "last"],
            "chart": ["adjusted_close", "close", "last"],
            "reference": ["close", "last"],
        },
    }
    assert captured["upsert"]["instrument_id"] == "513050-sh"
    assert captured["upsert"]["rows"] == [
        {
            "metric_family": "price",
            "quote_basis": "close",
            "as_of_date": market_data_ops.date(2024, 1, 2),
            "value": market_data_ops.Decimal("0.91"),
            "currency": "CNY",
            "provider": "tushare:fund_daily",
            "status": "complete",
        },
        {
            "metric_family": "price",
            "quote_basis": "adjusted_close",
            "as_of_date": market_data_ops.date(2024, 1, 2),
            "value": "0.91",
            "currency": "CNY",
            "provider": "tushare:fund_adj:qfq:latest_factor=1",
            "status": "complete",
        },
    ]


def test_empty_listed_security_response_preserves_existing_history(monkeypatch) -> None:
    captured: dict[str, object] = {}

    monkeypatch.setattr(
        market_data_ops,
        "_call_tushare_api",
        lambda **_kwargs: [],
    )

    def fail_if_written(**_kwargs: object) -> int:
        raise AssertionError("An empty provider response must not rewrite canonical history.")

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured.update(kwargs)
        return {"instrument_id": kwargs["instrument_id"]}

    monkeypatch.setattr(
        market_data_ops,
        "upsert_market_data_points",
        fail_if_written,
    )
    monkeypatch.setattr(market_data_ops, "upsert_price_bars", fail_if_written)
    monkeypatch.setattr(
        market_data_ops,
        "update_refresh_status",
        fake_update_refresh_status,
    )

    result = market_data_ops._refresh_from_tushare(
        instrument_id="513050-sh",
        instrument={
            "instrument_id": "513050-sh",
            "instrument_type": "etf",
            "currency": "CNY",
            "identifiers": [
                {
                    "identifier_type": "ticker",
                    "identifier_value": "513050.SH",
                    "is_primary": True,
                }
            ],
            "market_data": [
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2026-07-16",
                    "value": "1.145",
                    "status": "complete",
                },
                {
                    "metric_family": "price",
                    "quote_basis": "adjusted_close",
                    "as_of_date": "2026-07-16",
                    "value": "1.145",
                    "provider": "tushare:fund_adj:qfq:latest_factor=1",
                    "status": "complete",
                },
            ],
            "quote_selection_policy": {
                "trading": ["last", "close"],
                "valuation": ["close", "last"],
                "total_return": ["adjusted_close", "close", "last"],
                "chart": ["adjusted_close", "close", "last"],
                "reference": ["close", "last"],
            },
        },
        updated_by="test",
        full_history=True,
    )

    assert result == {"instrument_id": "513050-sh"}
    assert captured["status"] == "failed"
    assert "preserved" in str(captured["message"])


def test_tushare_listed_security_never_persists_complete_close_without_adjusted_close(
    monkeypatch,
) -> None:
    captured: dict[str, object] = {}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        if kwargs["api_name"] == "adj_factor":
            return [
                {"ts_code": "600000.SH", "trade_date": "20260714", "adj_factor": "1"},
            ]
        return [
            {"ts_code": "600000.SH", "trade_date": "20260714", "close": "10.00"},
            {"ts_code": "600000.SH", "trade_date": "20260715", "close": "10.10"},
        ]

    def fake_upsert_market_data_points(**kwargs: object) -> int:
        captured["upsert"] = kwargs
        return len(kwargs["rows"])

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured["refresh_status"] = kwargs
        return {"instrument_id": kwargs["instrument_id"]}

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(
        market_data_ops,
        "upsert_market_data_points",
        fake_upsert_market_data_points,
    )
    monkeypatch.setattr(market_data_ops, "update_refresh_status", fake_update_refresh_status)

    market_data_ops._refresh_from_tushare(
        instrument_id="600000-sh",
        instrument={
            "instrument_id": "600000-sh",
            "instrument_type": "equity",
            "currency": "CNY",
            "identifiers": [
                {
                    "identifier_type": "ticker",
                    "identifier_value": "600000.SH",
                    "is_primary": True,
                }
            ],
            "market_data": [],
            "quote_selection_policy": {
                "trading": ["last", "close"],
                "valuation": ["close", "last"],
                "total_return": ["adjusted_close", "close", "last"],
                "chart": ["adjusted_close", "close", "last"],
                "reference": ["close", "last"],
            },
        },
        updated_by="test",
        full_history=False,
    )

    points = captured["upsert"]["rows"]
    complete_close_dates = {
        point["as_of_date"]
        for point in points
        if point["quote_basis"] == "close" and point["status"] == "complete"
    }
    adjusted_close_dates = {
        point["as_of_date"]
        for point in points
        if point["quote_basis"] == "adjusted_close" and point["status"] == "complete"
    }
    assert complete_close_dates <= adjusted_close_dates
    assert next(
        point
        for point in points
        if point["quote_basis"] == "close"
        and point["as_of_date"] == market_data_ops.date(2026, 7, 15)
    )["status"] == "partial"
    assert captured["refresh_status"]["status"] == "partial"
    assert "missing adjustment factor" in captured["refresh_status"]["message"]
    assert "2026-07-15" in captured["refresh_status"]["message"]


def test_tushare_share_split_detection_is_review_only_until_issuer_confirmation() -> None:
    candidates = market_data_ops._detect_tushare_share_splits(
        factors={
            market_data_ops.date(2026, 3, 27): market_data_ops.Decimal("1"),
            market_data_ops.date(2026, 3, 30): market_data_ops.Decimal("2"),
        },
        close_by_date={
            market_data_ops.date(2026, 3, 27): market_data_ops.Decimal("1.660"),
            market_data_ops.date(2026, 3, 30): market_data_ops.Decimal("0.852"),
        },
    )

    assert len(candidates) == 1
    assert candidates[0]["effective_date"] == market_data_ops.date(2026, 3, 30)
    assert candidates[0]["new_units"] == "2"
    assert candidates[0]["old_units"] == "1"
    assert candidates[0]["status"] == "detected"
    assert candidates[0]["provenance"]["detection_method"] == "tushare_factor_price_continuity/v1"


def test_tushare_factor_change_without_inverse_raw_price_move_is_not_a_split() -> None:
    candidates = market_data_ops._detect_tushare_share_splits(
        factors={
            market_data_ops.date(2026, 3, 27): market_data_ops.Decimal("1"),
            market_data_ops.date(2026, 3, 30): market_data_ops.Decimal("2"),
        },
        close_by_date={
            market_data_ops.date(2026, 3, 27): market_data_ops.Decimal("1.00"),
            market_data_ops.date(2026, 3, 30): market_data_ops.Decimal("1.01"),
        },
    )

    assert candidates == []


def test_tushare_full_history_is_capped_at_2024(monkeypatch) -> None:
    captured: dict[str, object] = {}

    def fake_call_tushare_api(**kwargs: object) -> list[dict[str, object]]:
        captured["api_call"] = kwargs
        return []

    def fake_update_refresh_status(**kwargs: object) -> dict[str, object]:
        captured["refresh_status"] = kwargs
        return {"instrument_id": kwargs["instrument_id"]}

    monkeypatch.setattr(market_data_ops, "_call_tushare_api", fake_call_tushare_api)
    monkeypatch.setattr(market_data_ops, "update_refresh_status", fake_update_refresh_status)

    market_data_ops._refresh_from_tushare(
        instrument_id="000300-sh",
        instrument={
            "instrument_id": "000300-sh",
            "instrument_type": "index",
            "currency": "CNY",
            "identifiers": [
                {"identifier_type": "ticker", "identifier_value": "000300.SH", "is_primary": True}
            ],
            "market_data": [
                {
                    "metric_family": "price",
                    "quote_basis": "close",
                    "as_of_date": "2026-06-12",
                    "value": "4190.00",
                    "status": "complete",
                }
            ],
        },
        updated_by="test",
        full_history=True,
    )

    assert captured["api_call"]["api_name"] == "index_daily"
    assert captured["api_call"]["params"]["start_date"] == "20240101"


def test_parse_nav_rows_from_label_snapshot_matrix_extracts_nav_values() -> None:
    matrix = [
        ["招商证券股份有限公司_九慕稳健1号私募证券投资基金_专用表", None],
        ["日期：2026-04-14", None],
        ["期初单位净值", "1.0980"],
        ["单位净值", "1.1135"],
        ["累计派现金额", "0"],
        ["累计单位净值", "1.1135"],
    ]

    parsed = _parse_nav_rows_from_label_snapshot_matrix(matrix)

    assert len(parsed) == 1
    assert parsed[0]["as_of_date"] == "2026-04-14"
    assert str(parsed[0]["nav"]) == "1.1135"
    assert str(parsed[0]["cash_cumulative_nav"]) == "1.1135"
    assert parsed[0]["instrument_name"] == "九慕稳健1号私募证券投资基金"
    assert "currency" not in parsed[0]
    assert "frequency" not in parsed[0]


def test_label_snapshot_attachment_extracts_exact_filename_code_and_workbook_name() -> None:
    parsed = _parse_nav_rows_from_attachment(
        attachment_name=(
            "SBMM07_国泰君安期货CTA因子组合3号集合资产管理计划_"
            "3级科目估值表_20260626.xls"
        ),
        attachment_bytes=_workbook_bytes(
            [
                [
                    "国泰海通___国泰君安期货CTA因子组合3号集合资产管理计划"
                    "___专用表"
                ],
                ["估值日期：2026-06-26", "单位净值：1.0484"],
            ]
        ),
        parser_profile="label_nav_snapshot",
    )

    assert parsed == [
        {
            "as_of_date": "2026-06-26",
            "nav": Decimal("1.0484"),
            "instrument_code": "SBMM07",
            "instrument_name": "国泰君安期货CTA因子组合3号集合资产管理计划",
            "_source_sheet": "Sheet",
            "_source_sheets": ["Sheet"],
        }
    ]
