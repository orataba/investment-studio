from __future__ import annotations

from contextlib import AbstractContextManager
import csv
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP, localcontext
from io import BytesIO, StringIO
import hashlib
import json
import logging
import multiprocessing
import re
import signal
import threading
import time
import warnings
from typing import Any, Callable
from zipfile import BadZipFile, ZipFile

from portfolio_ops_instrument_core import (
    parse_positive_market_data_value,
    validate_nav_history_instrument_type,
)

from platform_app.core.settings import get_settings
from platform_app.db.session import get_session_factory
from platform_app.services.email_ingestion import (
    EmailIngestionBusyError,
    EmailIngestionError,
    ingest_email_nav,
)
from platform_app.services.email_ingestion.repository import EmailIngestionRepository
from platform_app.services.fund_nav_action_candidates import (
    FundNavActionCandidateRepository,
)
from platform_app.services.instrument_store import (
    get_price_bar_coverage,
    get_instrument,
    list_instrument_ids_with_nav_history_before,
    list_instruments,
    list_stale_current_fund_nav_projections,
    publish_fund_nav_history,
    StaleFundNavPublicationError,
    update_refresh_status,
    upsert_quote_selection_policy,
    upsert_corporate_action_event,
    upsert_market_data_points,
    upsert_price_bars,
)
from platform_app.services.nav_raw_store import (
    list_raw_nav_observations,
    record_raw_nav_observations,
)
from platform_app.services import tushare_client

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional dependency
    load_workbook = None

try:
    import xlrd
except ImportError:  # pragma: no cover - optional dependency
    xlrd = None

LOGGER = logging.getLogger("portfolio_ops.market_data_ops")


class MarketDataItemTimeout(TimeoutError):
    pass


class _BatchItemTimeout:
    def __init__(self, seconds: int, instrument_id: str) -> None:
        self.seconds = max(0, int(seconds or 0))
        self.instrument_id = instrument_id
        self._enabled = (
            self.seconds > 0
            and threading.current_thread() is threading.main_thread()
            and hasattr(signal, "SIGALRM")
            and hasattr(signal, "setitimer")
        )
        self._previous_handler: Any = None
        self._previous_timer: tuple[float, float] = (0.0, 0.0)

    def __enter__(self) -> None:
        if not self._enabled:
            return
        self._previous_handler = signal.getsignal(signal.SIGALRM)
        self._previous_timer = signal.getitimer(signal.ITIMER_REAL)
        signal.signal(signal.SIGALRM, self._raise_timeout)
        signal.setitimer(signal.ITIMER_REAL, self.seconds)

    def __exit__(self, exc_type: object, exc: object, tb: object) -> None:
        if not self._enabled:
            return
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, self._previous_handler)
        previous_delay, previous_interval = self._previous_timer
        if previous_delay > 0:
            signal.setitimer(signal.ITIMER_REAL, previous_delay, previous_interval)

    def _raise_timeout(self, signum: int, frame: object) -> None:
        raise MarketDataItemTimeout(
            f"Market data refresh timed out after {self.seconds} seconds for {self.instrument_id}."
        )


def market_data_item_timeout(instrument_id: str) -> AbstractContextManager[None]:
    settings = get_settings()
    return _BatchItemTimeout(settings.market_data_batch_item_timeout_seconds, instrument_id)

NAV_IMPORT_HEADER_MAP = {
    "date": "as_of_date",
    "日期": "as_of_date",
    "净值日期": "as_of_date",
    "业务日期": "as_of_date",
    "交易日期": "as_of_date",
    "估值基准日": "as_of_date",
    "asofdate": "as_of_date",
    "as_of_date": "as_of_date",
    "trade_date": "as_of_date",
    "navdate": "as_of_date",
    "nav": "nav",
    "净值": "nav",
    "单位净值": "nav",
    "单位净值元": "nav",
    "单位净值元份": "nav",
    "资产份额净值元": "nav",
    "实际净值": "nav",
    "navwithdividend": "nav_with_dividend",
    "nav_with_dividend": "nav_with_dividend",
    "累计净值": "cash_cumulative_nav",
    "累计净值元": "cash_cumulative_nav",
    "累计单位净值": "cash_cumulative_nav",
    "累计单位净值元": "cash_cumulative_nav",
    "累计单位净值元份": "cash_cumulative_nav",
    "资产份额累计净值元": "cash_cumulative_nav",
    "实际累计净值": "cash_cumulative_nav",
    "复权净值": "nav_with_dividend",
    "复权单位净值": "nav_with_dividend",
    "分红再投资净值": "nav_with_dividend",
    "分红再投资单位净值": "nav_with_dividend",
    "复利净值": "nav_with_dividend",
    "currency": "currency",
    "ccy": "currency",
    "币种": "currency",
    "frequency": "frequency",
    "freq": "frequency",
    "频率": "frequency",
    "productcode": "instrument_code",
    "产品代码": "instrument_code",
    "产品编码": "instrument_code",
    "instrumentcode": "instrument_code",
    "资产代码": "instrument_code",
    "fundcode": "instrument_code",
    "基金代码": "instrument_code",
    "productname": "instrument_name",
    "产品名称": "instrument_name",
    "instrumentname": "instrument_name",
    "资产名称": "instrument_name",
    "fundname": "instrument_name",
    "基金名称": "instrument_name",
    # Common administrator export headers retain an English translation after
    # punctuation is stripped by ``_normalize_nav_header``.
    "日期navasofdate": "as_of_date",
    "产品名称fundname": "instrument_name",
    "单位净值navshare": "nav",
    "累计单位净值accumulatednavshare": "cash_cumulative_nav",
    "协会备案编码fundfillingcode": "instrument_code",
    "协会备案编码fundfilingcode": "instrument_code",
}

LABEL_SNAPSHOT_FIELD_ALIASES = {
    "as_of_date": ("日期", "净值日期", "估值日期"),
    "nav": ("单位净值", "基金份额净值"),
    "cash_cumulative_nav": ("累计单位净值", "基金份额累计净值"),
    "nav_with_dividend": ("复权单位净值", "分红再投资净值", "复利净值"),
}
LABEL_SNAPSHOT_IDENTITY_ALIASES = {
    "instrument_code": ("产品代码", "产品编码", "基金代码"),
    "instrument_name": ("产品名称", "产品全称", "基金名称"),
}

TOTAL_RETURN_NAV_DECIMAL_PLACES = Decimal("0.0000000000000001")
FUND_NAV_FACTOR_DECIMAL_PLACES = Decimal("0.000000000000000001")
NAV_SEMANTIC_ZERO_TOLERANCE = Decimal("0.00000001")
MAX_WORKBOOK_ROWS = 100_000
MAX_WORKBOOK_COLUMNS = 256
MAX_WORKBOOK_CELLS = 1_000_000
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
TUSHARE_PROFILE_ALIASES = {"tushare", "tushare_pro", "tushare-pro"}
TUSHARE_PRICE_SUFFIXES = {"SH", "SZ"}
TUSHARE_INDEX_SUFFIXES = {"SH", "SZ", "CSI", "CNI"}
TUSHARE_HISTORY_START_DATE = date(2024, 1, 1)
TUSHARE_LISTED_SECURITY_QUOTE_SELECTION_POLICY: dict[str, list[str]] = {
    "trading": ["last", "close"],
    "valuation": ["close", "last"],
    "total_return": ["adjusted_close", "close", "last"],
    "chart": ["adjusted_close", "close", "last"],
    "reference": ["close", "last"],
}


def _normalize_nav_header(value: str) -> str:
    return re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", value.lower())


def _excel_serial_to_date(value: float) -> date | None:
    if value <= 0:
        return None
    try:
        return (datetime(1899, 12, 30) + timedelta(days=value)).date()
    except OverflowError:
        return None


def _parse_nav_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return _excel_serial_to_date(float(value))

    normalized = value.strip()
    if not normalized:
        return None
    if re.fullmatch(r"\d{8}", normalized):
        try:
            return datetime.strptime(normalized, "%Y%m%d").date()
        except ValueError:
            pass
    if re.fullmatch(r"\d+(?:\.\d+)?", normalized):
        return _excel_serial_to_date(float(normalized))
    for fmt in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%m-%d-%Y",
        "%d/%m/%Y",
        "%Y%m%d",
        "%Y-%m-%d %H:%M:%S",
        "%Y/%m/%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M",
        "%Y年%m月%d日",
        "%Y年%m月%d日 %H:%M:%S",
        "%Y年%m月%d日 %H:%M",
    ):
        try:
            return datetime.strptime(normalized, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(normalized)
    except ValueError:
        return None


def _parse_nav_decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    normalized = str(value).strip().replace(",", "")
    if not normalized:
        return None
    try:
        return parse_positive_market_data_value(normalized)
    except ValueError:
        return None


def _parse_nonnegative_decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    normalized = str(value).strip().replace(",", "")
    if not normalized:
        return None
    try:
        parsed = Decimal(normalized)
    except (ArithmeticError, ValueError):
        return None
    if not parsed.is_finite() or parsed < 0:
        return None
    return parsed


def _format_total_return_nav_decimal(value: Decimal) -> Decimal:
    return value.quantize(TOTAL_RETURN_NAV_DECIMAL_PLACES, rounding=ROUND_HALF_UP)


def _detect_delimiter(line: str) -> str:
    if "\t" in line:
        return "\t"
    if line.count(";") > line.count(","):
        return ";"
    return ","


def _matrix_from_text(raw_text: str) -> list[list[object]]:
    lines = [line for line in raw_text.replace("\r\n", "\n").split("\n") if line.strip()]
    if not lines:
        return []
    delimiter = _detect_delimiter(lines[0])
    return [
        [cell.strip() for cell in row]
        for row in csv.reader(StringIO("\n".join(lines)), delimiter=delimiter)
    ]


def _bounded_matrix(rows: Any) -> list[list[object]]:
    matrix: list[list[object]] = []
    cell_count = 0
    for row_index, row in enumerate(rows, start=1):
        if row_index > MAX_WORKBOOK_ROWS:
            raise ValueError(f"NAV workbook exceeds {MAX_WORKBOOK_ROWS} rows.")
        values = list(row)
        if len(values) > MAX_WORKBOOK_COLUMNS:
            raise ValueError(
                f"NAV workbook exceeds {MAX_WORKBOOK_COLUMNS} columns on row {row_index}."
            )
        cell_count += len(values)
        if cell_count > MAX_WORKBOOK_CELLS:
            raise ValueError(f"NAV workbook exceeds {MAX_WORKBOOK_CELLS} cells.")
        matrix.append(values)
    return matrix


def _xlsx_matrices(file_bytes: bytes) -> list[tuple[str, list[list[object]]]]:
    if load_workbook is None:
        return []
    try:
        with ZipFile(BytesIO(file_bytes)) as archive:
            expanded_size = sum(info.file_size for info in archive.infolist())
    except BadZipFile as error:
        raise ValueError("Invalid xlsx workbook.") from error
    if expanded_size > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError(
            f"NAV xlsx expands beyond {MAX_XLSX_UNCOMPRESSED_BYTES} bytes."
        )
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
        )
        workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        matrices: list[tuple[str, list[list[object]]]] = []
        workbook_cell_count = 0
        for sheet in workbook.worksheets:
            # Some administrators emit a bogus A1:A1 dimension even though the
            # worksheet contains a full table. Resetting dimensions keeps the
            # memory-safe read-only reader while making it inspect real cells.
            reset_dimensions = getattr(sheet, "reset_dimensions", None)
            if callable(reset_dimensions):
                reset_dimensions()
            matrix = _bounded_matrix(sheet.iter_rows(values_only=True))
            workbook_cell_count += sum(len(row) for row in matrix)
            if workbook_cell_count > MAX_WORKBOOK_CELLS:
                raise ValueError(
                    f"NAV workbook exceeds {MAX_WORKBOOK_CELLS} cells across sheets."
                )
            matrices.append((str(sheet.title), matrix))
        return matrices
    finally:
        workbook.close()


def _xls_matrices(file_bytes: bytes) -> list[tuple[str, list[list[object]]]]:
    if xlrd is None:
        return []
    diagnostics = StringIO()
    workbook = xlrd.open_workbook(
        file_contents=file_bytes,
        on_demand=True,
        logfile=diagnostics,
    )
    if diagnostic_text := diagnostics.getvalue().strip():
        LOGGER.debug("xlrd nonfatal workbook diagnostics: %s", diagnostic_text)
    try:
        matrices: list[tuple[str, list[list[object]]]] = []
        workbook_cell_count = 0
        for sheet in (
            workbook.sheet_by_index(index) for index in range(workbook.nsheets)
        ):
            matrix = _bounded_matrix(
                sheet.row_values(row_index) for row_index in range(sheet.nrows)
            )
            workbook_cell_count += sum(len(row) for row in matrix)
            if workbook_cell_count > MAX_WORKBOOK_CELLS:
                raise ValueError(
                    f"NAV workbook exceeds {MAX_WORKBOOK_CELLS} cells across sheets."
                )
            matrices.append((str(sheet.name), matrix))
        return matrices
    finally:
        workbook.release_resources()


def _extract_attachment_text_from_matrix(matrix: list[list[object]]) -> str:
    parts: list[str] = []
    for row in matrix:
        for value in row:
            text = "" if value is None else str(value).strip()
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_numeric_from_text(label: str, value: str) -> Decimal | None:
    normalized_label = label.strip()
    candidates = [
        value,
        value.replace("：", ":"),
    ]
    for candidate in candidates:
        if normalized_label and normalized_label in candidate:
            tail = candidate.split(normalized_label, 1)[-1].lstrip(":： ")
            parsed = _parse_nav_decimal(tail)
            if parsed is not None:
                return parsed
    return None


def _parse_nav_rows_from_matrix(matrix: list[list[object]]) -> list[dict[str, object]]:
    if not matrix:
        return []
    header: list[str] | None = None
    start_index = 0
    for idx, row in enumerate(matrix[:12]):
        mapped = [NAV_IMPORT_HEADER_MAP.get(_normalize_nav_header(str(cell)), "") for cell in row]
        if "as_of_date" in mapped and any(
            key in mapped for key in ("nav", "cash_cumulative_nav", "nav_with_dividend")
        ):
            header = mapped
            start_index = idx + 1
            break
    if header is None:
        return []

    rows: list[dict[str, object]] = []
    for raw_row in matrix[start_index:]:
        values = ["" if value is None else str(value).strip() for value in raw_row]
        if not any(values):
            continue
        row_data: dict[str, object] = {}
        for index, column in enumerate(header):
            if not column:
                continue
            cell = values[index] if index < len(values) else ""
            if column == "as_of_date":
                parsed_date = _parse_nav_date(cell)
                if parsed_date is not None:
                    row_data["as_of_date"] = parsed_date.isoformat()
            elif column in {"nav", "cash_cumulative_nav", "nav_with_dividend"}:
                parsed_value = _parse_nav_decimal(cell)
                if parsed_value is not None:
                    row_data[column] = parsed_value
                    if column == "nav_with_dividend":
                        row_data["_total_return_source_field"] = (
                            "provider_reinvested_nav_column"
                        )
            elif column == "currency" and cell:
                row_data["currency"] = cell.upper()
            elif column == "frequency" and cell:
                row_data["frequency"] = cell.lower()
            elif column in {"instrument_code", "instrument_name"} and cell:
                row_data[column] = cell
        if row_data.get("as_of_date") and any(
            row_data.get(key) is not None
            for key in ("nav", "cash_cumulative_nav", "nav_with_dividend")
        ):
            rows.append(row_data)
    rows.sort(key=lambda item: str(item["as_of_date"]))
    return rows


def _parse_nav_rows_from_label_snapshot_matrix(matrix: list[list[object]]) -> list[dict[str, object]]:
    if not matrix:
        return []

    found_date: date | None = None
    found_nav: Decimal | None = None
    found_cash_cumulative_nav: Decimal | None = None
    found_total_return_nav: Decimal | None = None
    found_instrument_code = ""
    found_instrument_name = ""
    normalized_aliases = {
        field: tuple(_normalize_nav_header(alias) for alias in aliases)
        for field, aliases in LABEL_SNAPSHOT_FIELD_ALIASES.items()
    }
    normalized_identity_aliases = {
        field: tuple(_normalize_nav_header(alias) for alias in aliases)
        for field, aliases in LABEL_SNAPSHOT_IDENTITY_ALIASES.items()
    }

    for row in matrix:
        string_values = ["" if value is None else str(value).strip() for value in row]
        if not any(string_values):
            continue
        for index, cell in enumerate(string_values):
            normalized_cell = _normalize_nav_header(cell)
            next_value = string_values[index + 1] if index + 1 < len(string_values) else ""

            if found_date is None and re.fullmatch(
                r"\s*\d{4}年\d{1,2}月\d{1,2}日\s*",
                cell,
            ):
                found_date = _parse_nav_date(cell)

            if (
                not found_instrument_code
                and normalized_cell
                in normalized_identity_aliases["instrument_code"]
                and next_value
            ):
                found_instrument_code = next_value.strip().upper()
            if (
                not found_instrument_name
                and normalized_cell
                in normalized_identity_aliases["instrument_name"]
                and next_value
            ):
                found_instrument_name = next_value.strip()

            if not found_instrument_code:
                for label in LABEL_SNAPSHOT_IDENTITY_ALIASES["instrument_code"]:
                    if re.match(rf"^\s*{re.escape(label)}\s*[：:]", cell):
                        found_instrument_code = re.sub(
                            rf"^\s*{re.escape(label)}\s*[：:]\s*",
                            "",
                            cell,
                        ).strip().upper()
                        break
            if not found_instrument_name:
                for label in LABEL_SNAPSHOT_IDENTITY_ALIASES["instrument_name"]:
                    if re.match(rf"^\s*{re.escape(label)}\s*[：:]", cell):
                        found_instrument_name = re.sub(
                            rf"^\s*{re.escape(label)}\s*[：:]\s*",
                            "",
                            cell,
                        ).strip()
                        break
            if not found_instrument_name:
                dedicated_table = re.fullmatch(
                    r".+?[_＿]+(?P<instrument_name>[^_＿]+?)[_＿]+专用表",
                    cell,
                )
                if dedicated_table is not None:
                    found_instrument_name = dedicated_table.group(
                        "instrument_name"
                    ).strip()

            if found_date is None and normalized_cell in normalized_aliases["as_of_date"]:
                found_date = _parse_nav_date(next_value)
            if found_nav is None and normalized_cell in normalized_aliases["nav"]:
                found_nav = _parse_nav_decimal(next_value)
            if (
                found_cash_cumulative_nav is None
                and normalized_cell in normalized_aliases["cash_cumulative_nav"]
            ):
                found_cash_cumulative_nav = _parse_nav_decimal(next_value)
            if (
                found_total_return_nav is None
                and normalized_cell in normalized_aliases["nav_with_dividend"]
            ):
                found_total_return_nav = _parse_nav_decimal(next_value)

            if found_date is None and any(alias in cell for alias in LABEL_SNAPSHOT_FIELD_ALIASES["as_of_date"]):
                found_date = _parse_nav_date(
                    re.sub(r"^.*?[：:]\s*", "", cell).strip()
                )
            if found_nav is None and not any(
                longer_label in cell
                for longer_label in (
                    "累计单位净值",
                    "复权单位净值",
                    "分红再投资净值",
                )
            ):
                found_nav = _extract_numeric_from_text("单位净值", cell)
            if found_cash_cumulative_nav is None:
                found_cash_cumulative_nav = _extract_numeric_from_text("累计单位净值", cell)
            if found_total_return_nav is None:
                for label in ("复权单位净值", "分红再投资净值", "复利净值"):
                    found_total_return_nav = _extract_numeric_from_text(label, cell)
                    if found_total_return_nav is not None:
                        break

    if found_date is None or found_nav is None:
        return []

    row: dict[str, object] = {
        "as_of_date": found_date.isoformat(),
        "nav": found_nav,
    }
    if found_instrument_code:
        row["instrument_code"] = found_instrument_code
    if found_instrument_name:
        row["instrument_name"] = found_instrument_name
    if found_cash_cumulative_nav is not None:
        row["cash_cumulative_nav"] = found_cash_cumulative_nav
    if found_total_return_nav is not None:
        row["nav_with_dividend"] = found_total_return_nav
        row["_total_return_source_field"] = "provider_reinvested_nav_label"
    return [row]


def _label_snapshot_filename_code(attachment_name: str) -> str:
    filename = attachment_name.replace("\\", "/").rsplit("/", 1)[-1].strip()
    match = re.match(r"^(?P<code>[A-Z][A-Z0-9]{4,15})(?=[_＿-])", filename)
    return match.group("code") if match is not None else ""


def _with_label_snapshot_filename_identity(
    rows: list[dict[str, object]],
    *,
    attachment_name: str,
) -> list[dict[str, object]]:
    filename_code = _label_snapshot_filename_code(attachment_name)
    if not filename_code:
        return rows
    enriched: list[dict[str, object]] = []
    for raw_row in rows:
        row = dict(raw_row)
        parsed_code = str(row.get("instrument_code") or "").strip().upper()
        if parsed_code and parsed_code != filename_code:
            raise ValueError(
                "Label NAV snapshot identity conflicts between workbook and filename."
            )
        row["instrument_code"] = filename_code
        enriched.append(row)
    return enriched


def _with_optional_label_snapshot_filename_identity(
    rows: list[dict[str, object]],
    *,
    attachment_name: str,
) -> list[dict[str, object]]:
    """Fill missing identity without overriding a workbook's explicit code."""
    filename_code = _label_snapshot_filename_code(attachment_name)
    if not filename_code:
        return rows
    enriched: list[dict[str, object]] = []
    for raw_row in rows:
        row = dict(raw_row)
        if not str(row.get("instrument_code") or "").strip():
            row["instrument_code"] = filename_code
        enriched.append(row)
    return enriched


def _parse_nav_rows_from_text(raw_text: str) -> list[dict[str, object]]:
    return _parse_nav_rows_from_matrix(_matrix_from_text(raw_text))


def _parse_nav_rows_from_workbook_matrices(
    matrices: list[tuple[str, list[list[object]]]],
    *,
    parser: Callable[[list[list[object]]], list[dict[str, object]]],
) -> list[dict[str, object]]:
    """Parse every sheet, retaining provenance and collapsing exact duplicates."""
    rows_by_signature: dict[str, dict[str, object]] = {}
    for sheet_name, matrix in matrices:
        for raw_row in parser(matrix):
            row = dict(raw_row)
            signature = json.dumps(
                {
                    key: str(value)
                    for key, value in sorted(row.items())
                    if key not in {"_source_sheet", "_source_sheets"}
                },
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
            )
            existing = rows_by_signature.get(signature)
            if existing is None:
                row["_source_sheet"] = sheet_name
                row["_source_sheets"] = [sheet_name]
                rows_by_signature[signature] = row
                continue
            source_sheets = list(existing.get("_source_sheets", []))
            if sheet_name not in source_sheets:
                source_sheets.append(sheet_name)
                existing["_source_sheets"] = source_sheets
    return sorted(
        rows_by_signature.values(),
        key=lambda row: (
            str(row.get("as_of_date") or ""),
            str(row.get("instrument_code") or ""),
            str(row.get("instrument_name") or ""),
            str(row.get("nav") or ""),
            str(row.get("cash_cumulative_nav") or ""),
            str(row.get("nav_with_dividend") or ""),
            str(row.get("_source_sheet") or ""),
        ),
    )


def _parse_nav_rows_from_xlsx(file_bytes: bytes) -> list[dict[str, object]]:
    return _parse_nav_rows_from_workbook_matrices(
        _xlsx_matrices(file_bytes),
        parser=_parse_nav_rows_from_matrix,
    )


def _parse_nav_rows_from_xls(file_bytes: bytes) -> list[dict[str, object]]:
    return _parse_nav_rows_from_workbook_matrices(
        _xls_matrices(file_bytes),
        parser=_parse_nav_rows_from_matrix,
    )


_XLS_OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
_ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


def _workbook_matrices_by_content(
    file_bytes: bytes,
) -> list[tuple[str, list[list[object]]]]:
    if file_bytes.startswith(_XLS_OLE_SIGNATURE):
        return _xls_matrices(file_bytes)
    if file_bytes.startswith(_ZIP_SIGNATURES):
        return _xlsx_matrices(file_bytes)
    raise ValueError("NAV workbook content is neither XLS nor XLSX.")


def _parse_nav_rows_from_workbook_content(
    file_bytes: bytes,
    *,
    parser: Callable[[list[list[object]]], list[dict[str, object]]],
) -> list[dict[str, object]]:
    return _parse_nav_rows_from_workbook_matrices(
        _workbook_matrices_by_content(file_bytes),
        parser=parser,
    )


def _decode_nav_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("NAV text file must be UTF-8 or GB18030 encoded.")


def _parse_nav_rows_from_uploaded_file(
    *,
    file_name: str,
    file_bytes: bytes,
) -> list[dict[str, object]]:
    lower_name = file_name.lower().strip()
    if lower_name.endswith((".csv", ".tsv", ".txt")):
        return _parse_nav_rows_from_text(_decode_nav_text(file_bytes))
    if lower_name.endswith((".xlsx", ".xls")):
        return _parse_nav_rows_from_workbook_content(
            file_bytes,
            parser=_parse_nav_rows_from_matrix,
        )
    raise ValueError("Unsupported NAV file type. Use csv, tsv, txt, xlsx, or xls.")


def _parse_nav_rows_from_attachment(
    *,
    attachment_name: str,
    attachment_bytes: bytes,
    parser_profile: str,
) -> list[dict[str, object]]:
    lower_name = attachment_name.lower()
    if parser_profile == "label_nav_snapshot":
        if lower_name.endswith((".xls", ".xlsx")):
            return _with_label_snapshot_filename_identity(
                _parse_nav_rows_from_workbook_content(
                    attachment_bytes,
                    parser=_parse_nav_rows_from_label_snapshot_matrix,
                ),
                attachment_name=attachment_name,
            )
        return []

    if lower_name.endswith((".csv", ".tsv", ".txt")):
        return _parse_nav_rows_from_text(_decode_nav_text(attachment_bytes))
    if lower_name.endswith((".xlsx", ".xls")):
        matrices = _workbook_matrices_by_content(attachment_bytes)
        rows = _parse_nav_rows_from_workbook_matrices(
            matrices,
            parser=_parse_nav_rows_from_matrix,
        )
        if rows:
            return rows
        return _with_optional_label_snapshot_filename_identity(
            _parse_nav_rows_from_workbook_matrices(
                matrices,
                parser=_parse_nav_rows_from_label_snapshot_matrix,
            ),
            attachment_name=attachment_name,
        )
    return []


def _extract_attachment_text(attachment_name: str, attachment_bytes: bytes) -> str:
    lower_name = attachment_name.lower()
    if lower_name.endswith((".csv", ".tsv", ".txt")):
        return _decode_nav_text(attachment_bytes)
    if lower_name.endswith((".xlsx", ".xls")):
        return "\n".join(
            _extract_attachment_text_from_matrix(matrix)
            for _sheet_name, matrix in _workbook_matrices_by_content(
                attachment_bytes
            )
        )
    return ""


def _normalize_text_token(value: str) -> str:
    return re.sub(r"[^0-9a-zA-Z\u4e00-\u9fff]+", "", value.lower())


def _normalize_import_status(rows: list[dict[str, object]], fallback: str) -> str:
    del rows
    return fallback


_NAV_VALUE_FIELDS = frozenset({"nav", "cash_cumulative_nav", "nav_with_dividend"})
_NAV_SOURCE_EVIDENCE_FIELDS = (
    "_email_candidate_route_id",
    "_email_attachment_name",
    "_email_folder",
    "_email_uid",
    "_email_sent_at",
    "_raw_observation_id",
    "_raw_captured_at",
    "_raw_source_kind",
    "_raw_source_ref",
    "_raw_source_provider",
    "_total_return_source_field",
    "_source_sheet",
    "_source_sheets",
)
_NAV_VALUE_STATUSES = frozenset({"complete", "partial", "unavailable"})


def _nav_value_source_metadata(
    row: dict[str, object],
    value_field: str,
) -> dict[str, object]:
    existing = row.get(f"_{value_field}_source_metadata")
    if isinstance(existing, dict):
        return dict(existing)
    metadata = {
        key: row[key]
        for key in _NAV_SOURCE_EVIDENCE_FIELDS
        if row.get(key) is not None and row.get(key) != ""
    }
    field_provider = row.get(f"_{value_field}_source_provider")
    if field_provider is not None and field_provider != "":
        metadata["_raw_source_provider"] = field_provider
    return metadata


def _nav_value_status(row: dict[str, object], value_field: str) -> str | None:
    if row.get(value_field) is None or row.get(value_field) == "":
        return None
    status = str(row.get(f"_{value_field}_status") or "").strip().lower()
    if status not in _NAV_VALUE_STATUSES:
        raise ValueError(
            f'{value_field} requires an explicit source status; got "{status}".'
        )
    return status


def _stamp_nav_value_statuses(
    rows: list[dict[str, object]],
    *,
    status: str,
) -> list[dict[str, object]]:
    normalized_status = status.strip().lower()
    if normalized_status not in _NAV_VALUE_STATUSES:
        raise ValueError(f'Unsupported NAV status "{status}".')
    stamped: list[dict[str, object]] = []
    for raw_row in rows:
        row = dict(raw_row)
        for value_field in _NAV_VALUE_FIELDS:
            if row.get(value_field) is not None and row.get(value_field) != "":
                row[f"_{value_field}_status"] = normalized_status
        stamped.append(row)
    return stamped


def _merge_rows_by_date(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    merged: dict[str, dict[str, object]] = {}
    row_sources: dict[str, tuple[int, str, str, int]] = {}
    revision_counts: dict[str, int] = {}

    def source_rank(row: dict[str, object]) -> tuple[int, str, str, int]:
        timestamp = str(
            row.get("_email_sent_at") or row.get("_raw_captured_at") or ""
        )
        source_kind = str(row.get("_raw_source_kind") or "")
        if source_kind == "manual_import":
            source_class = 3
        elif row.get("_email_candidate_route_id") is not None:
            source_class = 2
        elif source_kind == "api_observation":
            source_class = 1
        else:
            source_class = 0
        if row.get("_email_candidate_route_id") is not None:
            revision_scope = f"email:{row.get('_email_folder') or ''}"
            revision_sequence = int(row.get("_email_uid") or 0)
        elif row.get("_raw_observation_id") is not None:
            revision_scope = "raw-observation"
            revision_sequence = int(row.get("_raw_observation_id") or 0)
        else:
            revision_scope = ""
            revision_sequence = 0
        return (
            source_class,
            timestamp,
            revision_scope,
            revision_sequence,
        )

    def compare_source_rank(
        left: tuple[int, str, str, int],
        right: tuple[int, str, str, int],
    ) -> int | None:
        """Compare only revisions with a real chronological ordering."""
        if left[0] != right[0]:
            return 1 if left[0] > right[0] else -1
        if left[1] and right[1] and left[1] != right[1]:
            return 1 if left[1] > right[1] else -1
        if left[2] and left[2] == right[2] and left[3] != right[3]:
            return 1 if left[3] > right[3] else -1
        if left == right:
            return 0
        return None

    def normalized_value(value: object) -> tuple[str, object]:
        if value is None or value == "":
            return ("absent", "")
        parsed = _parse_nav_decimal(value)
        if parsed is not None:
            return ("decimal", parsed)
        return ("raw", str(value).strip())

    def observation_signature(row: dict[str, object]) -> tuple[object, ...]:
        return (
            str(row.get("currency") or "").strip().upper(),
            *(normalized_value(row.get(key)) for key in sorted(_NAV_VALUE_FIELDS)),
        )

    def prepared_row(row: dict[str, object], as_of_date: str) -> dict[str, object]:
        prepared = {
            key: value
            for key, value in row.items()
            if key != "as_of_date"
            and not key.endswith("_source_metadata")
            and value is not None
            and value != ""
        }
        prepared["as_of_date"] = as_of_date
        for value_field in _NAV_VALUE_FIELDS:
            if row.get(value_field) is None or row.get(value_field) == "":
                continue
            prepared[f"_{value_field}_source_metadata"] = (
                _nav_value_source_metadata(row, value_field)
            )
        return prepared

    for row in rows:
        as_of_date = str(row.get("as_of_date") or "").strip()
        if not as_of_date:
            continue
        rank = source_rank(row)
        candidate = prepared_row(row, as_of_date)
        existing = merged.get(as_of_date)
        if existing is None:
            merged[as_of_date] = candidate
            row_sources[as_of_date] = rank
            continue

        previous_rank = row_sources[as_of_date]
        same_observation = observation_signature(existing) == observation_signature(
            candidate
        )
        comparison = compare_source_rank(rank, previous_rank)
        if not same_observation:
            if comparison in (None, 0):
                raise ValueError(
                    f"Conflicting NAV observations for {as_of_date} lack a strict "
                    "source revision order."
                )
            revision_counts[as_of_date] = revision_counts.get(as_of_date, 0) + 1
        if comparison == 1:
            merged[as_of_date] = candidate
            row_sources[as_of_date] = rank
    for as_of_date, revision_count in revision_counts.items():
        merged[as_of_date]["_email_revision_count"] = revision_count
    return [merged[key] for key in sorted(merged.keys())]


@dataclass(frozen=True)
class FundNavPublication:
    rows: list[dict[str, object]]
    projection_run: dict[str, object]
    current_fund_nav_event_ids: list[str]
    current_fund_nav_reinvestment_evidence_ids: list[str]
    action_candidates: list[dict[str, object]]
    adjustment_factors: list[dict[str, object]]
    derived_total_return_count: int
    explicit_total_return_count: int

    @property
    def has_derived_total_return(self) -> bool:
        return self.derived_total_return_count > 0


FUND_NAV_PROJECTION_METHOD_VERSION = "fund_nav_reinvestment_projection/v6"
FUND_NAV_PROJECTION_CREATED_BY = "platform_fund_nav_projection_builder"
AUTO_CASH_DISTRIBUTION_MAX_ABS_INTERVAL_RETURN = Decimal("0.50")
CASH_REPORTING_RESET_CONFIRMATION_ROWS = 3
CUMULATIVE_NAV_SEMANTICS_MIN_VOTES = 3
CUMULATIVE_NAV_SEMANTICS_DOMINANCE = 4
CUMULATIVE_NAV_SEMANTICS_COMPARISON_LAGS = (1, 5, 20)


def _quantized_factor_level(value: Decimal) -> Decimal:
    with localcontext() as context:
        context.prec = 76
        return value.quantize(
            FUND_NAV_FACTOR_DECIMAL_PLACES,
            rounding=ROUND_HALF_UP,
        )


def _fund_nav_ledger_id(prefix: str, *parts: object) -> str:
    identity = "\\0".join(str(part) for part in parts)
    digest = hashlib.sha256(identity.encode("utf-8")).hexdigest()
    return f"{prefix}-{digest}"


def _row_source_provider(
    row: dict[str, object],
    *,
    fallback: str,
) -> str:
    raw_provider = str(row.get("_raw_source_provider") or "").strip()
    if raw_provider:
        source_ref = str(row.get("_raw_source_ref") or "").strip()
        return f"{raw_provider}|{source_ref}" if source_ref else raw_provider
    route_id = row.get("_email_candidate_route_id")
    if route_id is not None:
        return f"email-route:{route_id}"
    return fallback.strip() or "platform_nav_import"


def _row_lineage_evidence(row: dict[str, object]) -> dict[str, object]:
    evidence: dict[str, object] = {}
    for evidence_key in _NAV_SOURCE_EVIDENCE_FIELDS:
        if evidence_key == "_total_return_source_field":
            continue
        value = row.get(evidence_key)
        if value is not None and value != "":
            evidence[evidence_key.removeprefix("_")] = value
    return evidence


def _reported_decimal_uncertainty(value: object) -> Decimal:
    if value is None or value == "":
        return Decimal("0")
    try:
        parsed = Decimal(str(value).strip().replace(",", ""))
    except Exception:
        return Decimal("0")
    if not parsed.is_finite():
        return Decimal("0")
    exponent = parsed.as_tuple().exponent
    quantum = Decimal(1).scaleb(exponent) if exponent < 0 else Decimal(1)
    return abs(quantum) / Decimal(2)


def _cash_delta_uncertainty(
    previous_row: dict[str, object],
    current_row: dict[str, object],
) -> Decimal:
    return max(
        NAV_SEMANTIC_ZERO_TOLERANCE,
        sum(
            (
                _reported_decimal_uncertainty(previous_row.get("nav")),
                _reported_decimal_uncertainty(
                    previous_row.get("cash_cumulative_nav")
                ),
                _reported_decimal_uncertainty(current_row.get("nav")),
                _reported_decimal_uncertainty(
                    current_row.get("cash_cumulative_nav")
                ),
            ),
            Decimal("0"),
        ),
    )


def _cash_balance_zero_tolerance(row: dict[str, object]) -> Decimal:
    return max(
        NAV_SEMANTIC_ZERO_TOLERANCE,
        _reported_decimal_uncertainty(row.get("nav"))
        + _reported_decimal_uncertainty(row.get("cash_cumulative_nav")),
    )


def _cash_reporting_basis_reset_evidence(
    *,
    complete_rows: list[tuple[date, dict[str, object], Decimal]],
    current_index: int,
    expected_cash_balance: Decimal | None,
    event_ids: list[str],
) -> dict[str, object] | None:
    """Recognize a sustained provider reset from cash-cumulative to unit NAV.

    This is deliberately narrower than a generic negative cash-balance change:
    the pre-reset balance must be stable, unit NAV may not fall at the boundary,
    and the disclosed cash balance must remain zero for three observations.
    """

    if (
        expected_cash_balance is None
        or event_ids
        or current_index < 2
        or len(complete_rows) - current_index
        < CASH_REPORTING_RESET_CONFIRMATION_ROWS
    ):
        return None
    previous_date, previous_row, previous_cash_balance = complete_rows[
        current_index - 1
    ]
    prior_date, prior_row, prior_cash_balance = complete_rows[current_index - 2]
    current_date, current_row, current_cash_balance = complete_rows[current_index]
    boundary_uncertainty = _cash_delta_uncertainty(previous_row, current_row)
    if expected_cash_balance <= boundary_uncertainty:
        return None
    if (
        abs(prior_cash_balance - previous_cash_balance)
        > _cash_delta_uncertainty(prior_row, previous_row)
    ):
        return None
    previous_unit_nav = _parse_nav_decimal(previous_row.get("nav"))
    current_unit_nav = _parse_nav_decimal(current_row.get("nav"))
    if previous_unit_nav is None or current_unit_nav is None:
        return None
    unit_move_uncertainty = (
        _reported_decimal_uncertainty(previous_row.get("nav"))
        + _reported_decimal_uncertainty(current_row.get("nav"))
    )
    if current_unit_nav + unit_move_uncertainty < previous_unit_nav:
        return None

    confirmation_rows = complete_rows[
        current_index : current_index + CASH_REPORTING_RESET_CONFIRMATION_ROWS
    ]
    if any(
        abs(cash_balance) > _cash_balance_zero_tolerance(row)
        for _row_date, row, cash_balance in confirmation_rows
    ):
        return None
    with localcontext() as context:
        context.prec = 76
        unit_return = current_unit_nav / previous_unit_nav - Decimal("1")
    return {
        "kind": "cash_cumulative_reporting_basis_reset_to_unit_nav",
        "interval_start_date": previous_date.isoformat(),
        "interval_end_date": current_date.isoformat(),
        "stable_pre_reset_start_date": prior_date.isoformat(),
        "cash_balance_before": str(previous_cash_balance),
        "cash_balance_after": str(current_cash_balance),
        "unit_nav_before": str(previous_unit_nav),
        "unit_nav_after": str(current_unit_nav),
        "unit_return_at_reset": str(_quantized_factor_level(unit_return)),
        "confirmation_dates": [
            row_date.isoformat()
            for row_date, _row, _cash_balance in confirmation_rows
        ],
    }


def _nav_ratio_uncertainty(row: dict[str, object]) -> Decimal:
    unit_nav = _parse_nav_decimal(row.get("nav"))
    cumulative_nav = _parse_nav_decimal(row.get("cash_cumulative_nav"))
    if unit_nav is None or cumulative_nav is None or unit_nav <= 0:
        return Decimal("0")
    unit_uncertainty = _reported_decimal_uncertainty(row.get("nav"))
    cumulative_uncertainty = _reported_decimal_uncertainty(
        row.get("cash_cumulative_nav")
    )
    with localcontext() as context:
        context.prec = 76
        return max(
            NAV_SEMANTIC_ZERO_TOLERANCE,
            cumulative_uncertainty / unit_nav
            + abs(cumulative_nav) * unit_uncertainty / (unit_nav * unit_nav),
        )


def _infer_cumulative_nav_semantics(
    rows: list[dict[str, object]],
) -> dict[str, object]:
    """Distinguish cash-cumulative from dividend-reinvested disclosures.

    A cash-cumulative series keeps ``cumulative - unit`` stable between cash
    distributions.  A dividend-reinvested series keeps ``cumulative / unit``
    stable instead.  We require several informative post-divergence intervals
    and a decisive vote margin; otherwise the source remains ambiguous.
    """

    complete_rows: list[dict[str, object]] = []
    for row in rows:
        unit_nav = _parse_nav_decimal(row.get("nav"))
        cumulative_nav = _parse_nav_decimal(row.get("cash_cumulative_nav"))
        if (
            unit_nav is None
            or cumulative_nav is None
            or unit_nav <= 0
            or _nav_value_status(row, "nav") != "complete"
            or _nav_value_status(row, "cash_cumulative_nav") != "complete"
        ):
            continue
        complete_rows.append(row)

    cash_cumulative_votes = 0
    reinvested_votes = 0
    informative_intervals = 0
    comparison_pairs = (
        (complete_rows[current_index - lag], complete_rows[current_index])
        for current_index in range(1, len(complete_rows))
        for lag in CUMULATIVE_NAV_SEMANTICS_COMPARISON_LAGS
        if current_index >= lag
    )
    for previous_row, current_row in comparison_pairs:
        previous_unit = _parse_nav_decimal(previous_row.get("nav"))
        current_unit = _parse_nav_decimal(current_row.get("nav"))
        previous_cumulative = _parse_nav_decimal(
            previous_row.get("cash_cumulative_nav")
        )
        current_cumulative = _parse_nav_decimal(
            current_row.get("cash_cumulative_nav")
        )
        assert previous_unit is not None and current_unit is not None
        assert previous_cumulative is not None and current_cumulative is not None
        unit_move_uncertainty = (
            _reported_decimal_uncertainty(previous_row.get("nav"))
            + _reported_decimal_uncertainty(current_row.get("nav"))
        )
        if abs(current_unit - previous_unit) <= unit_move_uncertainty:
            continue
        previous_balance = previous_cumulative - previous_unit
        current_balance = current_cumulative - current_unit
        balance_uncertainty = _cash_delta_uncertainty(
            previous_row,
            current_row,
        )
        if (
            abs(previous_balance) <= balance_uncertainty
            and abs(current_balance) <= balance_uncertainty
        ):
            # Before the first distribution both semantics collapse to unit NAV.
            continue
        with localcontext() as context:
            context.prec = 76
            previous_ratio = previous_cumulative / previous_unit
            current_ratio = current_cumulative / current_unit
        difference_stable = (
            abs(current_balance - previous_balance) <= balance_uncertainty
        )
        ratio_stable = abs(current_ratio - previous_ratio) <= (
            _nav_ratio_uncertainty(previous_row)
            + _nav_ratio_uncertainty(current_row)
        )
        if difference_stable == ratio_stable:
            # A distribution boundary can move both measures; a flat unit NAV
            # can leave both stable. Neither interval is discriminating.
            continue
        informative_intervals += 1
        if difference_stable:
            cash_cumulative_votes += 1
        else:
            reinvested_votes += 1

    kind = "ambiguous"
    if cash_cumulative_votes >= max(
        CUMULATIVE_NAV_SEMANTICS_MIN_VOTES,
        reinvested_votes * CUMULATIVE_NAV_SEMANTICS_DOMINANCE,
    ):
        kind = "cash_non_reinvested"
    elif reinvested_votes >= max(
        CUMULATIVE_NAV_SEMANTICS_MIN_VOTES,
        cash_cumulative_votes * CUMULATIVE_NAV_SEMANTICS_DOMINANCE,
    ):
        kind = "dividend_reinvested"
    return {
        "kind": kind,
        "cash_cumulative_votes": cash_cumulative_votes,
        "dividend_reinvested_votes": reinvested_votes,
        "informative_intervals": informative_intervals,
        "complete_observation_count": len(complete_rows),
        "comparison_lags": list(CUMULATIVE_NAV_SEMANTICS_COMPARISON_LAGS),
    }


def _current_fund_nav_heads(
    instrument: dict[str, object] | None,
) -> tuple[list[dict[str, object]], dict[str, dict[str, object]]]:
    """Return exact current action/evidence heads without legacy status inference."""

    if not isinstance(instrument, dict):
        return [], {}

    events: list[dict[str, object]] = []
    seen_event_ids: set[str] = set()
    for index, raw_event in enumerate(
        list(instrument.get("fund_nav_events", [])),
        start=1,
    ):
        if not isinstance(raw_event, dict):
            raise ValueError(f"Current fund NAV event {index} must be an object.")
        event = dict(raw_event)
        event_id = str(event.get("fund_nav_event_id") or "").strip()
        effective_date = _parse_nav_date(event.get("effective_date"))
        event_type = str(event.get("event_type") or "").strip()
        if not event_id or effective_date is None:
            raise ValueError(
                f"Current fund NAV event {index} requires an id and effective_date."
            )
        if event_id in seen_event_ids:
            raise ValueError(f'Duplicate current fund NAV event id "{event_id}".')
        if event_type not in {"cash_distribution", "unit_split"}:
            raise ValueError(
                f'Current fund NAV event "{event_id}" has unsupported type '
                f'"{event_type}".'
            )
        seen_event_ids.add(event_id)
        event["fund_nav_event_id"] = event_id
        event["effective_date"] = effective_date.isoformat()
        events.append(event)

    events_by_date: dict[date, list[dict[str, object]]] = {}
    for event in events:
        event_date = _parse_nav_date(event["effective_date"])
        assert event_date is not None
        events_by_date.setdefault(event_date, []).append(event)
    for event_date, same_day_events in events_by_date.items():
        if len(same_day_events) == 1:
            same_day_events.sort(
                key=lambda item: str(item["fund_nav_event_id"]),
            )
            continue
        sequence_orders = [
            event.get("sequence_order") for event in same_day_events
        ]
        if (
            any(
                not isinstance(sequence_order, int) or sequence_order < 1
                for sequence_order in sequence_orders
            )
            or len(set(sequence_orders)) != len(sequence_orders)
        ):
            # Keep a deterministic diagnostic order, but do not invent the
            # non-commutative business order. Projection segments crossing the
            # date will be marked unavailable below.
            same_day_events.sort(
                key=lambda item: str(item["fund_nav_event_id"]),
            )
            for event in same_day_events:
                event["_sequence_ambiguous"] = True
            continue
        same_day_events.sort(key=lambda item: int(item["sequence_order"]))
    ordered_events = [
        event
        for event_date in sorted(events_by_date)
        for event in events_by_date[event_date]
    ]

    evidence_by_event_id: dict[str, dict[str, object]] = {}
    seen_evidence_ids: set[str] = set()
    for index, raw_evidence in enumerate(
        list(instrument.get("fund_nav_reinvestment_evidence", [])),
        start=1,
    ):
        if not isinstance(raw_evidence, dict):
            raise ValueError(
                f"Current fund NAV reinvestment evidence {index} must be an object."
            )
        evidence = dict(raw_evidence)
        evidence_id = str(
            evidence.get("fund_nav_reinvestment_evidence_id") or ""
        ).strip()
        event_id = str(evidence.get("fund_nav_event_id") or "").strip()
        if not evidence_id or not event_id:
            raise ValueError(
                "Current fund NAV reinvestment evidence requires evidence and event ids."
            )
        if evidence_id in seen_evidence_ids:
            raise ValueError(
                f'Duplicate current reinvestment evidence id "{evidence_id}".'
            )
        if event_id in evidence_by_event_id:
            raise ValueError(
                f'Current fund NAV event "{event_id}" has multiple evidence heads.'
            )
        if event_id not in seen_event_ids:
            raise ValueError(
                f'Current reinvestment evidence "{evidence_id}" references a non-current event.'
            )
        seen_evidence_ids.add(evidence_id)
        evidence["fund_nav_reinvestment_evidence_id"] = evidence_id
        evidence["fund_nav_event_id"] = event_id
        evidence_by_event_id[event_id] = evidence
    return ordered_events, evidence_by_event_id


def _fund_nav_events_between(
    events: list[dict[str, object]],
    *,
    after_date: date,
    through_date: date,
) -> list[dict[str, object]]:
    return [
        event
        for event in events
        if after_date
        < (_parse_nav_date(event.get("effective_date")) or date.min)
        <= through_date
    ]


def _fund_nav_source_observation_fingerprint(
    rows: list[dict[str, object]],
) -> str:
    """Hash only ordered durable source facts, never derived projection output."""

    observations: list[dict[str, object]] = []
    for row in sorted(rows, key=lambda item: str(item.get("as_of_date") or "")):
        observation: dict[str, object] = {
            "as_of_date": str(row.get("as_of_date") or "").strip(),
            "currency": str(row.get("currency") or "").strip().upper(),
        }
        for value_field in sorted(_NAV_VALUE_FIELDS):
            value = row.get(value_field)
            if value is None or value == "":
                continue
            observation[value_field] = str(value).strip()
            observation[f"{value_field}_status"] = _nav_value_status(
                row,
                value_field,
            )
            observation[f"{value_field}_source"] = _nav_value_source_metadata(
                row,
                value_field,
            )
        observations.append(observation)
    payload = {
        "contract": "fund_nav_source_observations/v1",
        "observations": observations,
    }
    return hashlib.sha256(
        json.dumps(
            payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()


def _cash_balance_action_candidate(
    *,
    instrument_id: str,
    previous_row: dict[str, object],
    current_row: dict[str, object],
    cash_balance_before: Decimal,
    cash_balance_after: Decimal,
    expected_cash_balance: Decimal | None,
    unit_nav: Decimal,
    cash_cumulative_nav: Decimal,
    unit_provider: str,
    cash_provider: str,
    unit_evidence: dict[str, object],
    cash_evidence: dict[str, object],
    confirmed_event_ids: list[str] | None = None,
) -> dict[str, object]:
    """Describe an observed interval anomaly without inventing a fund event.

    A cash-cumulative disclosure does not reveal the exact event date or the
    reinvestment NAV.  Those unknowns deliberately do not enter the canonical
    event/factor ledger.
    """

    interval_start = _parse_nav_date(previous_row.get("as_of_date"))
    interval_end = _parse_nav_date(current_row.get("as_of_date"))
    if interval_start is None or interval_end is None or interval_start >= interval_end:
        raise ValueError("Fund NAV action candidate requires an ordered date interval.")
    observed_delta = cash_balance_after - cash_balance_before
    uncertainty = _cash_delta_uncertainty(previous_row, current_row)
    candidate_type = (
        "cash_distribution_signal"
        if observed_delta > 0
        else "cash_balance_discontinuity"
    )
    revision_payload = {
        "interval_start_date": interval_start.isoformat(),
        "interval_end_date": interval_end.isoformat(),
        "cash_balance_before": str(cash_balance_before),
        "cash_balance_after": str(cash_balance_after),
        "expected_cash_balance": (
            str(expected_cash_balance) if expected_cash_balance is not None else None
        ),
        "measurement_uncertainty": str(uncertainty),
        "unit_nav": str(unit_nav),
        "cash_cumulative_nav": str(cash_cumulative_nav),
        "unit_provider": unit_provider,
        "cash_provider": cash_provider,
        "unit_evidence": unit_evidence,
        "cash_evidence": cash_evidence,
        "confirmed_event_ids": confirmed_event_ids or [],
    }
    source_revision = hashlib.sha256(
        json.dumps(
            revision_payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()
    return {
        "fund_nav_action_candidate_id": _fund_nav_ledger_id(
            "fund-nav-action-candidate",
            instrument_id,
            candidate_type,
            interval_start,
            interval_end,
            source_revision,
        ),
        "candidate_type": candidate_type,
        "interval_start_date": interval_start.isoformat(),
        "interval_end_date": interval_end.isoformat(),
        "observed_cash_balance_before": str(cash_balance_before),
        "observed_cash_balance_after": str(cash_balance_after),
        "observed_cash_delta": str(observed_delta),
        "expected_cash_balance": (
            str(expected_cash_balance) if expected_cash_balance is not None else None
        ),
        "measurement_uncertainty": str(uncertainty),
        "status": "open",
        "source_provider": cash_provider,
        "source_revision": source_revision,
        "source_evidence": revision_payload,
    }


def _cash_disclosure_intervals(
    *,
    instrument_id: str,
    rows: list[dict[str, object]],
    events: list[dict[str, object]],
    source_provider: str,
) -> tuple[dict[date, dict[str, object]], list[dict[str, object]]]:
    """Validate disclosed cash balances without deriving reinvestment returns."""

    complete_rows: list[tuple[date, dict[str, object], Decimal]] = []
    for raw_row in rows:
        row_date = _parse_nav_date(raw_row.get("as_of_date"))
        unit_nav = _parse_nav_decimal(raw_row.get("nav"))
        cash_cumulative_nav = _parse_nav_decimal(
            raw_row.get("cash_cumulative_nav")
        )
        if (
            row_date is None
            or unit_nav is None
            or cash_cumulative_nav is None
            or _nav_value_status(raw_row, "nav") != "complete"
            or _nav_value_status(raw_row, "cash_cumulative_nav") != "complete"
        ):
            continue
        complete_rows.append(
            (row_date, raw_row, cash_cumulative_nav - unit_nav)
        )

    intervals: dict[date, dict[str, object]] = {}
    candidates: list[dict[str, object]] = []
    if not complete_rows:
        return intervals, candidates
    first_date, _first_row, first_balance = complete_rows[0]
    intervals[first_date] = {
        "kind": "baseline",
        "valid": True,
        "cash_balance": first_balance,
        "event_ids": [],
    }

    previous_date, previous_row, previous_balance = complete_rows[0]
    for current_index, (
        current_date,
        current_row,
        current_balance,
    ) in enumerate(complete_rows[1:], start=1):
        interval_events = _fund_nav_events_between(
            events,
            after_date=previous_date,
            through_date=current_date,
        )
        interval_event_ids = [
            str(event["fund_nav_event_id"]) for event in interval_events
        ]
        interval_valid = not any(
            bool(event.get("_sequence_ambiguous")) for event in interval_events
        )
        expected_cash_balance = previous_balance
        if interval_valid:
            for event in interval_events:
                event_type = str(event.get("event_type") or "")
                if event_type == "cash_distribution":
                    cash_per_unit = _parse_nav_decimal(
                        event.get("cash_per_unit")
                    )
                    if cash_per_unit is None:
                        interval_valid = False
                        break
                    expected_cash_balance += cash_per_unit
                elif event_type == "unit_split":
                    unit_ratio = _parse_nav_decimal(event.get("unit_ratio"))
                    if unit_ratio is None:
                        interval_valid = False
                        break
                    with localcontext() as context:
                        context.prec = 76
                        expected_cash_balance /= unit_ratio
                else:  # pragma: no cover - current-head validation rejects this
                    interval_valid = False
                    break

        tolerance = _cash_delta_uncertainty(previous_row, current_row)
        reconciled = interval_valid and (
            abs(current_balance - expected_cash_balance) <= tolerance
        )
        reporting_basis_reset = (
            _cash_reporting_basis_reset_evidence(
                complete_rows=complete_rows,
                current_index=current_index,
                expected_cash_balance=(
                    expected_cash_balance if interval_valid else None
                ),
                event_ids=interval_event_ids,
            )
            if not reconciled
            else None
        )
        intervals[current_date] = {
            "kind": "interval",
            "valid": reconciled,
            "cash_balance": current_balance,
            "expected_cash_balance": (
                expected_cash_balance if interval_valid else None
            ),
            "observed_cash_delta": current_balance - previous_balance,
            "measurement_uncertainty": tolerance,
            "previous_unit_nav": _parse_nav_decimal(previous_row.get("nav")),
            "current_unit_nav": _parse_nav_decimal(current_row.get("nav")),
            "event_ids": interval_event_ids,
            "interval_start_date": previous_date.isoformat(),
            "reporting_basis_reset": reporting_basis_reset,
        }
        if not reconciled and reporting_basis_reset is None:
            unit_nav = _parse_nav_decimal(current_row.get("nav"))
            cash_cumulative_nav = _parse_nav_decimal(
                current_row.get("cash_cumulative_nav")
            )
            assert unit_nav is not None and cash_cumulative_nav is not None
            unit_source = _nav_value_source_metadata(current_row, "nav")
            cash_source = _nav_value_source_metadata(
                current_row,
                "cash_cumulative_nav",
            )
            candidates.append(
                _cash_balance_action_candidate(
                    instrument_id=instrument_id,
                    previous_row=previous_row,
                    current_row=current_row,
                    cash_balance_before=previous_balance,
                    cash_balance_after=current_balance,
                    expected_cash_balance=(
                        expected_cash_balance if interval_valid else None
                    ),
                    unit_nav=unit_nav,
                    cash_cumulative_nav=cash_cumulative_nav,
                    unit_provider=_row_source_provider(
                        unit_source,
                        fallback=source_provider,
                    ),
                    cash_provider=_row_source_provider(
                        cash_source,
                        fallback=source_provider,
                    ),
                    unit_evidence=_row_lineage_evidence(unit_source),
                    cash_evidence=_row_lineage_evidence(cash_source),
                    confirmed_event_ids=interval_event_ids,
                )
            )
        previous_date = current_date
        previous_row = current_row
        previous_balance = current_balance
    return intervals, candidates


def _auto_cash_distribution_terms(
    *,
    interval: dict[str, object] | None,
    current_unit_nav: Decimal,
) -> tuple[Decimal, Decimal, Decimal] | None:
    """Resolve a positive disclosed cash jump when both interval endpoints exist.

    The provider's cash-cumulative NAV reveals the per-unit distribution.  The
    first complete observation after that disclosure change supplies the
    reinvestment NAV used for an endpoint total-return calculation.  Negative
    jumps, intervals containing separately confirmed events, and implausible
    implied returns remain unresolved rather than being guessed.
    """

    if interval is None or bool(interval.get("valid")):
        return None
    if list(interval.get("event_ids") or []):
        return None
    current_balance = interval.get("cash_balance")
    expected_balance = interval.get("expected_cash_balance")
    uncertainty = interval.get("measurement_uncertainty")
    previous_unit_nav = interval.get("previous_unit_nav")
    if not all(
        isinstance(value, Decimal)
        for value in (
            current_balance,
            expected_balance,
            uncertainty,
            previous_unit_nav,
        )
    ):
        return None
    assert isinstance(current_balance, Decimal)
    assert isinstance(expected_balance, Decimal)
    assert isinstance(uncertainty, Decimal)
    assert isinstance(previous_unit_nav, Decimal)
    if current_unit_nav <= 0 or previous_unit_nav <= 0:
        return None
    cash_per_unit = current_balance - expected_balance
    if cash_per_unit <= uncertainty:
        return None
    with localcontext() as context:
        context.prec = 76
        implied_interval_return = (
            (current_unit_nav + cash_per_unit) / previous_unit_nav
        ) - Decimal("1")
        if (
            abs(implied_interval_return)
            > AUTO_CASH_DISTRIBUTION_MAX_ABS_INTERVAL_RETURN
        ):
            return None
        multiplier = _quantized_factor_level(
            Decimal("1") + cash_per_unit / current_unit_nav
        )
    return cash_per_unit, multiplier, implied_interval_return


def _build_fund_nav_publication(
    *,
    instrument_id: str,
    rows: list[dict[str, object]],
    source_provider: str = "platform_nav_import",
    instrument: dict[str, object] | None = None,
) -> FundNavPublication:
    """Build a fail-closed, revision-aware canonical fund NAV projection.

    Unit NAV is the only directly publishable price observation. A provider's
    cash-cumulative disclosure remains raw evidence. Total return is published
    only when it is either explicit provider evidence or unit NAV multiplied by
    a traceable reinvestment factor. A broken segment stays NA until an explicit
    provider total-return observation creates a new auditable re-anchor.
    """

    if instrument is None or "fund_nav_events" not in instrument:
        instrument = get_instrument(instrument_id)
    if instrument is None:
        raise ValueError(f'Instrument "{instrument_id}" no longer exists.')

    ordered_rows = sorted(
        (dict(row) for row in rows),
        key=lambda item: str(item.get("as_of_date") or ""),
    )
    source_fingerprint = _fund_nav_source_observation_fingerprint(ordered_rows)
    cumulative_nav_semantics = _infer_cumulative_nav_semantics(ordered_rows)
    events, evidence_by_event_id = _current_fund_nav_heads(instrument)
    current_event_ids = sorted(
        str(event["fund_nav_event_id"]) for event in events
    )
    current_evidence_ids = sorted(
        str(evidence["fund_nav_reinvestment_evidence_id"])
        for evidence in evidence_by_event_id.values()
    )
    if cumulative_nav_semantics["kind"] == "dividend_reinvested":
        # Some managers label a reinvested total-return series as “累计净值”.
        # Interpreting its balance changes as cash distributions would apply
        # the same dividend twice.
        cash_intervals: dict[date, dict[str, object]] = {}
        action_candidates: list[dict[str, object]] = []
    else:
        cash_intervals, action_candidates = _cash_disclosure_intervals(
            instrument_id=instrument_id,
            rows=ordered_rows,
            events=events,
            source_provider=source_provider,
        )

    canonical_rows: list[dict[str, object]] = []
    row_facts: list[dict[str, object]] = []
    seen_source_dates: set[date] = set()
    fallback_currency = str(instrument.get("currency") or "").strip().upper()
    for row_index, row in enumerate(ordered_rows, start=1):
        row_date = _parse_nav_date(row.get("as_of_date"))
        if row_date is None:
            raise ValueError(f"Fund NAV source row {row_index} has an invalid date.")
        if row_date in seen_source_dates:
            raise ValueError(
                f"Fund NAV source rows contain duplicate date {row_date.isoformat()}; "
                "durable revisions must be merged before projection."
            )
        seen_source_dates.add(row_date)
        currency = str(row.get("currency") or fallback_currency).strip().upper()
        unit_nav = _parse_nav_decimal(row.get("nav"))
        cash_cumulative_nav = _parse_nav_decimal(row.get("cash_cumulative_nav"))
        explicit_total_return = _parse_nav_decimal(row.get("nav_with_dividend"))
        unit_status = _nav_value_status(row, "nav")
        cash_status = _nav_value_status(row, "cash_cumulative_nav")
        total_return_status = _nav_value_status(row, "nav_with_dividend")
        if unit_nav is not None and not currency:
            raise ValueError(
                f"Fund NAV source row {row_index} has no canonical currency."
            )

        unit_source = _nav_value_source_metadata(row, "nav")
        cash_source = _nav_value_source_metadata(row, "cash_cumulative_nav")
        total_source = _nav_value_source_metadata(row, "nav_with_dividend")
        total_source_field = str(
            row.get("_total_return_source_field")
            or "provider_reinvested_nav"
        )
        if (
            explicit_total_return is None
            and cash_cumulative_nav is not None
            and cumulative_nav_semantics["kind"] == "dividend_reinvested"
        ):
            explicit_total_return = cash_cumulative_nav
            total_return_status = cash_status
            total_source = cash_source
            total_source_field = "cumulative_nav_dividend_reinvested"
        unit_provider = _row_source_provider(unit_source, fallback=source_provider)
        cash_provider = _row_source_provider(cash_source, fallback=source_provider)
        total_provider = _row_source_provider(total_source, fallback=source_provider)
        unit_evidence = _row_lineage_evidence(unit_source)
        cash_evidence = _row_lineage_evidence(cash_source)
        total_evidence = _row_lineage_evidence(total_source)
        if total_source_field == "cumulative_nav_dividend_reinvested":
            total_evidence = {
                **total_evidence,
                "cumulative_nav_semantics": cumulative_nav_semantics,
            }

        canonical: dict[str, object] | None = None
        if unit_nav is not None:
            canonical = {
                "as_of_date": row_date.isoformat(),
                "currency": currency,
                "nav": str(row.get("nav")).strip(),
                "nav_status": unit_status,
                "nav_source_provider": unit_provider,
                "nav_lineage": {
                    "kind": "provider_explicit",
                    "evidence": {
                        "source_field": "unit_nav",
                        **unit_evidence,
                    },
                },
            }
            canonical_rows.append(canonical)
        row_facts.append(
            {
                "raw": row,
                "date": row_date,
                "canonical": canonical,
                "unit_nav": unit_nav,
                "cash_cumulative_nav": cash_cumulative_nav,
                "explicit_total_return": explicit_total_return,
                "unit_status": unit_status,
                "cash_status": cash_status,
                "total_status": total_return_status,
                "unit_provider": unit_provider,
                "cash_provider": cash_provider,
                "total_provider": total_provider,
                "total_source_field": total_source_field,
                "unit_evidence": unit_evidence,
                "cash_evidence": cash_evidence,
                "total_evidence": total_evidence,
            }
        )

    adjustment_factors: list[dict[str, object]] = []
    explicit_count = 0
    derived_count = 0
    break_reasons: list[dict[str, object]] = []
    auto_cash_distributions: list[dict[str, object]] = []
    auto_cash_reporting_basis_resets: list[dict[str, object]] = []
    segment_live = False
    segment_anchor_date: date | None = None
    segment_factor_key: str | None = None
    segment_factor_level: Decimal | None = None
    segment_event_cutoff: date | None = None
    segment_origin: str | None = None
    segment_broken = False
    provider_normalization_scale: Decimal | None = None
    current_segment_start_date: date | None = None
    first_complete_cash_date = next(
        (
            fact["date"]
            for fact in row_facts
            if fact["unit_nav"] is not None
            and fact["cash_cumulative_nav"] is not None
            and fact["unit_status"] == "complete"
            and fact["cash_status"] == "complete"
        ),
        None,
    )

    for fact in row_facts:
        row_date = fact["date"]
        assert isinstance(row_date, date)
        canonical = fact["canonical"]
        unit_nav = fact["unit_nav"]
        cash_cumulative_nav = fact["cash_cumulative_nav"]
        explicit_total_return = fact["explicit_total_return"]
        unit_complete = unit_nav is not None and fact["unit_status"] == "complete"
        cash_complete = (
            cash_cumulative_nav is not None and fact["cash_status"] == "complete"
        )
        explicit_complete = (
            explicit_total_return is not None
            and fact["total_status"] == "complete"
            and unit_complete
            and canonical is not None
        )

        if explicit_complete:
            assert isinstance(unit_nav, Decimal)
            assert isinstance(explicit_total_return, Decimal)
            with localcontext() as context:
                context.prec = 76
                raw_provider_factor = _quantized_factor_level(
                    explicit_total_return / unit_nav
                )
            prior_complete_total_rows = [
                row
                for row in canonical_rows
                if row.get("nav_with_dividend") is not None
                and row.get("nav_with_dividend_status") == "complete"
            ]
            bridge_factor_level: Decimal | None = None
            bridge_cash_distribution: tuple[Decimal, Decimal, Decimal] | None = None
            if (
                prior_complete_total_rows
                and segment_live
                and not segment_broken
                and segment_origin
                in {"window_normalized", "cash_cumulative_derived"}
                and segment_factor_level is not None
                and segment_event_cutoff is not None
            ):
                bridge_interval = cash_intervals.get(row_date)
                bridge_events = _fund_nav_events_between(
                    events,
                    after_date=segment_event_cutoff,
                    through_date=row_date,
                )
                if (
                    bridge_interval is not None
                    and bool(bridge_interval.get("valid"))
                    and not bridge_events
                ):
                    bridge_factor_level = segment_factor_level
                elif (
                    cash_complete
                    and cumulative_nav_semantics["kind"]
                    == "cash_non_reinvested"
                ):
                    bridge_cash_distribution = _auto_cash_distribution_terms(
                        interval=bridge_interval,
                        current_unit_nav=unit_nav,
                    )
                    if bridge_cash_distribution is not None:
                        assert bridge_interval is not None
                        cash_per_unit, multiplier, implied_return = (
                            bridge_cash_distribution
                        )
                        bridge_factor_level = _quantized_factor_level(
                            segment_factor_level * multiplier
                        )
                        auto_cash_distributions.append(
                            {
                                "interval_start_date": str(
                                    bridge_interval.get("interval_start_date")
                                ),
                                "interval_end_date": row_date.isoformat(),
                                "cash_per_unit": str(cash_per_unit),
                                "reinvestment_nav": str(unit_nav),
                                "implied_interval_return": str(implied_return),
                                "resolved_by": "provider_total_return_overlap",
                            }
                        )
                        action_candidates[:] = [
                            candidate
                            for candidate in action_candidates
                            if not (
                                candidate.get("interval_start_date")
                                == bridge_interval.get("interval_start_date")
                                and candidate.get("interval_end_date")
                                == row_date.isoformat()
                            )
                        ]
                if bridge_factor_level is not None:
                    provider_normalization_scale = _quantized_factor_level(
                        bridge_factor_level / raw_provider_factor
                    )
            starts_unverified_segment = bool(prior_complete_total_rows) and (
                segment_broken
                or (
                    segment_origin
                    in {"window_normalized", "cash_cumulative_derived"}
                    and bridge_factor_level is None
                )
            )
            if starts_unverified_segment:
                for prior_row in prior_complete_total_rows:
                    # The persistence contract intentionally has no partial
                    # total-return quote: an unprovable point is represented by
                    # absence.  The historical projection run preserves the
                    # superseded derivation; the new run records the segment
                    # break below and publishes only the verified scale.
                    for field_name in (
                        "nav_with_dividend",
                        "nav_with_dividend_status",
                        "nav_with_dividend_source_provider",
                        "nav_with_dividend_lineage",
                    ):
                        prior_row.pop(field_name, None)
                adjustment_factors.clear()
                explicit_count = 0
                derived_count = 0
                auto_cash_distributions.clear()
                provider_normalization_scale = None
                break_reasons.append(
                    {
                        "as_of_date": row_date.isoformat(),
                        "reason": "provider_reanchor_without_verified_overlap",
                        "excluded_complete_points": len(prior_complete_total_rows),
                    }
                )
                current_segment_start_date = row_date
            if provider_normalization_scale is None:
                provider_normalization_scale = Decimal("1")
            factor_level = _quantized_factor_level(
                raw_provider_factor * provider_normalization_scale
            )
            factor_key = f"provider:{row_date.isoformat()}"
            total_evidence = dict(fact["total_evidence"])
            source_field = str(fact["total_source_field"])
            factor_evidence: dict[str, object] = {
                "source_field": source_field,
                "unit_nav": str(unit_nav),
                "provider_total_return_nav": str(explicit_total_return),
                "unit_nav_source_provider": str(fact["unit_provider"]),
                "total_return_source_provider": str(fact["total_provider"]),
                **total_evidence,
            }
            if provider_normalization_scale != Decimal("1"):
                factor_evidence.update(
                    {
                        "canonical_normalization_scale": str(
                            provider_normalization_scale
                        ),
                        "normalization_reason": (
                            "preserve_verified_pre_overlap_return_history"
                        ),
                        "overlap_provider_factor": str(raw_provider_factor),
                    }
                )
            adjustment_factors.append(
                {
                    "factor_logical_key": factor_key,
                    "as_of_date": row_date.isoformat(),
                    "factor_level": str(factor_level),
                    "factor_kind": "provider_implied",
                    "evidence_kind": "provider_total_return",
                    "method_version": FUND_NAV_PROJECTION_METHOD_VERSION,
                    "anchor_date": row_date.isoformat(),
                    "source_provider": str(fact["total_provider"]),
                    "evidence": factor_evidence,
                }
            )
            canonical_total = _format_total_return_nav_decimal(
                unit_nav * factor_level
            )
            canonical["nav_with_dividend"] = canonical_total
            canonical["nav_with_dividend_status"] = "complete"
            canonical["nav_with_dividend_source_provider"] = str(
                fact["total_provider"]
            )
            canonical["nav_with_dividend_lineage"] = {
                "kind": "provider_explicit",
                "evidence": {
                    "factor_logical_key": factor_key,
                    "source_field": source_field,
                    "provider_reported_total_return_nav": str(
                        explicit_total_return
                    ),
                    **(
                        {
                            "canonical_normalization_scale": str(
                                provider_normalization_scale
                            )
                        }
                        if provider_normalization_scale != Decimal("1")
                        else {}
                    ),
                    **total_evidence,
                },
            }
            explicit_count += 1
            segment_live = cash_complete
            segment_anchor_date = row_date
            segment_factor_key = factor_key
            segment_factor_level = factor_level
            segment_event_cutoff = row_date
            segment_origin = "provider_explicit"
            segment_broken = False
            if current_segment_start_date is None:
                current_segment_start_date = row_date
            continue

        if not unit_complete or not cash_complete or canonical is None:
            continue
        assert isinstance(unit_nav, Decimal)
        assert isinstance(cash_cumulative_nav, Decimal)
        cash_balance = cash_cumulative_nav - unit_nav

        if segment_anchor_date is None:
            is_first_complete_cash = row_date == first_complete_cash_date
            if is_first_complete_cash:
                factor_key = f"window:{row_date.isoformat()}"
                factor_level = Decimal("1")
                adjustment_factors.append(
                    {
                        "factor_logical_key": factor_key,
                        "as_of_date": row_date.isoformat(),
                        "factor_level": "1",
                        "factor_kind": "event_derived",
                        "evidence_kind": "window_normalized_anchor",
                        "method_version": FUND_NAV_PROJECTION_METHOD_VERSION,
                        "anchor_date": row_date.isoformat(),
                        "source_provider": "derived:cash_cumulative_nav",
                        "evidence": {
                            "normalization_scope": "observed_return_window",
                            "starting_cash_balance": str(cash_balance),
                            "calculation": "unit_nav_at_anchor=1.0_factor_level",
                        },
                    }
                )
                segment_live = True
                segment_anchor_date = row_date
                segment_factor_key = factor_key
                segment_factor_level = factor_level
                segment_event_cutoff = row_date
                segment_origin = "window_normalized"
                segment_broken = False
                current_segment_start_date = row_date
            else:
                continue

        if not segment_live:
            continue
        assert segment_anchor_date is not None
        assert segment_factor_key is not None
        assert segment_factor_level is not None
        assert segment_event_cutoff is not None

        interval_events = _fund_nav_events_between(
            events,
            after_date=segment_event_cutoff,
            through_date=row_date,
        )
        interval = cash_intervals.get(row_date)
        auto_cash_factor_applied = False
        if row_date != segment_anchor_date and (
            interval is None or not bool(interval.get("valid"))
        ):
            ambiguous_events = [
                event for event in interval_events if event.get("_sequence_ambiguous")
            ]
            auto_cash_terms = (
                None
                if (
                    ambiguous_events
                    or cumulative_nav_semantics["kind"]
                    != "cash_non_reinvested"
                )
                else _auto_cash_distribution_terms(
                    interval=interval,
                    current_unit_nav=unit_nav,
                )
            )
            reporting_basis_reset = (
                dict(interval.get("reporting_basis_reset") or {})
                if interval is not None
                else {}
            )
            if reporting_basis_reset:
                auto_cash_reporting_basis_resets.append(
                    reporting_basis_reset
                )
                segment_event_cutoff = row_date
                interval_events = []
                auto_cash_factor_applied = True
            elif auto_cash_terms is not None:
                assert interval is not None
                cash_per_unit, multiplier, implied_return = auto_cash_terms
                previous_factor_key = segment_factor_key
                previous_factor_level = segment_factor_level
                next_level = _quantized_factor_level(
                    previous_factor_level * multiplier
                )
                next_key = f"cash-disclosure:{row_date.isoformat()}"
                adjustment_factors.append(
                    {
                        "factor_logical_key": next_key,
                        "as_of_date": row_date.isoformat(),
                        "factor_level": str(next_level),
                        "factor_kind": "provider_implied",
                        "evidence_kind": "provider_cash_cumulative",
                        "method_version": FUND_NAV_PROJECTION_METHOD_VERSION,
                        "anchor_date": row_date.isoformat(),
                        "source_provider": str(fact["cash_provider"]),
                        "evidence": {
                            "source_field": "cash_cumulative_nav",
                            "interval_start_date": str(
                                interval.get("interval_start_date")
                            ),
                            "interval_end_date": row_date.isoformat(),
                            "cash_balance_before": str(
                                interval.get("expected_cash_balance")
                            ),
                            "cash_balance_after": str(
                                interval.get("cash_balance")
                            ),
                            "cash_per_unit": str(cash_per_unit),
                            "reinvestment_nav": str(unit_nav),
                            "reinvestment_nav_semantics": (
                                "first_complete_observation_after_"
                                "cash_disclosure_change"
                            ),
                            "implied_interval_return": str(implied_return),
                            "previous_factor_logical_key": previous_factor_key,
                            "previous_factor_level": str(previous_factor_level),
                            "calculation": (
                                "previous_factor*(1+cash_per_unit/"
                                "current_unit_nav)"
                            ),
                            **dict(fact["cash_evidence"]),
                        },
                    }
                )
                auto_cash_distributions.append(
                    {
                        "interval_start_date": str(
                            interval.get("interval_start_date")
                        ),
                        "interval_end_date": row_date.isoformat(),
                        "cash_per_unit": str(cash_per_unit),
                        "reinvestment_nav": str(unit_nav),
                        "implied_interval_return": str(implied_return),
                        "resolved_by": "cash_cumulative_disclosure",
                    }
                )
                action_candidates[:] = [
                    candidate
                    for candidate in action_candidates
                    if not (
                        candidate.get("interval_start_date")
                        == interval.get("interval_start_date")
                        and candidate.get("interval_end_date")
                        == row_date.isoformat()
                    )
                ]
                segment_factor_level = next_level
                segment_factor_key = next_key
                segment_anchor_date = row_date
                segment_event_cutoff = row_date
                segment_origin = "cash_cumulative_derived"
                segment_broken = False
                interval_events = []
                auto_cash_factor_applied = True
            else:
                segment_live = False
                segment_broken = True
            if ambiguous_events and not auto_cash_factor_applied:
                break_reasons.extend(
                    {
                        "as_of_date": str(event["effective_date"]),
                        "reason": "same_day_action_sequence_unverified",
                        "fund_nav_event_id": str(event["fund_nav_event_id"]),
                    }
                    for event in ambiguous_events
                )
            elif not auto_cash_factor_applied:
                break_reasons.append(
                    {
                        "as_of_date": row_date.isoformat(),
                        "reason": "cash_disclosure_conflict",
                    }
                )
            if not auto_cash_factor_applied:
                continue
        next_level = segment_factor_level
        next_key = segment_factor_key
        interval_factors: list[dict[str, object]] = []
        interval_valid = True
        for event in interval_events:
            event_id = str(event["fund_nav_event_id"])
            event_date = _parse_nav_date(event.get("effective_date"))
            assert event_date is not None
            if event.get("_sequence_ambiguous"):
                interval_valid = False
                break_reasons.append(
                    {
                        "as_of_date": event_date.isoformat(),
                        "reason": "same_day_action_sequence_unverified",
                        "fund_nav_event_id": event_id,
                    }
                )
                break
            evidence_id: str | None = None
            if str(event.get("event_type")) == "cash_distribution":
                cash_per_unit = _parse_nav_decimal(event.get("cash_per_unit"))
                evidence = evidence_by_event_id.get(event_id)
                reinvestment_nav = _parse_nav_decimal(
                    evidence.get("reinvestment_nav") if evidence else None
                )
                if cash_per_unit is None or reinvestment_nav is None:
                    interval_valid = False
                    break_reasons.append(
                        {
                            "as_of_date": event_date.isoformat(),
                            "reason": "reinvestment_nav_not_observed",
                            "fund_nav_event_id": event_id,
                        }
                    )
                    break
                evidence_id = str(
                    evidence["fund_nav_reinvestment_evidence_id"]
                )
                action_multiplier = _quantized_factor_level(
                    Decimal("1") + cash_per_unit / reinvestment_nav
                )
                calculation_evidence: dict[str, object] = {
                    "cash_per_unit": str(cash_per_unit),
                    "reinvestment_nav": str(reinvestment_nav),
                    "calculation": "previous_factor*(1+cash_per_unit/reinvestment_nav)",
                }
            else:
                unit_ratio = _parse_nav_decimal(event.get("unit_ratio"))
                if unit_ratio is None:
                    interval_valid = False
                    break_reasons.append(
                        {
                            "as_of_date": event_date.isoformat(),
                            "reason": "unit_split_ratio_not_observed",
                            "fund_nav_event_id": event_id,
                        }
                    )
                    break
                action_multiplier = _quantized_factor_level(unit_ratio)
                calculation_evidence = {
                    "unit_ratio": str(unit_ratio),
                    "calculation": "previous_factor*unit_ratio",
                }
            next_level = _quantized_factor_level(next_level * action_multiplier)
            factor_key = f"event:{event_id}"
            interval_factors.append(
                {
                    "factor_logical_key": factor_key,
                    "previous_factor_logical_key": next_key,
                    "as_of_date": event_date.isoformat(),
                    "factor_level": str(next_level),
                    "factor_kind": "event_derived",
                    "fund_nav_event_id": event_id,
                    "fund_nav_reinvestment_evidence_id": evidence_id,
                    "evidence_kind": "fund_nav_event",
                    "method_version": FUND_NAV_PROJECTION_METHOD_VERSION,
                    "anchor_date": segment_anchor_date.isoformat(),
                    "source_provider": str(
                        event.get("source") or "fund_nav_event"
                    ),
                    "evidence": {
                        "previous_factor_logical_key": next_key,
                        "action_multiplier": str(action_multiplier),
                        **calculation_evidence,
                    },
                }
            )
            next_key = factor_key
        if not interval_valid:
            segment_live = False
            segment_broken = True
            continue

        adjustment_factors.extend(interval_factors)
        segment_factor_level = next_level
        segment_factor_key = next_key
        segment_event_cutoff = row_date
        total_return_nav = _format_total_return_nav_decimal(
            unit_nav * segment_factor_level
        )
        canonical["nav_with_dividend"] = total_return_nav
        canonical["nav_with_dividend_status"] = "complete"
        canonical["nav_with_dividend_source_provider"] = (
            "derived:fund_nav_reinvestment_factor"
        )
        canonical["nav_with_dividend_lineage"] = {
            "kind": "derived_dividend_reinvestment",
            "method_version": FUND_NAV_PROJECTION_METHOD_VERSION,
            "anchor_date": segment_anchor_date.isoformat(),
            "evidence": {
                "factor_logical_key": segment_factor_key,
                "unit_nav": str(unit_nav),
                "cash_cumulative_source_provider": str(fact["cash_provider"]),
                "cash_cumulative_source_evidence": dict(fact["cash_evidence"]),
            },
        }
        derived_count += 1

    total_return_dates = [
        str(row["as_of_date"])
        for row in canonical_rows
        if row.get("nav_with_dividend") is not None
        and row.get("nav_with_dividend_status") == "complete"
    ]
    complete_unit_dates = [
        str(fact["date"])
        for fact in row_facts
        if fact["unit_nav"] is not None and fact["unit_status"] == "complete"
    ]
    total_return_date_set = set(total_return_dates)
    missing_total_return_dates = [
        unit_date
        for unit_date in complete_unit_dates
        if unit_date not in total_return_date_set
    ]
    if not total_return_dates:
        adjustment_factors = []
        projection_status = "unavailable"
        projection_kind = (
            "provider_explicit"
            if any(fact["explicit_total_return"] is not None for fact in row_facts)
            else "event_derived"
        )
        projection_anchor_date: str | None = None
        if not any(fact["unit_status"] == "complete" for fact in row_facts):
            unavailable_reason = "complete_unit_nav_unavailable"
        elif any(fact["explicit_total_return"] is not None for fact in row_facts):
            unavailable_reason = "provider_total_return_not_factorable"
        elif first_complete_cash_date is None:
            unavailable_reason = "cash_cumulative_evidence_unavailable"
        else:
            unavailable_reason = "window_normalization_anchor_unavailable"
    else:
        projection_status = (
            "partial" if missing_total_return_dates else "complete"
        )
        factor_kinds = {
            str(factor.get("factor_kind") or "")
            for factor in adjustment_factors
        }
        if factor_kinds == {"provider_implied", "event_derived"}:
            projection_kind = "hybrid_reanchored"
        elif factor_kinds == {"provider_implied"}:
            projection_kind = "provider_explicit"
        else:
            projection_kind = "event_derived"
        projection_anchor_date = (
            current_segment_start_date.isoformat()
            if current_segment_start_date is not None
            else min(str(factor["anchor_date"]) for factor in adjustment_factors)
        )
        unavailable_reason = None

    projection_evidence: dict[str, object] = {
        "contract": FUND_NAV_PROJECTION_METHOD_VERSION,
        "source_observation_count": len(ordered_rows),
        "unit_nav_observation_count": sum(
            fact["unit_nav"] is not None for fact in row_facts
        ),
        "cash_cumulative_observation_count": sum(
            fact["cash_cumulative_nav"] is not None for fact in row_facts
        ),
        "provider_total_return_observation_count": sum(
            fact["explicit_total_return"] is not None for fact in row_facts
        ),
        "published_total_return_dates": total_return_dates,
        "missing_total_return_dates": missing_total_return_dates,
        "factor_logical_keys": [
            str(factor["factor_logical_key"]) for factor in adjustment_factors
        ],
        "segment_breaks": break_reasons,
        "cumulative_nav_semantics": cumulative_nav_semantics,
    }
    if auto_cash_distributions:
        projection_evidence["auto_cash_distributions"] = (
            auto_cash_distributions
        )
    if auto_cash_reporting_basis_resets:
        projection_evidence["auto_cash_reporting_basis_resets"] = (
            auto_cash_reporting_basis_resets
        )
    if unavailable_reason is not None:
        projection_evidence["unavailable_reason"] = unavailable_reason
    normalized_source_provider = source_provider.strip() or "platform_nav_import"
    projection_run: dict[str, object] = {
        "source_observation_fingerprint": source_fingerprint,
        "projection_kind": projection_kind,
        "projection_status": projection_status,
        "method_version": FUND_NAV_PROJECTION_METHOD_VERSION,
        "anchor_date": projection_anchor_date,
        "source_provider": normalized_source_provider,
        "evidence": projection_evidence,
        "created_by": FUND_NAV_PROJECTION_CREATED_BY,
    }
    return FundNavPublication(
        rows=canonical_rows,
        projection_run=projection_run,
        current_fund_nav_event_ids=current_event_ids,
        current_fund_nav_reinvestment_evidence_ids=current_evidence_ids,
        action_candidates=action_candidates,
        adjustment_factors=adjustment_factors,
        derived_total_return_count=derived_count,
        explicit_total_return_count=explicit_count,
    )


def _filter_rows_for_instrument(
    *,
    instrument: dict[str, object],
    rows: list[dict[str, object]],
) -> list[dict[str, object]]:
    if not rows:
        return []

    has_row_identity = any(
        str(row.get("instrument_code") or "").strip() or str(row.get("instrument_name") or "").strip()
        for row in rows
    )
    if not has_row_identity:
        return rows

    identifier_candidates = {
        str(item.get("identifier_value") or "").strip().upper()
        for item in list(instrument.get("identifiers", []))
        if str(item.get("identifier_value") or "").strip()
    }
    identifier_candidates.add(str(instrument.get("instrument_id") or "").strip().upper())

    instrument_name = _normalize_text_token(str(instrument.get("instrument_name") or ""))
    filtered: list[dict[str, object]] = []
    for row in rows:
        row_code = str(row.get("instrument_code") or "").strip().upper()
        row_name = _normalize_text_token(str(row.get("instrument_name") or ""))
        code_match = bool(row_code) and row_code in identifier_candidates
        name_match = bool(row_name and instrument_name) and (
            row_name == instrument_name
            or row_name in instrument_name
            or instrument_name in row_name
        )
        if code_match or name_match:
            filtered.append(row)
    return filtered


def _apply_nav_row_currency(
    *,
    rows: list[dict[str, object]],
    instrument_currency: object,
) -> list[dict[str, object]]:
    expected_currency = str(instrument_currency or "").strip().upper()
    if not expected_currency:
        raise ValueError("The selected fund has no canonical currency.")

    normalized_rows: list[dict[str, object]] = []
    for row in rows:
        supplied_currency = str(row.get("currency") or "").strip().upper()
        if supplied_currency and supplied_currency != expected_currency:
            raise ValueError(
                f'NAV row currency "{supplied_currency}" does not match the selected '
                f'fund currency "{expected_currency}".'
            )
        normalized_rows.append({**row, "currency": expected_currency})
    return normalized_rows


def _nav_import_instrument(instrument_id: str) -> dict[str, object] | None:
    instrument = get_instrument(instrument_id)
    if instrument is None:
        return None
    validate_nav_history_instrument_type(
        instrument_type=str(instrument.get("instrument_type") or ""),
        instrument_id=instrument_id,
    )
    return instrument


def _normalize_nav_rows_for_instrument(
    *,
    instrument: dict[str, object],
    rows: list[dict[str, object]],
) -> list[dict[str, object]]:
    filtered_rows = _filter_rows_for_instrument(instrument=instrument, rows=rows)
    if not filtered_rows:
        has_row_identity = any(
            str(row.get("instrument_code") or "").strip()
            or str(row.get("instrument_name") or "").strip()
            for row in rows
        )
        if has_row_identity:
            raise ValueError(
                "Detected NAV rows, but none matched the selected instrument identifier or name."
            )
        raise ValueError("No NAV rows detected.")

    currency_rows = _apply_nav_row_currency(
        rows=filtered_rows,
        instrument_currency=instrument.get("currency"),
    )
    return _merge_rows_by_date(currency_rows)


def _publish_fund_nav_projection(
    *,
    instrument_id: str,
    load_source_rows: Callable[[dict[str, object]], list[dict[str, object]]],
    source_provider: str,
    refresh_status: str,
    updated_by: str | None,
    mode: str,
    message_factory: Callable[[FundNavPublication], str],
) -> tuple[dict[str, object], FundNavPublication]:
    """Build and atomically publish a current raw snapshot with stale-write retry."""

    for _attempt in range(3):
        instrument = get_instrument(instrument_id)
        if instrument is None:
            raise ValueError(f'Instrument "{instrument_id}" no longer exists.')
        source_rows = load_source_rows(instrument)
        publication = _build_fund_nav_publication(
            instrument_id=instrument_id,
            rows=source_rows,
            source_provider=source_provider,
            instrument=instrument,
        )
        message = message_factory(publication)
        try:
            persisted = publish_fund_nav_history(
                instrument_id=instrument_id,
                rows=publication.rows,
                projection_run=publication.projection_run,
                current_fund_nav_event_ids=(
                    publication.current_fund_nav_event_ids
                ),
                current_fund_nav_reinvestment_evidence_ids=(
                    publication.current_fund_nav_reinvestment_evidence_ids
                ),
                adjustment_factors=publication.adjustment_factors,
                expected_market_data_updated_at=(
                    str(instrument["market_data_updated_at"])
                    if instrument.get("market_data_updated_at") is not None
                    else None
                ),
                refresh_status=refresh_status,
                updated_by=updated_by,
                message=message,
                mode=mode,
            )
        except StaleFundNavPublicationError:
            continue
        if persisted is None:
            raise ValueError(f'Instrument "{instrument_id}" no longer exists.')
        record = persisted.get("record")
        if not isinstance(record, dict):
            raise ValueError(
                f'Fund NAV publication for "{instrument_id}" returned no record.'
            )
        FundNavActionCandidateRepository(get_session_factory()).project_current(
            instrument_id=instrument_id,
            candidates=publication.action_candidates,
        )
        return record, publication
    raise StaleFundNavPublicationError(
        f'Fund NAV publication for "{instrument_id}" remained stale after 3 retries.'
    )


def _load_all_durable_nav_source_rows(
    *,
    instrument_id: str,
    instrument: dict[str, object],
) -> list[dict[str, object]]:
    repository = EmailIngestionRepository(get_session_factory())
    source_rows = [
        *list_raw_nav_observations(instrument_id=instrument_id),
        *repository.matched_nav_history(instrument_id),
    ]
    source_settings = dict(instrument.get("source_settings", {}))
    if str(source_settings.get("source_mode") or "").strip().lower() == "email":
        history_start_date = get_settings().email_history_start_date
        source_rows = [
            row
            for row in source_rows
            if (row_date := _parse_nav_date(row.get("as_of_date"))) is None
            or row_date >= history_start_date
        ]
    return _merge_rows_by_date(
        _apply_nav_row_currency(
            rows=source_rows,
            instrument_currency=instrument.get("currency"),
        )
    )


def rebuild_stale_fund_nav_projections(
    *,
    updated_by: str | None,
    include_inactive: bool = False,
    instrument_ids: set[str] | None = None,
) -> dict[str, object]:
    """Rebuild only current NAV projections created by an older method contract."""

    instruments = list_instruments(include_inactive=include_inactive)
    stale_current = list_stale_current_fund_nav_projections(
        method_version=FUND_NAV_PROJECTION_METHOD_VERSION,
        include_inactive=include_inactive,
        instrument_ids=instrument_ids,
    )
    instruments_by_id = {
        str(instrument.get("instrument_id") or ""): instrument
        for instrument in instruments
    }
    targets = [
        (instruments_by_id[target["instrument_id"]], target)
        for target in stale_current
        if target["instrument_id"] in instruments_by_id
    ]

    LOGGER.info(
        "fund NAV projection reconciliation targets=%s skipped=%s method_version=%s",
        len(targets),
        len(instruments) - len(targets),
        FUND_NAV_PROJECTION_METHOD_VERSION,
    )
    results: list[dict[str, object]] = []
    for instrument, current_run in targets:
        instrument_id = str(instrument["instrument_id"])
        source_settings = dict(instrument.get("source_settings", {}))
        source_provider = str(
            current_run.get("source_provider") or "canonical_nav_reprojection"
        ).strip()
        try:
            with market_data_item_timeout(instrument_id):
                record, publication = _publish_fund_nav_projection(
                    instrument_id=instrument_id,
                    load_source_rows=lambda current_instrument: (
                        _load_all_durable_nav_source_rows(
                            instrument_id=instrument_id,
                            instrument=current_instrument,
                        )
                    ),
                    source_provider=source_provider,
                    refresh_status="refreshed",
                    updated_by=updated_by,
                    message_factory=lambda current_publication: (
                        f"Rebuilt {len(current_publication.rows)} canonical NAV dates "
                        f"under {FUND_NAV_PROJECTION_METHOD_VERSION}; "
                        f"{len(current_publication.action_candidates)} NAV action "
                        "signal(s) require review."
                    ),
                    mode="projection_reconciliation",
                )
            message = (
                f"Rebuilt {len(publication.rows)} canonical NAV dates under "
                f"{FUND_NAV_PROJECTION_METHOD_VERSION}."
            )
            status = "refreshed"
            result_instrument = record
        except Exception as error:
            LOGGER.exception(
                "fund NAV projection reconciliation failed instrument_id=%s",
                instrument_id,
            )
            message = (
                "Fund NAV projection reconciliation failed: "
                f"{type(error).__name__}: {error}"
            )
            result_instrument = update_refresh_status(
                instrument_id=instrument_id,
                status="failed",
                message=message,
                updated_by=updated_by,
                mode="projection_reconciliation",
            ) or instrument
            status = "failed"
        result_source_settings = dict(
            result_instrument.get("source_settings", source_settings)
        )
        results.append(
            {
                "instrument_id": instrument_id,
                "instrument_name": str(
                    result_instrument.get("instrument_name")
                    or instrument.get("instrument_name")
                    or ""
                ),
                "instrument_type": str(
                    result_instrument.get("instrument_type")
                    or instrument.get("instrument_type")
                    or ""
                ),
                "source_mode": result_source_settings.get("source_mode") or "manual",
                "source_api_profile": result_source_settings.get("source_api_profile") or "",
                "status": status,
                "message": message,
            }
        )

    return {
        "source": "fund_nav_projection",
        "refreshed_count": sum(item["status"] == "refreshed" for item in results),
        "skipped_count": len(instruments) - len(targets),
        "results": results,
    }


def _latest_nav_date_from_instrument(instrument: dict[str, object]) -> date | None:
    latest_nav_date: date | None = None
    for point in list(instrument.get("market_data", [])):
        if not isinstance(point, dict):
            continue
        if str(point.get("metric_family") or "").strip() != "nav":
            continue
        if str(point.get("status") or "").strip().lower() != "complete":
            continue
        current_date = _parse_nav_date(point.get("as_of_date"))
        if current_date is None:
            continue
        if latest_nav_date is None or current_date > latest_nav_date:
            latest_nav_date = current_date
    return latest_nav_date


def _latest_market_data_date_from_instrument(
    instrument: dict[str, object],
    *,
    metric_family: str,
    quote_bases: set[str],
) -> date | None:
    latest_date: date | None = None
    for point in list(instrument.get("market_data", [])):
        if not isinstance(point, dict):
            continue
        if str(point.get("metric_family") or "").strip() != metric_family:
            continue
        if str(point.get("quote_basis") or "").strip() not in quote_bases:
            continue
        if str(point.get("status") or "").strip().lower() != "complete":
            continue
        current_date = _parse_nav_date(point.get("as_of_date"))
        if current_date is None:
            continue
        if latest_date is None or current_date > latest_date:
            latest_date = current_date
    return latest_date


def _format_tushare_date(value: date) -> str:
    return value.strftime("%Y%m%d")


def _tushare_query_start_date(*, latest_date: date | None, full_history: bool) -> date:
    if full_history or latest_date is None or latest_date < TUSHARE_HISTORY_START_DATE:
        return TUSHARE_HISTORY_START_DATE
    return latest_date + timedelta(days=1)


def _tushare_profile_enabled(source_settings: dict[str, object]) -> bool:
    profile = str(source_settings.get("source_api_profile") or "").strip().lower()
    return profile in TUSHARE_PROFILE_ALIASES


def _tushare_identifier_code(instrument: dict[str, object]) -> str | None:
    identifiers = [
        item for item in list(instrument.get("identifiers", [])) if isinstance(item, dict)
    ]
    ordered_identifiers = sorted(
        identifiers,
        key=lambda item: (
            0 if bool(item.get("is_primary")) else 1,
            0 if str(item.get("identifier_type") or "").strip().lower() == "ticker" else 1,
        ),
    )
    for identifier in ordered_identifiers:
        raw_value = str(identifier.get("identifier_value") or "").strip().upper()
        if re.fullmatch(r"[0-9A-Z]{5,12}\.(?:OF|SH|SZ|CSI|CNI)", raw_value):
            return raw_value
    return None


def _policy_role_values(policy: dict[str, object], role: str) -> list[str]:
    raw_values = policy.get(role)
    if not isinstance(raw_values, list):
        return []
    values: list[str] = []
    for raw_value in raw_values:
        value = str(raw_value or "").strip()
        if value and value not in values:
            values.append(value)
    return values


def _ensure_tushare_listed_security_quote_policy(
    *,
    instrument_id: str,
    instrument: dict[str, object],
) -> None:
    current_policy = instrument.get("quote_selection_policy", {})
    if not isinstance(current_policy, dict):
        current_policy = {}
    target_policy = {
        role: list(values)
        for role, values in TUSHARE_LISTED_SECURITY_QUOTE_SELECTION_POLICY.items()
    }
    if all(
        _policy_role_values(current_policy, role) == values
        for role, values in target_policy.items()
    ):
        return
    upsert_quote_selection_policy(
        instrument_id=instrument_id,
        quote_selection_policy=target_policy,
    )


class TushareRefreshError(RuntimeError):
    pass


def _is_tushare_missing_value(value: object) -> bool:
    if value is None:
        return True
    try:
        return bool(value != value)
    except (TypeError, ValueError):
        return False


def _call_tushare_api(
    *,
    api_name: str,
    params: dict[str, object],
    fields: str,
) -> list[dict[str, object]]:
    settings = get_settings()
    if not settings.tushare_ready:
        raise TushareRefreshError("Tushare token is not configured. Set PORTFOLIO_OPS_PLATFORM_TUSHARE_TOKEN first.")

    try:
        pro = tushare_client.create_tushare_client(
            token=settings.tushare_token,
            api_url=settings.tushare_api_url,
            timeout_seconds=getattr(settings, "tushare_timeout_seconds", 30),
        )
        frame = tushare_client.invoke_tushare_api(
            client=pro,
            api_name=api_name,
            params=params,
            fields=fields,
        )
    except Exception as exc:
        safe_message = tushare_client.redact_tushare_error_message(
            exc,
            token=settings.tushare_token,
        )
        raise TushareRefreshError(
            f"Tushare SDK request failed: {safe_message}"
        ) from exc

    if frame is None:
        return []
    to_dict = getattr(frame, "to_dict", None)
    if not callable(to_dict):
        return []
    raw_rows = to_dict("records")
    if not isinstance(raw_rows, list):
        return []

    rows: list[dict[str, object]] = []
    for raw_row in raw_rows:
        if not isinstance(raw_row, dict):
            continue
        row: dict[str, object] = {}
        for key, value in raw_row.items():
            if _is_tushare_missing_value(value):
                row[str(key)] = None
            else:
                row[str(key)] = value
        rows.append(row)
    return rows


def _tushare_nav_rows(
    rows: list[dict[str, object]],
    *,
    instrument_currency: object,
    latest_date: date | None,
) -> list[dict[str, object]]:
    normalized_currency = str(instrument_currency or "").strip().upper()
    if not normalized_currency:
        raise ValueError("The selected fund has no canonical currency.")
    prepared_rows: list[dict[str, object]] = []
    for row in rows:
        point_date = _parse_nav_date(row.get("nav_date") or row.get("end_date") or row.get("ann_date"))
        if point_date is None:
            continue
        if point_date < TUSHARE_HISTORY_START_DATE:
            continue
        # Keep the latest stored date so provider corrections on that date are
        # retained as a new raw evidence revision.
        if latest_date is not None and point_date < latest_date:
            continue
        nav = _parse_nav_decimal(row.get("unit_nav"))
        cash_cumulative_nav = _parse_nav_decimal(row.get("accum_nav"))
        total_return_nav = _parse_nav_decimal(row.get("adj_nav"))
        if nav is None and cash_cumulative_nav is None and total_return_nav is None:
            continue
        prepared_rows.append(
            {
                "as_of_date": point_date.isoformat(),
                "nav": nav,
                "cash_cumulative_nav": cash_cumulative_nav,
                "nav_with_dividend": total_return_nav,
                "_total_return_source_field": "adj_nav",
                "currency": normalized_currency,
                "frequency": "daily",
            }
        )
    return _merge_rows_by_date(prepared_rows)


def _tushare_price_rows(
    rows: list[dict[str, object]],
    *,
    latest_date: date | None,
) -> list[dict[str, object]]:
    prepared_rows: list[dict[str, object]] = []
    for row in rows:
        point_date = _parse_nav_date(row.get("trade_date"))
        if point_date is None:
            continue
        if point_date < TUSHARE_HISTORY_START_DATE:
            continue
        if latest_date is not None and point_date <= latest_date:
            continue
        close_value = _parse_nav_decimal(row.get("close"))
        if close_value is None:
            continue
        prepared_rows.append(
            {
                "as_of_date": point_date,
                "value": close_value,
                "open": _parse_nav_decimal(row.get("open")),
                "high": _parse_nav_decimal(row.get("high")),
                "low": _parse_nav_decimal(row.get("low")),
                "close": close_value,
                "previous_close": _parse_nav_decimal(row.get("pre_close")),
                "volume": _parse_nonnegative_decimal(row.get("vol")),
                "turnover": _parse_nonnegative_decimal(row.get("amount")),
            }
        )
    return sorted(prepared_rows, key=lambda item: item["as_of_date"])


def _tushare_price_bar_rows(
    rows: list[dict[str, object]],
    *,
    api_name: str,
    adjustment_factors: dict[date, Decimal] | None = None,
) -> list[dict[str, object]]:
    bars: list[dict[str, object]] = []
    for row in rows:
        if not all(row.get(field_name) is not None for field_name in ("open", "high", "low", "close")):
            continue
        point_date = row.get("as_of_date")
        if not isinstance(point_date, date):
            continue
        volume = row.get("volume")
        turnover = row.get("turnover")
        adjustment_factor = (adjustment_factors or {}).get(point_date)
        bars.append(
            {
                "as_of_date": point_date,
                "open": row["open"],
                "high": row["high"],
                "low": row["low"],
                "close": row["close"],
                "previous_close": row.get("previous_close"),
                "volume": volume,
                "turnover": turnover,
                "adjustment_factor": adjustment_factor,
                "currency": "CNY",
                "volume_unit": "lot" if volume is not None else None,
                "turnover_unit": "thousand_cny" if turnover is not None else None,
                "provider": (
                    f"tushare:{api_name}+adjustment_factor"
                    if adjustment_factor is not None
                    else f"tushare:{api_name}"
                ),
                "status": (
                    "complete"
                    if volume is not None and turnover is not None
                    else "partial"
                ),
            }
        )
    return bars


def _tushare_adjustment_factors(
    rows: list[dict[str, object]],
) -> dict[date, Decimal]:
    factors: dict[date, Decimal] = {}
    for row in rows:
        point_date = _parse_nav_date(row.get("trade_date"))
        factor = _parse_nav_decimal(row.get("adj_factor"))
        if point_date is None or factor is None or factor <= 0:
            continue
        if point_date < TUSHARE_HISTORY_START_DATE:
            continue
        factors[point_date] = factor
    return factors


# Cash distributions also move Tushare adjustment factors, so a factor change
# by itself is not a share-split event.  We accept only conventional rational
# share ratios whose inverse mechanical move is visible in raw close while the
# adjusted series remains continuous.  Provider inference is never sufficient
# authority to change portfolio units: every candidate remains ``detected``
# until an issuer/exchange/CSD announcement confirms ratio and rounding.
TUSHARE_SHARE_SPLIT_RATIOS: tuple[tuple[Decimal, Decimal], ...] = (
    (Decimal("10"), Decimal("1")),
    (Decimal("5"), Decimal("1")),
    (Decimal("4"), Decimal("1")),
    (Decimal("3"), Decimal("1")),
    (Decimal("2"), Decimal("1")),
    (Decimal("3"), Decimal("2")),
    (Decimal("5"), Decimal("4")),
    (Decimal("4"), Decimal("5")),
    (Decimal("2"), Decimal("3")),
    (Decimal("1"), Decimal("2")),
    (Decimal("1"), Decimal("3")),
    (Decimal("1"), Decimal("4")),
    (Decimal("1"), Decimal("5")),
    (Decimal("1"), Decimal("10")),
)


def _detect_tushare_share_splits(
    *,
    factors: dict[date, Decimal],
    close_by_date: dict[date, Decimal],
) -> list[dict[str, object]]:
    if len(factors) < 2 or len(close_by_date) < 2:
        return []
    factor_dates = sorted(factors)
    close_dates = sorted(close_by_date)
    candidates: list[dict[str, object]] = []
    for index, effective_date in enumerate(factor_dates[1:], start=1):
        prior_factor_date = factor_dates[index - 1]
        factor_before = factors[prior_factor_date]
        factor_after = factors[effective_date]
        if factor_before <= 0 or factor_after <= 0:
            continue
        observed_ratio = factor_after / factor_before
        nearest_new, nearest_old = min(
            TUSHARE_SHARE_SPLIT_RATIOS,
            key=lambda ratio: abs(observed_ratio - (ratio[0] / ratio[1])),
        )
        canonical_ratio = nearest_new / nearest_old
        relative_ratio_error = abs(observed_ratio / canonical_ratio - Decimal("1"))
        if relative_ratio_error > Decimal("0.005"):
            continue

        prior_close_dates = [point_date for point_date in close_dates if point_date < effective_date]
        if effective_date not in close_by_date or not prior_close_dates:
            continue
        prior_close_date = prior_close_dates[-1]
        close_before = close_by_date[prior_close_date]
        close_after = close_by_date[effective_date]
        if close_before <= 0 or close_after <= 0:
            continue
        raw_return = close_after / close_before - Decimal("1")
        adjusted_return = (close_after * observed_ratio) / close_before - Decimal("1")
        if abs(raw_return) < Decimal("0.15") or abs(adjusted_return) > Decimal("0.25"):
            continue

        candidates.append(
            {
                "effective_date": effective_date,
                "new_units": _decimal_text(nearest_new),
                "old_units": _decimal_text(nearest_old),
                "status": "detected",
                "quantity_rounding": "exact",
                "quantity_precision": 0,
                "provenance": {
                    "detection_method": "tushare_factor_price_continuity/v1",
                    "factor_before_date": prior_factor_date.isoformat(),
                    "factor_before": _decimal_text(factor_before),
                    "factor_after": _decimal_text(factor_after),
                    "observed_factor_ratio": _decimal_text(observed_ratio),
                    "canonical_ratio": f"{_decimal_text(nearest_new)}:{_decimal_text(nearest_old)}",
                    "relative_ratio_error": _decimal_text(relative_ratio_error),
                    "prior_close_date": prior_close_date.isoformat(),
                    "close_before": _decimal_text(close_before),
                    "close_after": _decimal_text(close_after),
                    "raw_return": _decimal_text(raw_return),
                    "adjusted_return": _decimal_text(adjusted_return),
                },
            }
        )
    return candidates


def _existing_price_values(
    instrument: dict[str, object],
    *,
    quote_basis: str,
) -> dict[date, Decimal]:
    values: dict[date, Decimal] = {}
    for point in list(instrument.get("market_data", [])):
        if not isinstance(point, dict):
            continue
        if str(point.get("metric_family") or "").strip() != "price":
            continue
        if str(point.get("quote_basis") or "").strip() != quote_basis:
            continue
        if str(point.get("status") or "").strip().lower() == "unavailable":
            continue
        point_date = _parse_nav_date(point.get("as_of_date"))
        value = _parse_nav_decimal(point.get("value"))
        if point_date is not None and value is not None:
            values[point_date] = value
    return values


def _existing_complete_price_dates(
    instrument: dict[str, object],
    *,
    quote_basis: str,
) -> set[date]:
    dates: set[date] = set()
    for point in list(instrument.get("market_data", [])):
        if not isinstance(point, dict):
            continue
        if str(point.get("metric_family") or "").strip() != "price":
            continue
        if str(point.get("quote_basis") or "").strip() != quote_basis:
            continue
        if str(point.get("status") or "").strip() != "complete":
            continue
        point_date = _parse_nav_date(point.get("as_of_date"))
        if point_date is not None:
            dates.add(point_date)
    return dates


def _price_bar_repair_start_date(
    *,
    instrument_id: str,
    instrument: dict[str, object],
    require_adjustment_factor: bool,
) -> date | None:
    """Find the first canonical close when persisted OHLCV needs repair."""

    canonical_dates = sorted(
        point_date
        for point_date in _existing_price_values(
            instrument,
            quote_basis="close",
        )
        if point_date >= TUSHARE_HISTORY_START_DATE
    )
    if not canonical_dates:
        return None

    coverage = get_price_bar_coverage(instrument_id=instrument_id)
    row_count = int(coverage.get("row_count") or 0)
    factor_count = int(coverage.get("adjustment_factor_count") or 0)
    raw_first_bar_date = coverage.get("first_date")
    raw_latest_bar_date = coverage.get("latest_date")
    first_bar_date = (
        _parse_nav_date(raw_first_bar_date) if raw_first_bar_date is not None else None
    )
    latest_bar_date = (
        _parse_nav_date(raw_latest_bar_date)
        if raw_latest_bar_date is not None
        else None
    )
    coverage_incomplete = (
        row_count < len(canonical_dates)
        or first_bar_date is None
        or first_bar_date > canonical_dates[0]
        or latest_bar_date is None
        or latest_bar_date < canonical_dates[-1]
    )
    factor_coverage_incomplete = (
        require_adjustment_factor and factor_count < row_count
    )
    if coverage_incomplete or factor_coverage_incomplete:
        return canonical_dates[0]
    return None


QFQ_FACTOR_PATTERN = re.compile(r"(?:^|:)latest_factor=([0-9]+(?:\.[0-9]+)?)$")


def _stored_qfq_latest_factor(instrument: dict[str, object]) -> Decimal | None:
    latest_point: tuple[date, Decimal] | None = None
    for point in list(instrument.get("market_data", [])):
        if not isinstance(point, dict):
            continue
        if str(point.get("quote_basis") or "").strip() != "adjusted_close":
            continue
        if str(point.get("status") or "").strip() != "complete":
            continue
        provider = str(point.get("provider") or "")
        match = QFQ_FACTOR_PATTERN.search(provider)
        point_date = _parse_nav_date(point.get("as_of_date"))
        if match is None or point_date is None:
            continue
        factor = _parse_nav_decimal(match.group(1))
        if factor is not None and (latest_point is None or point_date > latest_point[0]):
            latest_point = (point_date, factor)
    return latest_point[1] if latest_point is not None else None


def _decimal_text(value: Decimal) -> str:
    return format(value.normalize(), "f")


def import_nav_text(
    *,
    instrument_id: str,
    raw_text: str,
    provider: str | None,
    status: str,
    updated_by: str | None,
) -> dict[str, object] | None:
    instrument = _nav_import_instrument(instrument_id)
    if instrument is None:
        return None
    rows = _parse_nav_rows_from_text(raw_text)
    normalized_rows = _normalize_nav_rows_for_instrument(
        instrument=instrument,
        rows=rows,
    )
    point_provider = provider or "platform_manual_import"
    record_raw_nav_observations(
        instrument_id=instrument_id,
        rows=normalized_rows,
        source_kind="manual_import",
        source_ref="text:" + hashlib.sha256(raw_text.encode("utf-8")).hexdigest(),
        provider=point_provider,
        status=status,
        evidence={"input_type": "text"},
    )
    record, _publication = _publish_fund_nav_projection(
        instrument_id=instrument_id,
        load_source_rows=lambda current_instrument: _load_all_durable_nav_source_rows(
            instrument_id=instrument_id,
            instrument=current_instrument,
        ),
        source_provider=point_provider,
        refresh_status="imported",
        updated_by=updated_by,
        message_factory=lambda publication: (
            f"Stored {len(normalized_rows)} raw NAV observations and rebuilt "
            f"{len(publication.rows)} canonical NAV dates; "
            f"{len(publication.action_candidates)} NAV action signal(s) require review."
        ),
        mode="manual",
    )
    return record


def preview_nav_import(
    *,
    instrument_id: str,
    raw_text: str | None = None,
    file_name: str | None = None,
    file_bytes: bytes | None = None,
) -> list[dict[str, object]] | None:
    instrument = _nav_import_instrument(instrument_id)
    if instrument is None:
        return None
    if raw_text is not None:
        rows = _parse_nav_rows_from_text(raw_text)
    elif file_name is not None and file_bytes is not None:
        rows = _parse_nav_rows_from_uploaded_file(file_name=file_name, file_bytes=file_bytes)
    else:
        raise ValueError("Provide NAV import text or a NAV file payload.")

    normalized_rows = _normalize_nav_rows_for_instrument(
        instrument=instrument,
        rows=rows,
    )
    return _build_fund_nav_publication(
        instrument_id=instrument_id,
        rows=_stamp_nav_value_statuses(normalized_rows, status="complete"),
        source_provider="preview_import",
        instrument=instrument,
    ).rows


def import_nav_file(
    *,
    instrument_id: str,
    file_name: str,
    file_bytes: bytes,
    provider: str | None,
    status: str,
    updated_by: str | None,
) -> dict[str, object] | None:
    instrument = _nav_import_instrument(instrument_id)
    if instrument is None:
        return None
    rows = _parse_nav_rows_from_uploaded_file(file_name=file_name, file_bytes=file_bytes)
    normalized_rows = _normalize_nav_rows_for_instrument(
        instrument=instrument,
        rows=rows,
    )
    point_provider = provider or f"platform_file_import:{file_name}"
    record_raw_nav_observations(
        instrument_id=instrument_id,
        rows=normalized_rows,
        source_kind="manual_import",
        source_ref="file:" + hashlib.sha256(file_bytes).hexdigest(),
        provider=point_provider,
        status=status,
        evidence={"file_name": file_name},
    )
    record, _publication = _publish_fund_nav_projection(
        instrument_id=instrument_id,
        load_source_rows=lambda current_instrument: _load_all_durable_nav_source_rows(
            instrument_id=instrument_id,
            instrument=current_instrument,
        ),
        source_provider=point_provider,
        refresh_status="imported",
        updated_by=updated_by,
        message_factory=lambda publication: (
            f"Stored {len(normalized_rows)} raw NAV observations from {file_name} "
            f"and rebuilt {len(publication.rows)} canonical NAV dates; "
            f"{len(publication.action_candidates)} NAV action signal(s) require review."
        ),
        mode="manual",
    )
    return record


def refresh_market_data(
    *,
    instrument_id: str,
    updated_by: str | None,
    full_history: bool = False,
    source: str = "configured",
) -> dict[str, object] | None:
    instrument = get_instrument(instrument_id)
    if instrument is None:
        return None

    source_settings = dict(instrument.get("source_settings", {}))
    requested_source = str(source or "configured").strip().lower()
    if requested_source == "email":
        _run_email_ingestion_batch(
            instruments=list_instruments(include_inactive=False),
            settings=get_settings(),
            updated_by=updated_by,
            full_history=full_history,
        )
        return get_instrument(instrument_id)
    if requested_source == "tushare":
        return _refresh_from_tushare(
            instrument_id=instrument_id,
            instrument=instrument,
            updated_by=updated_by,
            full_history=full_history,
        )

    source_mode = str(source_settings.get("source_mode") or "manual")
    if source_mode == "email":
        _run_email_ingestion_batch(
            instruments=list_instruments(include_inactive=False),
            settings=get_settings(),
            updated_by=updated_by,
            full_history=full_history,
        )
        return get_instrument(instrument_id)
    if source_mode == "api":
        profile = str(source_settings.get("source_api_profile") or "").strip()
        if profile.lower() in TUSHARE_PROFILE_ALIASES:
            return _refresh_from_tushare(
                instrument_id=instrument_id,
                instrument=instrument,
                updated_by=updated_by,
                full_history=full_history,
            )
        return update_refresh_status(
            instrument_id=instrument_id,
            status="blocked",
            message=(
                f"API refresh profile {profile or 'unconfigured profile'} is not wired in Database Dashboard yet."
            ),
            updated_by=updated_by,
            mode="api",
        )

    location = str(source_settings.get("source_location") or "Database Dashboard").strip()
    return update_refresh_status(
        instrument_id=instrument_id,
        status="awaiting_manual_import",
        message=f"Manual source active in Database Dashboard. Import canonical NAV from {location}.",
        updated_by=updated_by,
        mode="manual",
    )


def refresh_market_data_with_timeout(
    *,
    instrument_id: str,
    updated_by: str | None,
    full_history: bool = False,
    source: str = "configured",
) -> dict[str, object] | None:
    try:
        with market_data_item_timeout(instrument_id):
            return refresh_market_data(
                instrument_id=instrument_id,
                updated_by=updated_by,
                full_history=full_history,
                source=source,
            )
    except MarketDataItemTimeout as exc:
        normalized_source = str(source or "configured").strip().lower()
        refresh_mode = (
            "email"
            if normalized_source == "email"
            else "api" if normalized_source == "tushare" else None
        )
        LOGGER.warning(
            "market data item timed out instrument_id=%s source=%s message=%s",
            instrument_id,
            normalized_source,
            exc,
        )
        return update_refresh_status(
            instrument_id=instrument_id,
            status="failed",
            message=str(exc),
            updated_by=updated_by,
            mode=refresh_mode,
        )


def _refresh_tushare_listed_security(
    *,
    instrument_id: str,
    instrument: dict[str, object],
    ts_code: str,
    price_api_name: str,
    factor_api_name: str,
    updated_by: str | None,
    full_history: bool,
) -> dict[str, object] | None:
    """Refresh raw close and a latest-date-normalized forward-adjusted close.

    adjusted_close[t] = close[t] * factor[t] / factor[latest]. The latest
    adjusted close therefore equals the tradable close, while historical total
    returns reflect distributions and share adjustments. If the latest factor
    changes, all available history is recomputed from canonical raw closes.
    """
    if str(instrument.get("currency") or "").strip().upper() != "CNY":
        raise TushareRefreshError(
            f"Tushare listed-security price data for {ts_code} requires a CNY instrument."
        )
    _ensure_tushare_listed_security_quote_policy(
        instrument_id=instrument_id,
        instrument=instrument,
    )
    today = date.today()
    latest_close_date = _latest_market_data_date_from_instrument(
        instrument,
        metric_family="price",
        quote_bases={"close"},
    )
    repair_start = (
        None
        if full_history
        else _price_bar_repair_start_date(
            instrument_id=instrument_id,
            instrument=instrument,
            require_adjustment_factor=True,
        )
    )
    if full_history or latest_close_date is None:
        query_start = TUSHARE_HISTORY_START_DATE
    elif repair_start is not None:
        query_start = repair_start
    else:
        query_start = max(
            TUSHARE_HISTORY_START_DATE,
            min(latest_close_date, today) - timedelta(days=7),
        )
    params: dict[str, object] = {
        "ts_code": ts_code,
        "start_date": _format_tushare_date(query_start),
        "end_date": _format_tushare_date(today),
    }
    close_rows = _tushare_price_rows(
        _call_tushare_api(
            api_name=price_api_name,
            params=params,
            fields=(
                "ts_code,trade_date,open,high,low,close,pre_close,vol,amount"
            ),
        ),
        latest_date=None,
    )
    factor_rows = _call_tushare_api(
        api_name=factor_api_name,
        params=params,
        fields="ts_code,trade_date,adj_factor",
    )
    if not close_rows and not factor_rows:
        return update_refresh_status(
            instrument_id=instrument_id,
            status="failed",
            message=(
                f"Tushare returned an empty response for both {price_api_name} and "
                f"{factor_api_name} for established listed security {ts_code}; "
                "existing canonical price history was preserved without rewriting it."
            ),
            updated_by=updated_by,
            mode="api",
        )
    factors = _tushare_adjustment_factors(factor_rows)
    close_by_date = _existing_price_values(instrument, quote_basis="close")
    existing_adjusted_by_date = _existing_price_values(
        instrument,
        quote_basis="adjusted_close",
    )
    existing_adjusted_dates = _existing_complete_price_dates(
        instrument,
        quote_basis="adjusted_close",
    )
    fetched_close_by_date = {
        row["as_of_date"]: Decimal(str(row["value"]))
        for row in close_rows
        if isinstance(row.get("as_of_date"), date)
    }
    close_by_date.update(fetched_close_by_date)
    unpaired_existing_dates = set(close_by_date).difference(existing_adjusted_dates)
    previous_latest_factor = _stored_qfq_latest_factor(instrument)
    current_latest_factor = factors[max(factors)] if factors else previous_latest_factor
    factor_changed = (
        current_latest_factor is not None
        and (
            previous_latest_factor is None
            or current_latest_factor != previous_latest_factor
        )
    )
    candidate_adjusted_dates = (
        set(close_by_date)
        if full_history or factor_changed
        else set(fetched_close_by_date).union(unpaired_existing_dates)
    )
    missing_candidate_factors = candidate_adjusted_dates.difference(factors)
    if query_start > TUSHARE_HISTORY_START_DATE and (
        full_history or factor_changed or missing_candidate_factors
    ):
        full_factor_rows = _call_tushare_api(
            api_name=factor_api_name,
            params={
                "ts_code": ts_code,
                "start_date": _format_tushare_date(TUSHARE_HISTORY_START_DATE),
                "end_date": _format_tushare_date(today),
            },
            fields="ts_code,trade_date,adj_factor",
        )
        factors = _tushare_adjustment_factors(full_factor_rows)
        if factors:
            current_latest_factor = factors[max(factors)]
            factor_changed = (
                previous_latest_factor is None
                or current_latest_factor != previous_latest_factor
            )

    adjusted_dates = (
        sorted(close_by_date)
        if full_history or factor_changed
        else sorted(set(fetched_close_by_date).union(unpaired_existing_dates))
    )
    adjusted_points: list[dict[str, object]] = []
    adjusted_complete_dates: set[date] = set()
    missing_adjustment_dates: set[date] = set()
    if current_latest_factor is not None:
        factor_text = _decimal_text(current_latest_factor)
        for point_date in adjusted_dates:
            factor = factors.get(point_date)
            close_value = close_by_date.get(point_date)
            if factor is None or close_value is None:
                missing_adjustment_dates.add(point_date)
                continue
            adjusted_close = close_value * factor / current_latest_factor
            adjusted_complete_dates.add(point_date)
            adjusted_points.append(
                {
                    "metric_family": "price",
                    "quote_basis": "adjusted_close",
                    "as_of_date": point_date,
                    "value": _decimal_text(adjusted_close),
                    "currency": "CNY",
                    "provider": (
                        f"tushare:{factor_api_name}:qfq:latest_factor={factor_text}"
                    ),
                    "status": "complete",
                }
            )
    else:
        missing_adjustment_dates.update(adjusted_dates)

    points: list[dict[str, object]] = [
        {
            "metric_family": "price",
            "quote_basis": "close",
            "as_of_date": point_date,
            "value": value,
            "currency": "CNY",
            "provider": f"tushare:{price_api_name}",
            "status": (
                "complete" if point_date in adjusted_complete_dates else "partial"
            ),
        }
        for point_date, value in fetched_close_by_date.items()
    ]
    for point_date in sorted(
        adjusted_complete_dates.intersection(unpaired_existing_dates).difference(
            fetched_close_by_date
        )
    ):
        close_value = close_by_date.get(point_date)
        if close_value is None:
            continue
        points.append(
            {
                "metric_family": "price",
                "quote_basis": "close",
                "as_of_date": point_date,
                "value": close_value,
                "currency": "CNY",
                "provider": f"tushare:{price_api_name}",
                "status": "complete",
            }
        )
    for point_date in sorted(missing_adjustment_dates.difference(fetched_close_by_date)):
        close_value = close_by_date.get(point_date)
        if close_value is None:
            continue
        points.append(
            {
                "metric_family": "price",
                "quote_basis": "close",
                "as_of_date": point_date,
                "value": close_value,
                "currency": "CNY",
                "provider": f"tushare:{price_api_name}",
                "status": "partial",
            }
        )
    for point_date in sorted(missing_adjustment_dates.intersection(existing_adjusted_by_date)):
        points.append(
            {
                "metric_family": "price",
                "quote_basis": "adjusted_close",
                "as_of_date": point_date,
                "value": existing_adjusted_by_date[point_date],
                "currency": "CNY",
                "provider": f"tushare:{factor_api_name}:stale_missing_factor",
                "status": "partial",
            }
        )
    points.extend(adjusted_points)

    changed_count = upsert_market_data_points(
        instrument_id=instrument_id,
        rows=points,
    )
    if changed_count is None:
        return None
    bar_changed_count = upsert_price_bars(
        instrument_id=instrument_id,
        rows=_tushare_price_bar_rows(
            close_rows,
            api_name=price_api_name,
            adjustment_factors=factors,
        ),
    )
    if bar_changed_count is None:
        return None
    detected_actions = _detect_tushare_share_splits(
        factors=factors,
        close_by_date=close_by_date,
    )
    persisted_actions: list[dict[str, object]] = []
    for action in detected_actions:
        persisted_action = upsert_corporate_action_event(
            instrument_id=instrument_id,
            action_type="share_split",
            effective_date=action["effective_date"],
            new_units=action["new_units"],
            old_units=action["old_units"],
            source=f"tushare:{factor_api_name}",
            status=str(action["status"]),
            quantity_rounding=str(action["quantity_rounding"]),
            quantity_precision=int(action["quantity_precision"]),
            external_event_id=(
                f"{ts_code}:{action['effective_date'].isoformat()}:share_split"
            ),
            provenance=dict(action["provenance"]),
        )
        if isinstance(persisted_action, dict):
            persisted_actions.append(persisted_action)
    if missing_adjustment_dates:
        missing_dates_text = ", ".join(
            point_date.isoformat()
            for point_date in sorted(missing_adjustment_dates)[:10]
        )
        if len(missing_adjustment_dates) > 10:
            missing_dates_text += f", +{len(missing_adjustment_dates) - 10} more"
        return update_refresh_status(
            instrument_id=instrument_id,
            status="partial",
            message=(
                f"Tushare listed-security refresh is partial for {ts_code}: "
                f"missing adjustment factor for {missing_dates_text}. "
                "Affected raw closes were stored as partial; no complete close was "
                "stored without a same-date adjusted_close."
            ),
            updated_by=updated_by,
            mode="api",
        )
    if changed_count == 0 and bar_changed_count == 0:
        return update_refresh_status(
            instrument_id=instrument_id,
            status="no_new_data",
            message=f"Tushare returned no changed listed-security rows for {ts_code}.",
            updated_by=updated_by,
            mode="api",
        )
    adjusted_count = sum(
        1 for point in points if point.get("quote_basis") == "adjusted_close"
    )
    confirmed_action_count = sum(
        1 for action in persisted_actions if action.get("status") == "confirmed"
    )
    detected_action_count = sum(
        1 for action in persisted_actions if action.get("status") == "detected"
    )
    return update_refresh_status(
        instrument_id=instrument_id,
        status="refreshed",
        message=(
            f"Updated {changed_count} market-data points for {ts_code}: "
            f"raw close for valuation/trading and {adjusted_count} qfq adjusted closes "
            f"for charts and total return; {bar_changed_count} raw OHLCV bars; "
            f"{confirmed_action_count} confirmed and {detected_action_count} review-required "
            "share-adjustment event(s)."
        ),
        updated_by=updated_by,
        mode="api",
    )


def _refresh_from_tushare(
    *,
    instrument_id: str,
    instrument: dict[str, object],
    updated_by: str | None,
    full_history: bool,
) -> dict[str, object] | None:
    ts_code = _tushare_identifier_code(instrument)
    if ts_code is None:
        return update_refresh_status(
            instrument_id=instrument_id,
            status="blocked",
            message="Tushare refresh requires a Tushare ts_code identifier such as 018654.OF or 000300.SH.",
            updated_by=updated_by,
            mode="api",
        )

    instrument_type = str(instrument.get("instrument_type") or "").strip().lower()
    suffix = ts_code.rsplit(".", 1)[-1]
    try:
        if instrument_type == "fund" and suffix == "OF":
            latest_date = None if full_history else _latest_nav_date_from_instrument(instrument)
            rows = _call_tushare_api(
                api_name="fund_nav",
                params={"ts_code": ts_code},
                fields="ts_code,ann_date,end_date,nav_date,unit_nav,accum_nav,adj_nav,update_flag",
            )
            nav_rows = _tushare_nav_rows(
                rows,
                instrument_currency=instrument.get("currency"),
                latest_date=latest_date,
            )
            if not nav_rows:
                since_text = f" since {latest_date.isoformat()}" if latest_date else ""
                return update_refresh_status(
                    instrument_id=instrument_id,
                    status="no_new_data",
                    message=f"Tushare fund_nav returned no new NAV rows for {ts_code}{since_text}.",
                    updated_by=updated_by,
                    mode="api",
                )
            record_raw_nav_observations(
                instrument_id=instrument_id,
                rows=nav_rows,
                source_kind="api_observation",
                source_ref="tushare:fund_nav",
                provider="tushare:fund_nav",
                status="complete",
                evidence={"ts_code": ts_code, "source_field": "adj_nav"},
            )
            record, _publication = _publish_fund_nav_projection(
                instrument_id=instrument_id,
                load_source_rows=lambda current_instrument: _load_all_durable_nav_source_rows(
                    instrument_id=instrument_id,
                    instrument=current_instrument,
                ),
                source_provider="tushare:fund_nav",
                refresh_status="imported",
                updated_by=updated_by,
                message_factory=lambda publication: (
                    f"Stored {len(nav_rows)} Tushare NAV observations for {ts_code} "
                    f"and rebuilt {len(publication.rows)} canonical NAV dates; "
                    f"{len(publication.action_candidates)} NAV action signal(s) require review."
                ),
                mode="api",
            )
            return record

        if instrument_type in {"fund", "etf"} and suffix in TUSHARE_PRICE_SUFFIXES:
            return _refresh_tushare_listed_security(
                instrument_id=instrument_id,
                instrument=instrument,
                ts_code=ts_code,
                price_api_name="fund_daily",
                factor_api_name="fund_adj",
                updated_by=updated_by,
                full_history=full_history,
            )

        if instrument_type == "equity" and suffix in TUSHARE_PRICE_SUFFIXES:
            return _refresh_tushare_listed_security(
                instrument_id=instrument_id,
                instrument=instrument,
                ts_code=ts_code,
                price_api_name="daily",
                factor_api_name="adj_factor",
                updated_by=updated_by,
                full_history=full_history,
            )

        if instrument_type == "index" and suffix in TUSHARE_INDEX_SUFFIXES:
            if str(instrument.get("currency") or "").strip().upper() != "CNY":
                raise TushareRefreshError(
                    f"Tushare index price data for {ts_code} requires a CNY instrument."
                )
            latest_date = None if full_history else _latest_market_data_date_from_instrument(
                instrument,
                metric_family="price",
                quote_bases={"close"},
            )
            repair_start = (
                None
                if full_history
                else _price_bar_repair_start_date(
                    instrument_id=instrument_id,
                    instrument=instrument,
                    require_adjustment_factor=False,
                )
            )
            params = {"ts_code": ts_code}
            query_start = (
                repair_start
                if repair_start is not None and not full_history
                else _tushare_query_start_date(
                    latest_date=latest_date,
                    full_history=full_history,
                )
            )
            params["start_date"] = _format_tushare_date(query_start)
            params["end_date"] = _format_tushare_date(date.today())
            rows = _call_tushare_api(
                api_name="index_daily",
                params=params,
                fields=(
                    "ts_code,trade_date,open,high,low,close,pre_close,vol,amount"
                ),
            )
            return _upsert_tushare_price_rows(
                instrument_id=instrument_id,
                ts_code=ts_code,
                api_name="index_daily",
                rows=_tushare_price_rows(
                    rows,
                    latest_date=None if repair_start is not None else latest_date,
                ),
                updated_by=updated_by,
            )
    except TushareRefreshError as exc:
        return update_refresh_status(
            instrument_id=instrument_id,
            status="failed",
            message=f"Tushare refresh failed for {ts_code}: {exc}",
            updated_by=updated_by,
            mode="api",
        )

    return update_refresh_status(
        instrument_id=instrument_id,
        status="blocked",
        message=f"Tushare refresh is not supported for {instrument_type or 'instrument'} code {ts_code}.",
        updated_by=updated_by,
        mode="api",
    )


def _upsert_tushare_price_rows(
    *,
    instrument_id: str,
    ts_code: str,
    api_name: str,
    rows: list[dict[str, object]],
    updated_by: str | None,
) -> dict[str, object] | None:
    if not rows:
        return update_refresh_status(
            instrument_id=instrument_id,
            status="no_new_data",
            message=f"Tushare {api_name} returned no new close rows for {ts_code}.",
            updated_by=updated_by,
            mode="api",
        )

    changed_count = upsert_market_data_points(
        instrument_id=instrument_id,
        rows=[
            {
                "metric_family": "price",
                "quote_basis": "close",
                "as_of_date": row["as_of_date"],
                "value": row["value"],
                "currency": "CNY",
                "provider": f"tushare:{api_name}",
                "status": "complete",
            }
            for row in rows
        ],
    )
    if changed_count is None:
        return None
    bar_changed_count = upsert_price_bars(
        instrument_id=instrument_id,
        rows=_tushare_price_bar_rows(rows, api_name=api_name),
    )
    if bar_changed_count is None:
        return None
    if changed_count == 0 and bar_changed_count == 0:
        return update_refresh_status(
            instrument_id=instrument_id,
            status="no_new_data",
            message=f"Tushare {api_name} returned no changed close rows for {ts_code}.",
            updated_by=updated_by,
            mode="api",
        )
    return update_refresh_status(
        instrument_id=instrument_id,
        status="refreshed",
        message=(
            f"Imported {changed_count} close rows and {bar_changed_count} raw OHLCV bars "
            f"from Tushare {api_name} for {ts_code}."
        ),
        updated_by=updated_by,
        mode="api",
    )


def _matches_batch_source(instrument: dict[str, object], source: str) -> bool:
    source_settings = dict(instrument.get("source_settings", {}))
    source_mode = str(source_settings.get("source_mode") or "manual").strip().lower()
    normalized_source = source.strip().lower()
    if normalized_source == "email":
        return source_mode == "email"
    if normalized_source == "tushare":
        return source_mode == "api" and _tushare_profile_enabled(source_settings)
    if normalized_source == "all":
        return source_mode == "email" or (source_mode == "api" and _tushare_profile_enabled(source_settings))
    return False


def _tushare_refresh_process_worker(
    send_connection: Any,
    instrument_id: str,
    updated_by: str | None,
    full_history: bool,
) -> None:
    """Run one refresh in a killable process, isolating a stuck SDK request."""
    try:
        record = refresh_market_data(
            instrument_id=instrument_id,
            updated_by=updated_by,
            full_history=full_history,
            source="configured",
        )
        send_connection.send({"ok": True, "record": record})
    except BaseException as exc:  # pragma: no cover - defensive process boundary
        send_connection.send(
            {
                "ok": False,
                "error": f"{type(exc).__name__}: {exc}",
            }
        )
    finally:
        send_connection.close()


def _stop_refresh_process(process: multiprocessing.Process) -> None:
    if not process.is_alive():
        process.join(timeout=0.1)
        return
    process.terminate()
    process.join(timeout=2)
    if process.is_alive():
        process.kill()
        process.join(timeout=2)


def _run_tushare_refresh_process_batch(
    *,
    targets: list[dict[str, object]],
    updated_by: str | None,
    full_history: bool,
    settings: Any,
) -> list[dict[str, object] | None]:
    """Refresh Tushare targets concurrently with hard item and batch deadlines.

    Threads cannot be cancelled safely while blocked inside requests. A fresh,
    daemon child process per instrument gives the scheduler a real termination
    boundary and also releases database/network resources if the provider hangs.
    """
    if not targets:
        return []
    context = multiprocessing.get_context("spawn")
    max_workers = min(max(1, int(getattr(settings, "tushare_batch_max_workers", 4))), len(targets))
    configured_item_timeout = int(
        getattr(settings, "market_data_batch_item_timeout_seconds", 300) or 0
    )
    request_timeout = max(1, int(getattr(settings, "tushare_timeout_seconds", 30) or 30))
    item_timeout = (
        configured_item_timeout
        if configured_item_timeout > 0
        else max(60, request_timeout * 4)
    )
    configured_batch_timeout = int(
        getattr(settings, "tushare_batch_timeout_seconds", 3600) or 0
    )
    batch_timeout = (
        configured_batch_timeout
        if configured_batch_timeout > 0
        else item_timeout * ((len(targets) + max_workers - 1) // max_workers) + 60
    )
    batch_started = time.monotonic()
    ordered_results: list[dict[str, object] | None] = [None] * len(targets)
    pending_indexes = list(range(len(targets)))
    running: dict[int, tuple[multiprocessing.Process, Any, float]] = {}

    def failed_record(index: int, message: str) -> dict[str, object] | None:
        instrument_id = str(targets[index].get("instrument_id") or "")
        return update_refresh_status(
            instrument_id=instrument_id,
            status="failed",
            message=message,
            updated_by=updated_by,
            mode="api",
        )

    def start_one(index: int) -> None:
        instrument_id = str(targets[index].get("instrument_id") or "")
        receive_connection, send_connection = context.Pipe(duplex=False)
        process = context.Process(
            target=_tushare_refresh_process_worker,
            args=(send_connection, instrument_id, updated_by, full_history),
            name=f"tushare-refresh-{instrument_id}",
            daemon=True,
        )
        try:
            process.start()
        except Exception as exc:
            receive_connection.close()
            send_connection.close()
            ordered_results[index] = failed_record(
                index,
                f"Unable to start isolated Tushare refresh: {exc}",
            )
            return
        send_connection.close()
        running[index] = (process, receive_connection, time.monotonic())

    while pending_indexes or running:
        now = time.monotonic()
        batch_expired = now - batch_started >= batch_timeout
        if not batch_expired:
            while pending_indexes and len(running) < max_workers:
                start_one(pending_indexes.pop(0))

        completed_indexes: list[int] = []
        for index, (process, receive_connection, started_at) in list(running.items()):
            payload: dict[str, object] | None = None
            try:
                if receive_connection.poll():
                    candidate = receive_connection.recv()
                    payload = candidate if isinstance(candidate, dict) else None
                elif not process.is_alive():
                    payload = {
                        "ok": False,
                        "error": f"worker exited with code {process.exitcode} without a result",
                    }
            except (EOFError, OSError) as exc:
                payload = {"ok": False, "error": f"worker channel failed: {exc}"}

            timed_out = now - started_at >= item_timeout
            if payload is None and not timed_out and not batch_expired:
                continue

            if payload is not None and payload.get("ok") is True:
                record = payload.get("record")
                ordered_results[index] = record if isinstance(record, dict) else None
            else:
                instrument_id = str(targets[index].get("instrument_id") or "")
                if batch_expired:
                    message = (
                        f"Tushare batch deadline exceeded after {batch_timeout} seconds "
                        f"while refreshing {instrument_id}."
                    )
                elif timed_out:
                    message = (
                        f"Tushare refresh timed out after {item_timeout} seconds for "
                        f"{instrument_id}; the isolated worker was terminated."
                    )
                else:
                    message = f"Tushare refresh worker failed for {instrument_id}: {payload.get('error') if payload else 'unknown error'}"
                ordered_results[index] = failed_record(index, message)
            _stop_refresh_process(process)
            receive_connection.close()
            completed_indexes.append(index)

        for index in completed_indexes:
            running.pop(index, None)

        if batch_expired:
            for index in pending_indexes:
                instrument_id = str(targets[index].get("instrument_id") or "")
                ordered_results[index] = failed_record(
                    index,
                    f"Tushare batch deadline exceeded after {batch_timeout} seconds before {instrument_id} could start.",
                )
            pending_indexes.clear()
        if pending_indexes or running:
            time.sleep(0.05)

    return ordered_results


def _email_parser_adapter(
    attachment_name: str,
    attachment_bytes: bytes,
    parser_profile: str,
) -> list[dict[str, object]]:
    return _parse_nav_rows_from_attachment(
        attachment_name=attachment_name,
        attachment_bytes=attachment_bytes,
        parser_profile=parser_profile,
    )


def _email_attachment_text_adapter(
    attachment_name: str,
    attachment_bytes: bytes,
) -> str:
    return _extract_attachment_text(attachment_name, attachment_bytes)


def _email_batch_result_item(
    *,
    instrument: dict[str, object],
    status: str,
    message: str,
) -> dict[str, object]:
    source_settings = dict(instrument.get("source_settings", {}))
    return {
        "instrument_id": instrument["instrument_id"],
        "instrument_name": instrument["instrument_name"],
        "instrument_type": instrument["instrument_type"],
        "source_mode": source_settings.get("source_mode") or "manual",
        "source_api_profile": source_settings.get("source_api_profile") or "",
        "status": status,
        "message": message,
    }


def _run_email_ingestion_batch(
    *,
    instruments: list[dict[str, object]],
    settings: Any,
    updated_by: str | None,
    full_history: bool,
) -> dict[str, object]:
    instruments_by_id = {
        str(instrument.get("instrument_id") or ""): instrument
        for instrument in instruments
    }
    configured_email_ids = {
        instrument_id
        for instrument_id, instrument in instruments_by_id.items()
        if str(dict(instrument.get("source_settings", {})).get("source_mode") or "")
        .strip()
        .lower()
        == "email"
    }
    try:
        ingestion = ingest_email_nav(
            settings=settings,
            instruments=instruments,
            parse_attachment=_email_parser_adapter,
            extract_attachment_text=_email_attachment_text_adapter,
            full_history=full_history,
        )
    except EmailIngestionError as error:
        failure_status = (
            "failed" if isinstance(error, EmailIngestionBusyError) else "blocked"
        )
        results: list[dict[str, object]] = []
        for instrument_id in sorted(configured_email_ids):
            instrument = instruments_by_id[instrument_id]
            refreshed = update_refresh_status(
                instrument_id=instrument_id,
                status=failure_status,
                message=str(error),
                updated_by=updated_by,
                mode="email",
            )
            if refreshed is not None:
                results.append(
                    _email_batch_result_item(
                        instrument=refreshed,
                        status=failure_status,
                        message=str(error),
                    )
                )
        return {
            "source": "email",
            "refreshed_count": 0,
            "skipped_count": len(instruments) - len(configured_email_ids),
            "results": results,
            "folders": [],
        }

    repository = EmailIngestionRepository(get_session_factory())
    results_by_id: dict[str, dict[str, object]] = {}
    boundary_rebuild_ids = list_instrument_ids_with_nav_history_before(
        instrument_ids=configured_email_ids,
        before_date=settings.email_history_start_date,
    )
    publication_ids = set(ingestion.batches).union(boundary_rebuild_ids)
    for instrument_id in sorted(publication_ids):
        batch = ingestion.batches.get(instrument_id)
        instrument = instruments_by_id.get(instrument_id)
        if instrument is None:
            continue
        try:
            source_names = list(dict.fromkeys(batch.attachment_names if batch else []))
            if instrument_id in boundary_rebuild_ids:
                message_factory = lambda publication: (
                    f"Rebuilt {len(publication.rows)} canonical NAV dates within "
                    f"the email history boundary {settings.email_history_start_date.isoformat()}; "
                    f"{len(publication.action_candidates)} NAV action signal(s) require review."
                )
            else:
                message_factory = lambda publication: (
                    f"Rebuilt {len(publication.rows)} canonical NAV dates from "
                    f"{len(source_names)} durable email attachments; "
                    f"{len(publication.action_candidates)} NAV action signal(s) require review."
                )
            record, publication = _publish_fund_nav_projection(
                instrument_id=instrument_id,
                load_source_rows=lambda current_instrument: (
                    _load_all_durable_nav_source_rows(
                        instrument_id=instrument_id,
                        instrument=current_instrument,
                    )
                ),
                source_provider="email",
                refresh_status="imported",
                updated_by=updated_by,
                message_factory=message_factory,
                mode="email",
            )
            message = message_factory(publication)
            repository.mark_imported(batch.route_ids if batch else [])
            results_by_id[instrument_id] = _email_batch_result_item(
                instrument=record,
                status="imported",
                message=message,
            )
        except Exception as error:
            LOGGER.exception(
                "email canonical NAV publication failed instrument_id=%s",
                instrument_id,
            )
            message = f"Email NAV publication failed: {type(error).__name__}: {error}"
            record = update_refresh_status(
                instrument_id=instrument_id,
                status="failed",
                message=message,
                updated_by=updated_by,
                mode="email",
            )
            if record is not None:
                results_by_id[instrument_id] = _email_batch_result_item(
                    instrument=record,
                    status="failed",
                    message=message,
                )

    failed_folder_names = {
        folder.folder_name for folder in ingestion.failed_folders
    }
    for instrument_id in sorted(configured_email_ids - set(results_by_id)):
        instrument = instruments_by_id[instrument_id]
        source_settings = dict(instrument.get("source_settings", {}))
        source_folder = str(source_settings.get("source_location") or "").strip()
        if source_folder in failed_folder_names:
            status = "failed"
            message = f'Configured email folder "{source_folder}" failed to scan.'
        else:
            status = "no_new_data"
            message = "Email scan completed with no new exact-routed NAV observations."
        record = update_refresh_status(
            instrument_id=instrument_id,
            status=status,
            message=message,
            updated_by=updated_by,
            mode="email",
        )
        if record is not None:
            results_by_id[instrument_id] = _email_batch_result_item(
                instrument=record,
                status=status,
                message=message,
            )

    results = list(results_by_id.values())
    for folder in ingestion.failed_folders:
        results.append(
            {
                "instrument_id": f"email-folder:{folder.folder_name}",
                "instrument_name": folder.folder_name,
                "instrument_type": "other",
                "source_mode": "email",
                "source_api_profile": "",
                "status": "failed",
                "message": folder.error,
            }
        )
    return {
        "source": "email",
        "refreshed_count": sum(
            1 for item in results if item["status"] in {"imported", "refreshed"}
        ),
        "skipped_count": max(
            0,
            len(instruments) - len(configured_email_ids.union(ingestion.batches)),
        ),
        "results": results,
        "folders": [
            {
                "folder_name": folder.folder_name,
                "uid_validity": folder.uid_validity,
                "discovered_messages": folder.discovered_messages,
                "fetched_messages": folder.fetched_messages,
                "stored_attachments": folder.stored_attachments,
                "parsed_attachments": folder.parsed_attachments,
                "ignored_messages": folder.ignored_messages,
                "failed_messages": folder.failed_messages,
                "error": folder.error,
            }
            for folder in ingestion.folders
        ],
        "unmatched_candidates": ingestion.unmatched_candidates,
        "ambiguous_candidates": ingestion.ambiguous_candidates,
        "invalid_candidates": ingestion.invalid_candidates,
        "out_of_scope_candidates": ingestion.out_of_scope_candidates,
    }


def refresh_market_data_batch(
    *,
    source: str,
    updated_by: str | None,
    full_history: bool = False,
    include_inactive: bool = False,
) -> dict[str, object]:
    settings = get_settings()
    normalized_source = str(source or "all").strip().lower()
    if normalized_source == "configured":
        normalized_source = "all"
    instruments = list_instruments(include_inactive=include_inactive)
    targets = [item for item in instruments if _matches_batch_source(item, normalized_source)]
    LOGGER.info(
        "market data batch targets source=%s target_count=%s skipped_count=%s",
        normalized_source,
        len(targets),
        len(instruments) - len(targets),
    )
    if normalized_source == "email":
        return _run_email_ingestion_batch(
            instruments=instruments,
            settings=settings,
            updated_by=updated_by,
            full_history=full_history,
        )
    if normalized_source == "all":
        tushare_result = refresh_market_data_batch(
            source="tushare",
            updated_by=updated_by,
            full_history=full_history,
            include_inactive=include_inactive,
        )
        email_result = _run_email_ingestion_batch(
            instruments=instruments,
            settings=settings,
            updated_by=updated_by,
            full_history=full_history,
        )
        combined_results = [
            *list(tushare_result.get("results", [])),
            *list(email_result.get("results", [])),
        ]
        return {
            **email_result,
            "source": "all",
            "refreshed_count": sum(
                1
                for item in combined_results
                if item.get("status") in {"imported", "refreshed"}
            ),
            "skipped_count": min(
                int(tushare_result.get("skipped_count", 0)),
                int(email_result.get("skipped_count", 0)),
            ),
            "results": combined_results,
        }
    if normalized_source == "tushare" and targets:
        max_workers = min(getattr(settings, "tushare_batch_max_workers", 4), len(targets))
        LOGGER.info(
            "tushare batch executing in isolated processes target_count=%s max_workers=%s item_timeout=%s batch_timeout=%s",
            len(targets),
            max_workers,
            getattr(settings, "market_data_batch_item_timeout_seconds", 300),
            getattr(settings, "tushare_batch_timeout_seconds", 3600),
        )
        ordered_results = _run_tushare_refresh_process_batch(
            targets=targets,
            updated_by=updated_by,
            full_history=full_history,
            settings=settings,
        )
        results = []
        for refreshed in ordered_results:
            if refreshed is None:
                continue
            source_settings = dict(refreshed.get("source_settings", {}))
            refresh_status = dict(refreshed.get("refresh_status", {}))
            results.append(
                {
                    "instrument_id": refreshed["instrument_id"],
                    "instrument_name": refreshed["instrument_name"],
                    "instrument_type": refreshed["instrument_type"],
                    "source_mode": source_settings.get("source_mode") or "manual",
                    "source_api_profile": source_settings.get("source_api_profile") or "",
                    "status": refresh_status.get("status") or "idle",
                    "message": refresh_status.get("message") or "",
                }
            )
        return {
            "source": normalized_source,
            "refreshed_count": sum(
                1 for item in results if item["status"] in {"imported", "refreshed"}
            ),
            "skipped_count": len(instruments) - len(targets),
            "results": results,
        }
    return {
        "source": normalized_source,
        "refreshed_count": 0,
        "skipped_count": len(instruments),
        "results": [],
    }
